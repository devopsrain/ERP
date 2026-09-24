"""
Quality Management Routes — /quality

Inspection plans & specifications, the six tender inspection forms
(raw material, cable in-process, wire insulation, final product / certificate
of analysis, wire packing summary, AAC/ABC delivery report), calibration
management, customer complaints, CAPA, NCR, internal audits and the report
suite (HTML + Excel, SPC JSON for Chart.js).
"""
from __future__ import annotations

import json
import logging
import os
import re
import tempfile
from datetime import date, datetime, timedelta
from decimal import Decimal

from fastapi import APIRouter, Depends, Request
from fastapi.responses import FileResponse, JSONResponse, RedirectResponse, Response

from deps import current_company, flash, login_required, require_auth, template_context, validate_upload
from quality_data_store import inspections_for_order, open_quality_issues, quality_store
from quality_forms import (
    AUDIT_RESULTS, AUDIT_STATUSES, CAPA_SOURCES, CAPA_STATUSES, COMPLAINT_STATUSES, COMPLAINT_TYPES,
    EQUIPMENT_STATUSES, GRANULARITIES, INSPECTION_KINDS, INSPECTION_STATUSES, KIND_LABELS, NCR_DISPOSITIONS,
    NCR_ITEM_KINDS, NCR_STATUSES, NCR_TYPES, OVERALL_RESULTS, PRODUCT_DEFECTS, RM_DISPOSITIONS, SPEC_APPLIES_TO,
    SPEC_KINDS, SPEC_STATUSES, default_lines, defective_rate, group_stats, lines_from_form, period_key,
    spc_stats, to_date, to_num, trend_series, yield_stats,
)
from template_engine import templates

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/quality", tags=["quality"])
manager_required = require_auth("manager")

_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOC_MODULE = "quality"
ATTACH_EXTS = {"pdf", "doc", "docx", "xls", "xlsx", "csv", "jpg", "jpeg", "png", "gif", "webp", "zip", "txt"}
ATTACH_KINDS = set(INSPECTION_KINDS) | {"equipment", "complaint", "capa", "ncr", "audit", "spec"}


@router.on_event("startup")
async def _startup():
    quality_store.ensure_schema()


# ── helpers ───────────────────────────────────────────────────────

def _actor(request: Request) -> str:
    return request.session.get("username", "") or ""


def _ctx(request: Request, active_page: str = "", **extra) -> dict:
    ctx = template_context(request)
    ctx.update(
        active_page=active_page, kinds=INSPECTION_KINDS, kind_labels=KIND_LABELS,
        inspection_statuses=INSPECTION_STATUSES, overall_results=OVERALL_RESULTS, spec_kinds=SPEC_KINDS,
        spec_applies_to=SPEC_APPLIES_TO, spec_statuses=SPEC_STATUSES, rm_dispositions=RM_DISPOSITIONS,
        product_defects=PRODUCT_DEFECTS, equipment_statuses=EQUIPMENT_STATUSES, complaint_types=COMPLAINT_TYPES,
        complaint_statuses=COMPLAINT_STATUSES, capa_statuses=CAPA_STATUSES, capa_sources=CAPA_SOURCES,
        ncr_types=NCR_TYPES, ncr_item_kinds=NCR_ITEM_KINDS, ncr_dispositions=NCR_DISPOSITIONS,
        ncr_statuses=NCR_STATUSES, audit_statuses=AUDIT_STATUSES, audit_results=AUDIT_RESULTS,
        granularities=GRANULARITIES, today=date.today(),
    )
    ctx.update(extra)
    return ctx


async def _form(request: Request) -> dict:
    form = await request.form()
    return {k: v for k, v in form.items() if not hasattr(v, "filename")}


def _json_default(o):
    if isinstance(o, Decimal):
        return float(o)
    if isinstance(o, (date, datetime)):
        return o.isoformat()
    return str(o)


def _dates(request: Request, default_days: int = 90):
    qp = request.query_params
    d_to = to_date(qp.get("to")) or date.today()
    d_from = to_date(qp.get("from")) or (d_to - timedelta(days=default_days))
    return d_from, d_to


def _documents(cid: str, entity_kind: str, entity_id: str) -> list:
    try:
        import document_storage as storage
        return storage.documents_for(cid, DOC_MODULE, f"{entity_kind}-{entity_id}")
    except Exception as e:
        logger.debug("quality documents unavailable (%s/%s): %s", entity_kind, entity_id, e)
        return []


def _entity_url(entity_kind: str, entity_id: str) -> str:
    if entity_kind in INSPECTION_KINDS:
        return f"/quality/inspections/{entity_kind}/{entity_id}"
    return {"equipment": f"/quality/calibration/{entity_id}", "complaint": f"/quality/complaints/{entity_id}",
            "capa": f"/quality/capa/{entity_id}", "ncr": f"/quality/ncr/{entity_id}",
            "audit": f"/quality/audits/{entity_id}", "spec": f"/quality/specs/{entity_id}"}.get(entity_kind, "/quality/")


def _xlsx(title: str, headers: list, rows: list, filename: str, subtitle: str = "") -> FileResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\\/*?:\[\]]", "-", title)[:30].strip() or "Report"   # openpyxl sheet-name rules
    ws.append([title]); ws["A1"].font = Font(bold=True, size=13)
    if subtitle:
        ws.append([subtitle])
    ws.append([])
    ws.append(headers)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([(float(v) if isinstance(v, Decimal) else v.isoformat() if isinstance(v, (date, datetime)) else v) for v in r])
    for i, h in enumerate(headers):
        ws.column_dimensions[chr(65 + i) if i < 26 else "A" + chr(65 + i - 26)].width = max(12, min(40, len(str(h)) + 4))
    fd, path = tempfile.mkstemp(suffix=".xlsx"); os.close(fd)
    wb.save(path)
    return FileResponse(path, filename=filename, media_type=_XLSX)


def _fmt(v, nd=3):
    if v is None:
        return ""
    if isinstance(v, (int,)):
        return v
    if isinstance(v, (float, Decimal)):
        return round(float(v), nd)
    return v


def _notify(cid: str, text: str) -> None:
    try:
        from telegram_bot import notify_topic
        notify_topic(cid, "daily_digest", text)
    except Exception:
        pass


# ── dashboard ─────────────────────────────────────────────────────

@router.get("/", name="quality_dashboard")
async def dashboard(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    stats = quality_store.dashboard(cid)
    chart = {"kind_labels": [k["label"] for k in stats["by_kind"]],
             "kind_passed": [k["passed"] for k in stats["by_kind"]],
             "kind_failed": [k["failed"] for k in stats["by_kind"]],
             "kind_other": [k["count"] - k["passed"] - k["failed"] for k in stats["by_kind"]]}
    return templates.TemplateResponse("quality/dashboard.html", _ctx(
        request, "dashboard", stats=stats, chart_json=json.dumps(chart, default=_json_default)))


# ── JSON API for sibling modules / dashboards ─────────────────────

@router.get("/api/open-issues", name="quality_api_open_issues")
async def api_open_issues(request: Request, user=Depends(login_required)):
    return JSONResponse(open_quality_issues(current_company(request)))


@router.get("/api/order/{order_number}", name="quality_api_order_inspections")
async def api_order_inspections(order_number: str, request: Request, user=Depends(login_required)):
    data = inspections_for_order(current_company(request), order_number)
    return JSONResponse(json.loads(json.dumps(data, default=_json_default)))


# ── attachments (photos, test results, calibration certificates) ──

@router.post("/attach/{entity_kind}/{entity_id}", name="quality_attach_upload")
async def attach_upload(entity_kind: str, entity_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    back = _entity_url(entity_kind, entity_id) + "#files"
    if entity_kind not in ATTACH_KINDS or not quality_store.entity_exists(entity_kind, entity_id, cid):
        flash(request, "Record not found", "error")
        return RedirectResponse("/quality/", status_code=303)
    form = await request.form()
    upload = form.get("file")
    tag = (form.get("kind") or "attachment").strip()[:40]
    if upload is None or not getattr(upload, "filename", ""):
        flash(request, "No file selected", "error")
        return RedirectResponse(back, status_code=303)
    content = await upload.read()
    ok, err = validate_upload(upload.filename, content, allowed_exts=ATTACH_EXTS)
    if not ok:
        flash(request, err, "error")
        return RedirectResponse(back, status_code=303)
    try:
        import document_storage as storage
        doc = storage.store_document(cid, DOC_MODULE, f"{entity_kind}-{entity_id}", upload.filename, content,
                                     content_type=getattr(upload, "content_type", None), uploaded_by=_actor(request),
                                     tags=[tag, entity_kind], ip=(request.client.host if request.client else ""))
        if entity_kind == "equipment" and tag == "certificate":
            quality_store.set_equipment_certificate(entity_id, cid, doc["id"])
        flash(request, "File attached" + (" (stored locally — Nextcloud unavailable)" if doc.get("fallback_error") else ""), "success")
    except Exception as e:
        logger.error("quality upload failed: %s", e)
        flash(request, f"Upload failed: {e}", "error")
    return RedirectResponse(back, status_code=303)


@router.get("/attach/{entity_kind}/{entity_id}/{doc_id}", name="quality_attach_file")
async def attach_file(entity_kind: str, entity_id: str, doc_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    try:
        import document_storage as storage
        doc = storage.get_document(doc_id, cid)
        if doc.get("module") != DOC_MODULE or doc.get("entity_id") != storage.safe_segment(f"{entity_kind}-{entity_id}") or doc.get("deleted_at"):
            raise storage.DocumentNotFound("Document not found")
        data, filename, ctype = storage.open_document(doc_id, cid, by=_actor(request),
                                                      ip=(request.client.host if request.client else ""))
        inline = request.query_params.get("download") != "1"
        return Response(content=data, media_type=ctype, headers={
            "Content-Disposition": storage.content_disposition(filename, inline=inline),
            "Content-Length": str(len(data)), "X-Content-Type-Options": "nosniff", "Cache-Control": "private, no-store"})
    except Exception as e:
        flash(request, f"Could not open file: {e}", "error")
        return RedirectResponse(_entity_url(entity_kind, entity_id) + "#files", status_code=303)


@router.post("/attach/{entity_kind}/{entity_id}/{doc_id}/delete", name="quality_attach_delete")
async def attach_delete(entity_kind: str, entity_id: str, doc_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    try:
        import document_storage as storage
        doc = storage.get_document(doc_id, cid)
        if doc.get("module") != DOC_MODULE or doc.get("entity_id") != storage.safe_segment(f"{entity_kind}-{entity_id}"):
            raise storage.DocumentNotFound("Document not found")
        ok = storage.delete_document(doc_id, by=_actor(request), company_id=cid)
        flash(request, "File removed" if ok else "Could not remove file", "success" if ok else "error")
    except Exception as e:
        flash(request, f"Could not remove file: {e}", "error")
    return RedirectResponse(_entity_url(entity_kind, entity_id) + "#files", status_code=303)


# ── inspection plans / specifications ─────────────────────────────

@router.get("/specs", name="quality_spec_list")
async def spec_list(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    qp = request.query_params
    return templates.TemplateResponse("quality/spec_list.html", _ctx(
        request, "specs", specs=quality_store.list_spec_sets(cid, qp.get("applies_to") or None, qp.get("status") or None),
        applies_filter=qp.get("applies_to") or "", status_filter=qp.get("status") or ""))


@router.get("/specs/new", name="quality_spec_new_get")
async def spec_new_get(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    return templates.TemplateResponse("quality/spec_form.html", _ctx(
        request, "specs", spec={"applies_to": request.query_params.get("applies_to") or "final", "version": "1", "status": "draft"},
        is_edit=False, lookups=quality_store.lookups(cid)))


@router.post("/specs/new", name="quality_spec_new_post")
async def spec_new_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    data["created_by"] = _actor(request)
    if not (data.get("name") or "").strip():
        flash(request, "Specification name is required", "error")
        return RedirectResponse("/quality/specs/new", status_code=303)
    s = quality_store.create_spec_set(cid, data)
    if s:
        flash(request, "Specification created — add its parameters below", "success")
        return RedirectResponse(f"/quality/specs/{s['id']}", status_code=303)
    flash(request, "Failed to create specification", "error")
    return RedirectResponse("/quality/specs/new", status_code=303)


@router.get("/specs/{spec_id}", name="quality_spec_detail")
async def spec_detail(spec_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    s = quality_store.get_spec_set(spec_id, cid)
    if not s:
        flash(request, "Specification not found", "error")
        return RedirectResponse("/quality/specs", status_code=303)
    return templates.TemplateResponse("quality/spec_detail.html", _ctx(
        request, "specs", spec=s, documents=_documents(cid, "spec", spec_id), lookups=quality_store.lookups(cid)))


@router.post("/specs/{spec_id}/edit", name="quality_spec_edit")
async def spec_edit(spec_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    ok = quality_store.update_spec_set(spec_id, cid, data)
    flash(request, "Specification updated" if ok else "Failed to update specification", "success" if ok else "error")
    return RedirectResponse(f"/quality/specs/{spec_id}", status_code=303)


@router.post("/specs/{spec_id}/status", name="quality_spec_status")
async def spec_status(spec_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    data = await _form(request)
    status = data.get("status", "")
    if status not in SPEC_STATUSES:
        flash(request, "Invalid status", "error")
    else:
        ok = quality_store.set_spec_status(spec_id, cid, status)
        flash(request, f"Specification marked {status}" if ok else "Failed to update status", "success" if ok else "error")
    return RedirectResponse(f"/quality/specs/{spec_id}", status_code=303)


@router.post("/specs/{spec_id}/params", name="quality_spec_param_add")
async def spec_param_add(spec_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    if not quality_store.get_spec_set(spec_id, cid):
        flash(request, "Specification not found", "error")
        return RedirectResponse("/quality/specs", status_code=303)
    if not (data.get("parameter") or "").strip():
        flash(request, "Parameter name is required", "error")
    elif quality_store.add_spec_param(spec_id, cid, data):
        flash(request, "Parameter added", "success")
    else:
        flash(request, "Failed to add parameter", "error")
    return RedirectResponse(f"/quality/specs/{spec_id}#params", status_code=303)


@router.post("/specs/{spec_id}/params/{param_id}/delete", name="quality_spec_param_delete")
async def spec_param_delete(spec_id: str, param_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    ok = quality_store.delete_spec_param(param_id, cid)
    flash(request, "Parameter removed" if ok else "Failed to remove parameter", "success" if ok else "error")
    return RedirectResponse(f"/quality/specs/{spec_id}#params", status_code=303)


# ── inspections (generic over kind) ───────────────────────────────

def _kind_or_none(kind: str):
    return kind if kind in INSPECTION_KINDS else None


def _inspection_list_ctx(request: Request, cid: str, kind: str) -> dict:
    qp = request.query_params
    return dict(kind=kind, cfg=INSPECTION_KINDS[kind],
                inspections=quality_store.list_inspections(kind, cid, qp.get("status") or None, qp.get("result") or None,
                                                           (qp.get("q") or "").strip() or None,
                                                           to_date(qp.get("from")), to_date(qp.get("to"))),
                status_filter=qp.get("status") or "", result_filter=qp.get("result") or "",
                search=qp.get("q") or "", date_from=qp.get("from") or "", date_to=qp.get("to") or "")


@router.get("/inspections/{kind}", name="quality_inspection_list")
async def inspection_list(kind: str, request: Request, user=Depends(login_required)):
    if not _kind_or_none(kind):
        flash(request, "Unknown inspection type", "error")
        return RedirectResponse("/quality/", status_code=303)
    cid = current_company(request)
    return templates.TemplateResponse("quality/inspection_list.html", _ctx(request, f"insp_{kind}", **_inspection_list_ctx(request, cid, kind)))


@router.get("/inspections/{kind}/export.xlsx", name="quality_inspection_export")
async def inspection_export(kind: str, request: Request, user=Depends(login_required)):
    if not _kind_or_none(kind):
        return RedirectResponse("/quality/", status_code=303)
    cid = current_company(request)
    c = _inspection_list_ctx(request, cid, kind)
    cfg = INSPECTION_KINDS[kind]
    cols = ["ref_no", "inspection_date"] + [f["name"] for f in cfg["fields"]] + ["overall_result", "status"] + list(cfg["signatures"])
    labels = ["Ref no.", "Date"] + [f["label"] for f in cfg["fields"]] + ["Result", "Status"] + [s.replace("_", " ").capitalize() for s in cfg["signatures"]]
    rows = [[_fmt(r.get(c)) for c in cols] for r in c["inspections"]]
    return _xlsx(cfg["label"], labels, rows, f"quality_{kind}_{date.today().isoformat()}.xlsx")


@router.get("/inspections/{kind}/new", name="quality_inspection_new_get")
async def inspection_new_get(kind: str, request: Request, user=Depends(login_required)):
    if not _kind_or_none(kind):
        flash(request, "Unknown inspection type", "error")
        return RedirectResponse("/quality/", status_code=303)
    cid = current_company(request)
    cfg = INSPECTION_KINDS[kind]
    qp = request.query_params
    spec_id = qp.get("spec_set_id") or ""
    spec = quality_store.get_spec_set(spec_id, cid) if spec_id else None
    insp = {"inspection_date": date.today().isoformat(), "spec_set_id": spec_id, "prepared_by": _actor(request)}
    for f in cfg["fields"]:  # prefill from query (e.g. links from manufacturing / procurement)
        if qp.get(f["name"]):
            insp[f["name"]] = qp.get(f["name"])
    return templates.TemplateResponse("quality/inspection_form.html", _ctx(
        request, f"insp_{kind}", kind=kind, cfg=cfg, insp=insp, is_edit=False,
        lines=default_lines(kind, spec["params"] if spec else None),
        specs=quality_store.list_spec_sets(cid, cfg["applies_to"], "active"), lookups=quality_store.lookups(cid)))


@router.post("/inspections/{kind}/new", name="quality_inspection_new_post")
async def inspection_new_post(kind: str, request: Request, user=Depends(login_required)):
    if not _kind_or_none(kind):
        return RedirectResponse("/quality/", status_code=303)
    cid = current_company(request)
    data = await _form(request)
    insp = quality_store.create_inspection(kind, cid, data, lines_from_form(data), _actor(request))
    if insp:
        msg = f"{INSPECTION_KINDS[kind]['label']} {insp['ref_no']} saved — result: {insp['overall_result']}"
        flash(request, msg, "success" if insp["overall_result"] != "fail" else "error")
        if insp["overall_result"] == "fail":
            _notify(cid, f"⚠️ Quality: {INSPECTION_KINDS[kind]['label']} {insp['ref_no']} FAILED"
                         + (f" (order {insp.get('order_number')})" if insp.get("order_number") else ""))
        return RedirectResponse(f"/quality/inspections/{kind}/{insp['id']}", status_code=303)
    flash(request, "Failed to save inspection", "error")
    return RedirectResponse(f"/quality/inspections/{kind}/new", status_code=303)


@router.get("/inspections/final/{inspection_id}/certificate", name="quality_certificate")
async def certificate(inspection_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    insp = quality_store.get_inspection("final", inspection_id, cid)
    if not insp:
        flash(request, "Inspection not found", "error")
        return RedirectResponse("/quality/inspections/final", status_code=303)
    company = None
    try:
        from tenant_data_store import tenant_store
        company = tenant_store.get_tenant(cid)
    except Exception:
        company = None
    return templates.TemplateResponse("quality/certificate.html", _ctx(request, "insp_final", insp=insp, company=company or {}))


@router.get("/inspections/{kind}/{inspection_id}", name="quality_inspection_detail")
async def inspection_detail(kind: str, inspection_id: str, request: Request, user=Depends(login_required)):
    if not _kind_or_none(kind):
        return RedirectResponse("/quality/", status_code=303)
    cid = current_company(request)
    insp = quality_store.get_inspection(kind, inspection_id, cid)
    if not insp:
        flash(request, "Inspection not found", "error")
        return RedirectResponse(f"/quality/inspections/{kind}", status_code=303)
    ncrs = [n for n in quality_store.list_ncrs(cid) if n.get("source_inspection_id") == inspection_id]
    return templates.TemplateResponse("quality/inspection_detail.html", _ctx(
        request, f"insp_{kind}", kind=kind, cfg=INSPECTION_KINDS[kind], insp=insp,
        documents=_documents(cid, kind, inspection_id), linked_ncrs=ncrs,
        yield_info=yield_stats(insp.get("input_length_m"), insp.get("total_length_m"), insp.get("scrap_kg"), insp.get("under_length_m")) if kind == "packing" else None))


@router.get("/inspections/{kind}/{inspection_id}/edit", name="quality_inspection_edit_get")
async def inspection_edit_get(kind: str, inspection_id: str, request: Request, user=Depends(login_required)):
    if not _kind_or_none(kind):
        return RedirectResponse("/quality/", status_code=303)
    cid = current_company(request)
    insp = quality_store.get_inspection(kind, inspection_id, cid)
    if not insp:
        flash(request, "Inspection not found", "error")
        return RedirectResponse(f"/quality/inspections/{kind}", status_code=303)
    if insp["status"] == "approved":
        flash(request, "Approved inspections cannot be edited", "error")
        return RedirectResponse(f"/quality/inspections/{kind}/{inspection_id}", status_code=303)
    cfg = INSPECTION_KINDS[kind]
    return templates.TemplateResponse("quality/inspection_form.html", _ctx(
        request, f"insp_{kind}", kind=kind, cfg=cfg, insp=insp, is_edit=True, lines=insp["lines"],
        specs=quality_store.list_spec_sets(cid, cfg["applies_to"], "active"), lookups=quality_store.lookups(cid)))


@router.post("/inspections/{kind}/{inspection_id}/edit", name="quality_inspection_edit_post")
async def inspection_edit_post(kind: str, inspection_id: str, request: Request, user=Depends(login_required)):
    if not _kind_or_none(kind):
        return RedirectResponse("/quality/", status_code=303)
    cid = current_company(request)
    data = await _form(request)
    ok = quality_store.update_inspection(kind, inspection_id, cid, data, lines_from_form(data))
    flash(request, "Inspection updated" if ok else "Failed to update inspection (approved records are read-only)",
          "success" if ok else "error")
    return RedirectResponse(f"/quality/inspections/{kind}/{inspection_id}", status_code=303)


@router.post("/inspections/{kind}/{inspection_id}/submit", name="quality_inspection_submit")
async def inspection_submit(kind: str, inspection_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    ok = _kind_or_none(kind) and quality_store.set_inspection_status(kind, inspection_id, cid, "submitted", _actor(request))
    flash(request, "Inspection submitted for approval" if ok else "Failed to submit", "success" if ok else "error")
    return RedirectResponse(f"/quality/inspections/{kind}/{inspection_id}", status_code=303)


@router.post("/inspections/{kind}/{inspection_id}/approve", name="quality_inspection_approve")
async def inspection_approve(kind: str, inspection_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    data = await _form(request)
    target = "draft" if data.get("action") == "reopen" else "approved"
    ok = _kind_or_none(kind) and quality_store.set_inspection_status(kind, inspection_id, cid, target, _actor(request))
    flash(request, ("Inspection approved" if target == "approved" else "Inspection reopened") if ok else "Failed to update status",
          "success" if ok else "error")
    return RedirectResponse(f"/quality/inspections/{kind}/{inspection_id}", status_code=303)


# ── calibration ───────────────────────────────────────────────────

@router.get("/calibration", name="quality_equipment_list")
async def equipment_list(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    qp = request.query_params
    rows = quality_store.list_equipment(cid, qp.get("status") or None, (qp.get("q") or "").strip() or None)
    if qp.get("export") == "xlsx":
        return _xlsx("Calibration register", ["Tag", "Equipment", "Location", "Frequency (months)", "Last calibration",
                                              "Next due", "Status", "Performed by", "Provider", "Certificate", "Standard"],
                     [[r["equipment_tag"], r["equipment_name"], r["location"], r["calibration_frequency_months"],
                       r.get("last_calibration_date"), r.get("next_due_date"), r["status"], r["performed_by"],
                       r["provider"], r["certificate_number"], r["calibration_standard"]] for r in rows],
                     f"calibration_register_{date.today().isoformat()}.xlsx")
    counts = {s: sum(1 for r in quality_store.list_equipment(cid) if r["status"] == s) for s in EQUIPMENT_STATUSES}
    return templates.TemplateResponse("quality/equipment_list.html", _ctx(
        request, "calibration", equipment=rows, status_filter=qp.get("status") or "", search=qp.get("q") or "", counts=counts))


@router.get("/calibration/new", name="quality_equipment_new_get")
async def equipment_new_get(request: Request, user=Depends(login_required)):
    return templates.TemplateResponse("quality/equipment_form.html", _ctx(
        request, "calibration", equipment={"calibration_frequency_months": 12, "performed_by": "external"}, is_edit=False))


@router.post("/calibration/new", name="quality_equipment_new_post")
async def equipment_new_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    data["recorded_by"] = _actor(request)
    e = quality_store.create_equipment(cid, data)
    if e:
        flash(request, f"Equipment {e['equipment_tag']} registered", "success")
        return RedirectResponse(f"/quality/calibration/{e['id']}", status_code=303)
    flash(request, "Failed to register equipment — name and tag are required and the tag must be unique", "error")
    return RedirectResponse("/quality/calibration/new", status_code=303)


@router.get("/calibration/{equipment_id}", name="quality_equipment_detail")
async def equipment_detail(equipment_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    e = quality_store.get_equipment(equipment_id, cid)
    if not e:
        flash(request, "Equipment not found", "error")
        return RedirectResponse("/quality/calibration", status_code=303)
    return templates.TemplateResponse("quality/equipment_detail.html", _ctx(
        request, "calibration", equipment=e, documents=_documents(cid, "equipment", equipment_id)))


@router.get("/calibration/{equipment_id}/edit", name="quality_equipment_edit_get")
async def equipment_edit_get(equipment_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    e = quality_store.get_equipment(equipment_id, cid)
    if not e:
        flash(request, "Equipment not found", "error")
        return RedirectResponse("/quality/calibration", status_code=303)
    return templates.TemplateResponse("quality/equipment_form.html", _ctx(request, "calibration", equipment=e, is_edit=True))


@router.post("/calibration/{equipment_id}/edit", name="quality_equipment_edit_post")
async def equipment_edit_post(equipment_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    ok = quality_store.update_equipment(equipment_id, cid, data)
    flash(request, "Equipment updated" if ok else "Failed to update equipment", "success" if ok else "error")
    return RedirectResponse(f"/quality/calibration/{equipment_id}", status_code=303)


@router.post("/calibration/{equipment_id}/calibrate", name="quality_equipment_calibrate")
async def equipment_calibrate(equipment_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    rec = quality_store.add_calibration(equipment_id, cid, data, _actor(request))
    flash(request, "Calibration recorded" if rec else "Failed to record calibration", "success" if rec else "error")
    return RedirectResponse(f"/quality/calibration/{equipment_id}", status_code=303)


# ── complaints ────────────────────────────────────────────────────

@router.get("/complaints", name="quality_complaint_list")
async def complaint_list(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    qp = request.query_params
    rows = quality_store.list_complaints(cid, qp.get("status") or None, (qp.get("q") or "").strip() or None,
                                         to_date(qp.get("from")), to_date(qp.get("to")))
    if qp.get("export") == "xlsx":
        return _xlsx("Customer complaints", ["Complaint no.", "Date received", "Customer", "Product", "Batch", "Type",
                                             "Status", "Responsible dept.", "Root cause", "Linked CAPA", "Closed"],
                     [[r["complaint_no"], r["date_received"], r["customer_name"], r["product_details"], r["batch_number"],
                       r["complaint_type"], r["status"], r["responsible_department"], r["root_cause_analysis"],
                       r.get("linked_capa_id") or "", r.get("closed_at")] for r in rows],
                     f"complaints_{date.today().isoformat()}.xlsx")
    return templates.TemplateResponse("quality/complaint_list.html", _ctx(
        request, "complaints", complaints=rows, status_filter=qp.get("status") or "", search=qp.get("q") or "",
        date_from=qp.get("from") or "", date_to=qp.get("to") or ""))


@router.get("/complaints/new", name="quality_complaint_new_get")
async def complaint_new_get(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    return templates.TemplateResponse("quality/complaint_form.html", _ctx(
        request, "complaints", complaint={"date_received": date.today().isoformat(), "complaint_type": "other"},
        is_edit=False, lookups=quality_store.lookups(cid)))


@router.post("/complaints/new", name="quality_complaint_new_post")
async def complaint_new_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    c = quality_store.create_complaint(cid, data, _actor(request))
    if c:
        flash(request, f"Complaint {c['complaint_no']} registered", "success")
        return RedirectResponse(f"/quality/complaints/{c['id']}", status_code=303)
    flash(request, "Failed to register complaint", "error")
    return RedirectResponse("/quality/complaints/new", status_code=303)


@router.get("/complaints/{complaint_id}", name="quality_complaint_detail")
async def complaint_detail(complaint_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    c = quality_store.get_complaint(complaint_id, cid)
    if not c:
        flash(request, "Complaint not found", "error")
        return RedirectResponse("/quality/complaints", status_code=303)
    capa = quality_store.get_capa(c["linked_capa_id"], cid) if c.get("linked_capa_id") else None
    return templates.TemplateResponse("quality/complaint_detail.html", _ctx(
        request, "complaints", complaint=c, capa=capa, documents=_documents(cid, "complaint", complaint_id),
        lookups=quality_store.lookups(cid)))


@router.post("/complaints/{complaint_id}/edit", name="quality_complaint_edit")
async def complaint_edit(complaint_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    ok = quality_store.update_complaint(complaint_id, cid, data)
    flash(request, "Complaint updated" if ok else "Failed to update complaint (closed records are read-only)", "success" if ok else "error")
    return RedirectResponse(f"/quality/complaints/{complaint_id}", status_code=303)


@router.post("/complaints/{complaint_id}/close", name="quality_complaint_close")
async def complaint_close(complaint_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    data = await _form(request)
    ok = quality_store.close_complaint(complaint_id, cid, _actor(request), data.get("resolution") or "")
    flash(request, "Complaint closed" if ok else "Failed to close complaint", "success" if ok else "error")
    return RedirectResponse(f"/quality/complaints/{complaint_id}", status_code=303)


@router.post("/complaints/{complaint_id}/raise-capa", name="quality_complaint_raise_capa")
async def complaint_raise_capa(complaint_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    c = quality_store.get_complaint(complaint_id, cid)
    if not c:
        flash(request, "Complaint not found", "error")
        return RedirectResponse("/quality/complaints", status_code=303)
    capa = quality_store.create_capa(cid, {
        "problem_description": f"Customer complaint {c['complaint_no']} ({c['complaint_type']}) — {c['customer_name']}: {c['description'][:500]}",
        "possible_cause": c.get("root_cause_analysis") or "", "source": "complaint", "source_id": complaint_id,
        "responsible_person": c.get("responsible_department") or "", "target_date": (date.today() + timedelta(days=30)).isoformat(),
    }, _actor(request))
    if capa:
        flash(request, f"CAPA {capa['capa_no']} raised from complaint", "success")
        return RedirectResponse(f"/quality/capa/{capa['id']}", status_code=303)
    flash(request, "Failed to raise CAPA", "error")
    return RedirectResponse(f"/quality/complaints/{complaint_id}", status_code=303)


# ── CAPA ──────────────────────────────────────────────────────────

def _capa_filters(request: Request) -> dict:
    qp = request.query_params
    return dict(status=qp.get("status") or None, responsible=(qp.get("responsible") or "").strip() or None,
                requested_by=(qp.get("requested_by") or "").strip() or None, source=qp.get("source") or None,
                q=(qp.get("q") or "").strip() or None, date_from=to_date(qp.get("from")), date_to=to_date(qp.get("to")))


@router.get("/capa", name="quality_capa_list")
async def capa_list(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    f = _capa_filters(request)
    rows = quality_store.list_capa(cid, **f)
    qp = request.query_params
    if qp.get("export") == "xlsx":
        return _xlsx("CAPA register", ["CAPA no.", "Initiated", "Source", "Type", "Problem", "Proposed action", "Responsible",
                                       "Requested by", "Target date", "Completion", "Status", "Days overdue", "Approved by"],
                     [[r["capa_no"], r["initiated_date"], r["source"], r["action_type"], r["problem_description"],
                       r["proposed_action"], r["responsible_person"], r["requested_by"], r.get("target_date"),
                       r.get("completion_date"), r["effective_status"], r["days_overdue"], r["approved_by"]] for r in rows],
                     f"capa_register_{date.today().isoformat()}.xlsx")
    counts: dict = {}
    for r in quality_store.list_capa(cid):
        counts[r["effective_status"]] = counts.get(r["effective_status"], 0) + 1
    return templates.TemplateResponse("quality/capa_list.html", _ctx(
        request, "capa", capas=rows, counts=counts, status_filter=qp.get("status") or "", source_filter=qp.get("source") or "",
        responsible=qp.get("responsible") or "", requested_by=qp.get("requested_by") or "", search=qp.get("q") or "",
        date_from=qp.get("from") or "", date_to=qp.get("to") or ""))


@router.get("/capa/new", name="quality_capa_new_get")
async def capa_new_get(request: Request, user=Depends(login_required)):
    qp = request.query_params
    capa = {"initiated_date": date.today().isoformat(), "requested_by": _actor(request), "nc_kind": "actual",
            "action_type": "corrective", "source": qp.get("source") or "other", "source_id": qp.get("source_id") or "",
            "problem_description": qp.get("problem_description") or "",
            "target_date": (date.today() + timedelta(days=30)).isoformat()}
    return templates.TemplateResponse("quality/capa_form.html", _ctx(request, "capa", capa=capa, is_edit=False))


@router.post("/capa/new", name="quality_capa_new_post")
async def capa_new_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    if not (data.get("problem_description") or "").strip():
        flash(request, "Problem description is required", "error")
        return RedirectResponse("/quality/capa/new", status_code=303)
    c = quality_store.create_capa(cid, data, _actor(request))
    if c:
        flash(request, f"CAPA {c['capa_no']} opened", "success")
        return RedirectResponse(f"/quality/capa/{c['id']}", status_code=303)
    flash(request, "Failed to open CAPA", "error")
    return RedirectResponse("/quality/capa/new", status_code=303)


@router.get("/capa/reminders", name="quality_capa_reminders")
async def capa_reminders(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    rep = quality_store.capa_reminder_report(cid)
    if request.query_params.get("export") == "xlsx":
        rows = [[r["capa_no"], r["effective_status"], r["requested_by"], r["responsible_person"], r.get("target_date"),
                 r["days_overdue"], r["problem_description"][:200]] for r in rep["past_due"] + rep["pending"]]
        return _xlsx("CAPA reminders", ["CAPA no.", "Status", "Requested by", "Assignee", "Target date", "Days overdue", "Problem"],
                     rows, f"capa_reminders_{date.today().isoformat()}.xlsx")
    return templates.TemplateResponse("quality/capa_reminders.html", _ctx(request, "capa_reminders", report=rep))


@router.get("/capa/{capa_id}", name="quality_capa_detail")
async def capa_detail(capa_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    c = quality_store.get_capa(capa_id, cid)
    if not c:
        flash(request, "CAPA not found", "error")
        return RedirectResponse("/quality/capa", status_code=303)
    approval = None
    try:
        from approval_hooks import approval_status
        approval = approval_status("capa", capa_id, cid)
    except Exception:
        approval = None
    source = None
    if c.get("source") == "complaint" and c.get("source_id"):
        source = quality_store.get_complaint(c["source_id"], cid)
    elif c.get("source") == "ncr" and c.get("source_id"):
        source = quality_store.get_ncr(c["source_id"], cid)
    return templates.TemplateResponse("quality/capa_detail.html", _ctx(
        request, "capa", capa=c, approval=approval, source=source, documents=_documents(cid, "capa", capa_id)))


@router.post("/capa/{capa_id}/edit", name="quality_capa_edit")
async def capa_edit(capa_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    ok = quality_store.update_capa(capa_id, cid, data)
    flash(request, "CAPA updated" if ok else "Failed to update CAPA (closed records are read-only)", "success" if ok else "error")
    return RedirectResponse(f"/quality/capa/{capa_id}", status_code=303)


@router.post("/capa/{capa_id}/request-approval", name="quality_capa_request_approval")
async def capa_request_approval(capa_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    c = quality_store.get_capa(capa_id, cid)
    if not c:
        flash(request, "CAPA not found", "error")
        return RedirectResponse("/quality/capa", status_code=303)
    try:
        from approval_hooks import request_approval
        rid = request_approval(cid, "capa", capa_id, f"CAPA {c['capa_no']}: {c['problem_description'][:80]}", 0,
                               _actor(request), payload={"capa_no": c["capa_no"], "responsible": c["responsible_person"]})
    except Exception:
        rid = None
    flash(request, "Approval requested through the approval workflow" if rid else
          "No approval workflow is configured for CAPA — a manager can approve directly below", "success" if rid else "info")
    return RedirectResponse(f"/quality/capa/{capa_id}", status_code=303)


@router.post("/capa/{capa_id}/approve", name="quality_capa_approve")
async def capa_approve(capa_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    ok = quality_store.approve_capa(capa_id, cid, _actor(request))
    flash(request, "CAPA approved" if ok else "Failed to approve CAPA", "success" if ok else "error")
    return RedirectResponse(f"/quality/capa/{capa_id}", status_code=303)


@router.post("/capa/{capa_id}/close", name="quality_capa_close")
async def capa_close(capa_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    data = await _form(request)
    ok = quality_store.close_capa(capa_id, cid, _actor(request), data.get("effectiveness_check") or "")
    flash(request, "CAPA closed" if ok else "Failed to close CAPA", "success" if ok else "error")
    return RedirectResponse(f"/quality/capa/{capa_id}", status_code=303)


# ── NCR ───────────────────────────────────────────────────────────

@router.get("/ncr", name="quality_ncr_list")
async def ncr_list(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    qp = request.query_params
    rows = quality_store.list_ncrs(cid, qp.get("status") or None, qp.get("item_kind") or None,
                                   (qp.get("q") or "").strip() or None, to_date(qp.get("from")), to_date(qp.get("to")))
    if qp.get("export") == "xlsx":
        return _xlsx("NCR register", ["NCR no.", "Date", "Type", "Item kind", "Product / material", "Lot", "Location",
                                      "Description", "Inspector", "Disposition", "Status", "Linked CAPA"],
                     [[r["ncr_no"], r["date_of_inspection"], r["ncr_type"], r["item_kind"], r["product_or_material"],
                       r["lot_no"], r["location"], r["description"], r["inspector_name"], r.get("disposition") or "",
                       r["status"], r.get("linked_capa_id") or ""] for r in rows],
                     f"ncr_register_{date.today().isoformat()}.xlsx")
    return templates.TemplateResponse("quality/ncr_list.html", _ctx(
        request, "ncr", ncrs=rows, status_filter=qp.get("status") or "", item_kind_filter=qp.get("item_kind") or "",
        search=qp.get("q") or "", date_from=qp.get("from") or "", date_to=qp.get("to") or ""))


@router.get("/ncr/new", name="quality_ncr_new_get")
async def ncr_new_get(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    qp = request.query_params
    ncr = {"date_of_inspection": date.today().isoformat(), "ncr_type": "inspection", "item_kind": "finished_good",
           "inspector_name": _actor(request)}
    if qp.get("kind") in INSPECTION_KINDS and qp.get("inspection_id"):
        ncr.update(quality_store.ncr_prefill_from_inspection(qp.get("kind"), qp.get("inspection_id"), cid))
    return templates.TemplateResponse("quality/ncr_form.html", _ctx(request, "ncr", ncr=ncr, is_edit=False))


@router.post("/ncr/new", name="quality_ncr_new_post")
async def ncr_new_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    if not (data.get("description") or "").strip():
        flash(request, "Description is required", "error")
        return RedirectResponse("/quality/ncr/new", status_code=303)
    n = quality_store.create_ncr(cid, data, _actor(request))
    if n:
        flash(request, f"NCR {n['ncr_no']} raised", "success")
        _notify(cid, f"🚫 Quality: NCR {n['ncr_no']} raised — {n['product_or_material']} ({n['item_kind']})")
        return RedirectResponse(f"/quality/ncr/{n['id']}", status_code=303)
    flash(request, "Failed to raise NCR", "error")
    return RedirectResponse("/quality/ncr/new", status_code=303)


@router.get("/ncr/{ncr_id}", name="quality_ncr_detail")
async def ncr_detail(ncr_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    n = quality_store.get_ncr(ncr_id, cid)
    if not n:
        flash(request, "NCR not found", "error")
        return RedirectResponse("/quality/ncr", status_code=303)
    capa = quality_store.get_capa(n["linked_capa_id"], cid) if n.get("linked_capa_id") else None
    return templates.TemplateResponse("quality/ncr_detail.html", _ctx(
        request, "ncr", ncr=n, capa=capa, documents=_documents(cid, "ncr", ncr_id)))


@router.post("/ncr/{ncr_id}/edit", name="quality_ncr_edit")
async def ncr_edit(ncr_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    ok = quality_store.update_ncr(ncr_id, cid, data)
    flash(request, "NCR updated" if ok else "Failed to update NCR (closed records are read-only)", "success" if ok else "error")
    return RedirectResponse(f"/quality/ncr/{ncr_id}", status_code=303)


@router.post("/ncr/{ncr_id}/disposition", name="quality_ncr_disposition")
async def ncr_disposition(ncr_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    data = await _form(request)
    ok = quality_store.disposition_ncr(ncr_id, cid, data.get("disposition", ""), _actor(request), data.get("note") or "")
    flash(request, "Disposition recorded" if ok else "Failed to record disposition", "success" if ok else "error")
    return RedirectResponse(f"/quality/ncr/{ncr_id}", status_code=303)


@router.post("/ncr/{ncr_id}/close", name="quality_ncr_close")
async def ncr_close(ncr_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    ok = quality_store.close_ncr(ncr_id, cid, _actor(request))
    flash(request, "NCR closed" if ok else "Record a disposition before closing the NCR", "success" if ok else "error")
    return RedirectResponse(f"/quality/ncr/{ncr_id}", status_code=303)


@router.post("/ncr/{ncr_id}/raise-capa", name="quality_ncr_raise_capa")
async def ncr_raise_capa(ncr_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    n = quality_store.get_ncr(ncr_id, cid)
    if not n:
        flash(request, "NCR not found", "error")
        return RedirectResponse("/quality/ncr", status_code=303)
    capa = quality_store.create_capa(cid, {
        "problem_description": f"NCR {n['ncr_no']} — {n['product_or_material']}: {n['description'][:500]}",
        "possible_cause": n.get("measurement") or "", "source": "ncr", "source_id": ncr_id,
        "target_date": (date.today() + timedelta(days=30)).isoformat(),
    }, _actor(request))
    if capa:
        flash(request, f"CAPA {capa['capa_no']} raised from NCR", "success")
        return RedirectResponse(f"/quality/capa/{capa['id']}", status_code=303)
    flash(request, "Failed to raise CAPA", "error")
    return RedirectResponse(f"/quality/ncr/{ncr_id}", status_code=303)


# ── internal audits ───────────────────────────────────────────────

@router.get("/audits", name="quality_audit_list")
async def audit_list(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    qp = request.query_params
    rows = quality_store.list_audits(cid, qp.get("status") or None, to_date(qp.get("from")), to_date(qp.get("to")))
    if qp.get("export") == "xlsx":
        return _xlsx("Internal audits", ["Audit no.", "Standard", "Scope", "Auditee dept.", "Auditor", "Planned", "Actual",
                                         "Status", "Checklist items", "Non-conformities"],
                     [[r["audit_no"], r["standard"], r["scope"], r["auditee_department"], r["auditor"], r.get("planned_date"),
                       r.get("actual_date"), r["status"], r["item_count"], r["nc_count"]] for r in rows],
                     f"audits_{date.today().isoformat()}.xlsx")
    return templates.TemplateResponse("quality/audit_list.html", _ctx(
        request, "audits", audits=rows, status_filter=qp.get("status") or "", date_from=qp.get("from") or "", date_to=qp.get("to") or ""))


@router.get("/audits/new", name="quality_audit_new_get")
async def audit_new_get(request: Request, user=Depends(login_required)):
    return templates.TemplateResponse("quality/audit_form.html", _ctx(
        request, "audits", audit={"standard": "ISO 9001:2015", "auditor": _actor(request)}, is_edit=False))


@router.post("/audits/new", name="quality_audit_new_post")
async def audit_new_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    a = quality_store.create_audit(cid, data, _actor(request))
    if a:
        flash(request, f"Audit {a['audit_no']} planned — add checklist items", "success")
        return RedirectResponse(f"/quality/audits/{a['id']}", status_code=303)
    flash(request, "Failed to create audit", "error")
    return RedirectResponse("/quality/audits/new", status_code=303)


@router.get("/audits/{audit_id}", name="quality_audit_detail")
async def audit_detail(audit_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    a = quality_store.get_audit(audit_id, cid)
    if not a:
        flash(request, "Audit not found", "error")
        return RedirectResponse("/quality/audits", status_code=303)
    summary = {r: sum(1 for c in a["checklist"] if c["result"] == r) for r in AUDIT_RESULTS}
    return templates.TemplateResponse("quality/audit_detail.html", _ctx(
        request, "audits", audit=a, summary=summary, documents=_documents(cid, "audit", audit_id)))


@router.post("/audits/{audit_id}/edit", name="quality_audit_edit")
async def audit_edit(audit_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    ok = quality_store.update_audit(audit_id, cid, data)
    flash(request, "Audit updated" if ok else "Failed to update audit (closed audits are read-only)", "success" if ok else "error")
    return RedirectResponse(f"/quality/audits/{audit_id}", status_code=303)


@router.post("/audits/{audit_id}/status", name="quality_audit_status")
async def audit_status(audit_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    data = await _form(request)
    status = data.get("status", "")
    ok = status in AUDIT_STATUSES and quality_store.set_audit_status(audit_id, cid, status)
    flash(request, f"Audit marked {status}" if ok else "Failed to update audit status", "success" if ok else "error")
    return RedirectResponse(f"/quality/audits/{audit_id}", status_code=303)


@router.post("/audits/{audit_id}/checklist", name="quality_audit_item_add")
async def audit_item_add(audit_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    if not quality_store.get_audit(audit_id, cid):
        flash(request, "Audit not found", "error")
        return RedirectResponse("/quality/audits", status_code=303)
    ok = quality_store.add_checklist_item(audit_id, cid, data)
    flash(request, "Checklist item added" if ok else "Checklist item text is required", "success" if ok else "error")
    return RedirectResponse(f"/quality/audits/{audit_id}#checklist", status_code=303)


@router.post("/audits/{audit_id}/checklist/{item_id}/edit", name="quality_audit_item_edit")
async def audit_item_edit(audit_id: str, item_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    ok = quality_store.update_checklist_item(item_id, cid, data)
    flash(request, "Checklist item updated" if ok else "Failed to update item", "success" if ok else "error")
    return RedirectResponse(f"/quality/audits/{audit_id}#checklist", status_code=303)


@router.post("/audits/{audit_id}/checklist/{item_id}/delete", name="quality_audit_item_delete")
async def audit_item_delete(audit_id: str, item_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    ok = quality_store.delete_checklist_item(item_id, cid)
    flash(request, "Checklist item removed" if ok else "Failed to remove item", "success" if ok else "error")
    return RedirectResponse(f"/quality/audits/{audit_id}#checklist", status_code=303)


@router.post("/audits/{audit_id}/checklist/{item_id}/raise-capa", name="quality_audit_item_raise_capa")
async def audit_item_raise_capa(audit_id: str, item_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    a = quality_store.get_audit(audit_id, cid)
    item = quality_store.get_checklist_item(item_id, cid)
    if not a or not item:
        flash(request, "Audit item not found", "error")
        return RedirectResponse("/quality/audits", status_code=303)
    capa = quality_store.create_capa(cid, {
        "problem_description": f"Audit {a['audit_no']} ({a['standard']}) finding [{item['result']}] {item.get('requirement_ref') or ''}: {item['item']}",
        "possible_cause": item.get("evidence") or "", "source": "audit", "source_id": item_id,
        "responsible_person": a.get("auditee_department") or "", "target_date": (date.today() + timedelta(days=30)).isoformat(),
        "action_type": "preventive" if item["result"] == "observation" else "corrective",
    }, _actor(request))
    if capa:
        flash(request, f"CAPA {capa['capa_no']} raised from audit finding", "success")
        return RedirectResponse(f"/quality/capa/{capa['id']}", status_code=303)
    flash(request, "Failed to raise CAPA", "error")
    return RedirectResponse(f"/quality/audits/{audit_id}#checklist", status_code=303)


# ── reports ───────────────────────────────────────────────────────

def _report_filters(request: Request) -> dict:
    qp = request.query_params
    d_from, d_to = _dates(request)
    return dict(date_from=d_from, date_to=d_to, product=(qp.get("product") or "").strip(),
                supplier=(qp.get("supplier") or "").strip(), machine=(qp.get("machine") or "").strip(),
                parameter=(qp.get("parameter") or "").strip(), granularity=qp.get("granularity") if qp.get("granularity") in GRANULARITIES else "monthly",
                order_number=(qp.get("order") or "").strip())


def _report_response(request: Request, template: str, key: str, title: str, columns: list, rows: list,
                     filters: dict, filter_fields: tuple, chart: dict = None, **extra):
    cid = current_company(request)
    if request.query_params.get("export") == "xlsx":
        return _xlsx(title, [c[1] for c in columns], [[_fmt(r.get(c[0])) for c in columns] for r in rows],
                     f"quality_{key}_{date.today().isoformat()}.xlsx",
                     subtitle=f"{filters['date_from']} → {filters['date_to']}")
    return templates.TemplateResponse(template, _ctx(
        request, f"report_{key}", report_key=key, report_title=title, columns=columns, rows=rows, filters=filters,
        filter_fields=filter_fields, options=quality_store.filter_options(cid),
        chart_json=json.dumps(chart, default=_json_default) if chart else "", **extra))


@router.get("/reports", name="quality_reports_index")
async def reports_index(request: Request, user=Depends(login_required)):
    return templates.TemplateResponse("quality/reports_index.html", _ctx(request, "reports"))


@router.get("/reports/parameter-stats", name="quality_report_parameter_stats")
async def report_parameter_stats(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    f = _report_filters(request)
    rows = quality_store.measurement_rows(cid, f["date_from"], f["date_to"], f["product"] or None, f["parameter"] or None, f["machine"] or None)
    stats = group_stats(rows, f["granularity"])
    cols = [("period", "Period"), ("product", "Product"), ("parameter", "Parameter"), ("unit", "Unit"), ("n", "n"),
            ("mean", "Mean"), ("stddev", "Std. deviation"), ("min", "Min"), ("max", "Max")]
    return _report_response(request, "quality/report_table.html", "parameter_stats",
                            "Periodic mean & standard deviation per parameter", cols, stats, f,
                            ("date", "granularity", "product", "parameter", "machine"),
                            description="Mean and sample standard deviation of every measured parameter per product, grouped by the selected period.")


@router.get("/reports/stability", name="quality_report_stability")
async def report_stability(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    f = _report_filters(request)
    rows = quality_store.measurement_rows(cid, f["date_from"], f["date_to"], f["product"] or None, f["parameter"] or None, f["machine"] or None)
    series: dict = {}
    for r in rows:
        series.setdefault((r["product"] or "—", r["parameter"]), []).append(r)
    charts = []
    table = []
    for (product, parameter), pts in sorted(series.items()):
        ts = trend_series(pts)
        st = spc_stats([p["value"] for p in ts])
        charts.append({"label": f"{product} — {parameter}", "labels": [p["date"] for p in ts],
                       "values": [p["value"] for p in ts], "running_mean": [p["running_mean"] for p in ts],
                       "mean": st["mean"], "ucl": st["ucl"], "lcl": st["lcl"]})
        table.append({"product": product, "parameter": parameter, "unit": pts[0].get("unit") or "", "n": st["n"],
                      "first": ts[0]["value"] if ts else None, "last": ts[-1]["value"] if ts else None,
                      "mean": st["mean"], "stddev": st["stddev"],
                      "drift_pct": ((ts[-1]["value"] - ts[0]["value"]) / abs(ts[0]["value"]) * 100.0) if ts and ts[0]["value"] else None,
                      "cv_pct": (st["stddev"] / abs(st["mean"]) * 100.0) if st["mean"] and st["stddev"] is not None else None})
    cols = [("product", "Product"), ("parameter", "Parameter"), ("unit", "Unit"), ("n", "n"), ("first", "First"), ("last", "Last"),
            ("mean", "Mean"), ("stddev", "Std. deviation"), ("drift_pct", "Drift %"), ("cv_pct", "CV %")]
    return _report_response(request, "quality/report_stability.html", "stability", "Product stability analysis", cols, table, f,
                            ("date", "product", "parameter", "machine"), chart={"series": charts[:12]})


@router.get("/reports/supplier", name="quality_report_supplier")
async def report_supplier(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    f = _report_filters(request)
    rows = quality_store.supplier_compliance(cid, f["date_from"], f["date_to"], f["supplier"] or None)
    cols = [("supplier_name", "Supplier"), ("inspections", "Inspections"), ("passed", "Passed"), ("failed", "Failed"),
            ("conditional", "Conditional"), ("pass_rate", "Pass rate %"), ("accepted", "Accepted"), ("rejections", "Rejections"),
            ("returns", "Returned"), ("replacements", "Replacements requested"), ("rating", "Rating"), ("last_inspection", "Last inspection")]
    chart = {"labels": [r["supplier_name"] or "—" for r in rows], "pass_rate": [r["pass_rate"] or 0 for r in rows],
             "rejections": [int(r["rejections"] or 0) for r in rows]}
    return _report_response(request, "quality/report_table.html", "supplier", "Supplier compliance & reputability", cols, rows, f,
                            ("date", "supplier"), chart=chart, chart_kind="supplier",
                            description="Pass rate, rejections, returns and replacement requests per supplier from raw-material inspections. Rating: A ≥ 95 %, B ≥ 80 %, C below.")


@router.get("/reports/yield", name="quality_report_yield")
async def report_yield(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    f = _report_filters(request)
    rows = []
    tot = {"input_m": 0.0, "output_m": 0.0, "loss_m": 0.0, "scrap_kg": 0.0, "under_length_m": 0.0, "rolls": 0}
    for p in quality_store.packing_rows(cid, f["date_from"], f["date_to"], f["product"] or None):
        y = yield_stats(p.get("input_length_m"), p.get("total_length_m"), p.get("scrap_kg"), p.get("under_length_m"))
        rows.append({"inspection_date": p["inspection_date"], "ref_no": p["ref_no"], "order_number": p["order_number"],
                     "product": p["product_type_mm2"], "colour": p["colour"], "rolls": int(to_num(p.get("rolls_count")) or 0), **y})
        for k in ("input_m", "output_m", "loss_m", "scrap_kg", "under_length_m"):
            tot[k] += y[k]
        tot["rolls"] += rows[-1]["rolls"]
    totals = {**tot, "yield_pct": (tot["output_m"] / tot["input_m"] * 100.0) if tot["input_m"] else None,
              "loss_pct": (tot["loss_m"] / tot["input_m"] * 100.0) if tot["input_m"] else None}
    cols = [("inspection_date", "Date"), ("ref_no", "Ref"), ("order_number", "Order"), ("product", "Product"), ("colour", "Colour"),
            ("rolls", "Rolls"), ("input_m", "Input (m)"), ("output_m", "Output (m)"), ("loss_m", "Loss (m)"), ("yield_pct", "Yield %"),
            ("under_length_m", "Under-length (m)"), ("scrap_kg", "Scrap (kg)")]
    by_product: dict = {}
    for r in rows:
        b = by_product.setdefault(r["product"] or "—", {"input": 0.0, "output": 0.0, "scrap": 0.0})
        b["input"] += r["input_m"]; b["output"] += r["output_m"]; b["scrap"] += r["scrap_kg"]
    chart = {"labels": list(by_product), "input": [b["input"] for b in by_product.values()],
             "output": [b["output"] for b in by_product.values()], "scrap": [b["scrap"] for b in by_product.values()]}
    return _report_response(request, "quality/report_table.html", "yield", "Material yield & analysis", cols, rows, f,
                            ("date", "product"), chart=chart, chart_kind="yield", totals=totals,
                            description="Input length vs packed output and scrap from the wire packing summaries.")


def _spc_payload(cid: str, f: dict, lsl, usl) -> dict:
    rows = quality_store.measurement_rows(cid, f["date_from"], f["date_to"], f["product"] or None,
                                          f["parameter"] or None, f["machine"] or None, order_number=f["order_number"] or None)
    values = [float(r["value"]) for r in rows]
    if lsl is None and usl is None and rows:  # borrow the inspection's own spec window when present
        lo = [float(r["spec_min"]) for r in rows if r.get("spec_min") is not None]
        hi = [float(r["spec_max"]) for r in rows if r.get("spec_max") is not None]
        lsl = min(lo) if lo else None
        usl = max(hi) if hi else None
    st = spc_stats(values, lsl, usl)
    return {"labels": [f"{r['date']} {r['ref']}" for r in rows], "values": values,
            "refs": [{"id": r["inspection_id"], "kind": r["kind"], "ref": r["ref"]} for r in rows],
            "unit": rows[0]["unit"] if rows else "", "stats": st,
            "mean": st["mean"], "ucl": st["ucl"], "lcl": st["lcl"], "lsl": st["lsl"], "usl": st["usl"]}


@router.get("/reports/spc", name="quality_report_spc")
async def report_spc(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    f = _report_filters(request)
    qp = request.query_params
    payload = _spc_payload(cid, f, to_num(qp.get("lsl")), to_num(qp.get("usl"))) if f["parameter"] else None
    rows = []
    if payload:
        for i, (lbl, v) in enumerate(zip(payload["labels"], payload["values"])):
            rows.append({"point": i + 1, "label": lbl, "value": v,
                         "flag": "out of control" if i in payload["stats"]["out_of_control"] else
                                 "out of spec" if i in payload["stats"]["out_of_spec"] else ""})
    cols = [("point", "#"), ("label", "Inspection"), ("value", "Value"), ("flag", "Flag")]
    return _report_response(request, "quality/report_spc.html", "spc", "SPC control chart", cols, rows, f,
                            ("date", "parameter", "product", "machine", "order"), chart=payload or {},
                            spc=payload, lsl=qp.get("lsl") or "", usl=qp.get("usl") or "")


@router.get("/reports/spc/data.json", name="quality_report_spc_data")
async def report_spc_data(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    f = _report_filters(request)
    qp = request.query_params
    payload = _spc_payload(cid, f, to_num(qp.get("lsl")), to_num(qp.get("usl")))
    return JSONResponse(json.loads(json.dumps(payload, default=_json_default)))


@router.get("/reports/defects", name="quality_report_defects")
async def report_defects(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    f = _report_filters(request)
    buckets: dict = {}
    for r in quality_store.result_rows(cid, f["date_from"], f["date_to"], product=f["product"] or None):
        key = (period_key(r["date"], f["granularity"]), r["product"] or "—")
        b = buckets.setdefault(key, {"period": key[0], "product": key[1], "inspected": 0, "failed": 0, "conditional": 0, "passed": 0})
        b["inspected"] += 1
        bucket_key = {"fail": "failed", "conditional": "conditional", "pass": "passed"}.get(r["overall_result"])
        if bucket_key:
            b[bucket_key] += 1
    rows = []
    for k in sorted(buckets):
        b = buckets[k]
        b["defective_pct"] = defective_rate(b["inspected"], b["failed"])
        rows.append(b)
    cols = [("period", "Period"), ("product", "Product"), ("inspected", "Volume inspected"), ("passed", "Passed"),
            ("conditional", "Conditional"), ("failed", "Defective"), ("defective_pct", "% defective")]
    per_period: dict = {}
    for b in rows:
        p = per_period.setdefault(b["period"], {"inspected": 0, "failed": 0})
        p["inspected"] += b["inspected"]; p["failed"] += b["failed"]
    chart = {"labels": list(per_period), "volume": [p["inspected"] for p in per_period.values()],
             "pct": [defective_rate(p["inspected"], p["failed"]) or 0 for p in per_period.values()]}
    return _report_response(request, "quality/report_table.html", "defects", "Volume vs. percentage defective", cols, rows, f,
                            ("date", "granularity", "product"), chart=chart, chart_kind="defects",
                            description="Inspections performed per period against the share that failed.")


@router.get("/reports/lot-history", name="quality_report_lot_history")
async def report_lot_history(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    ref = (request.query_params.get("ref") or "").strip()
    events = quality_store.lot_history(cid, ref) if ref else []
    if request.query_params.get("export") == "xlsx" and ref:
        return _xlsx(f"Lot history {ref}", ["Date", "Type", "Ref no.", "Product", "Result", "Status", "Detail"],
                     [[e["date"], e["label"], e["ref_no"], e["product"], e["result"], e["status"], e["detail"]] for e in events],
                     f"lot_history_{ref}_{date.today().isoformat()}.xlsx")
    return templates.TemplateResponse("quality/report_lot_history.html", _ctx(request, "report_lot_history", ref=ref, events=events))


@router.get("/reports/nc-summary", name="quality_report_nc_summary")
async def report_nc_summary(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    f = _report_filters(request)
    s = quality_store.nc_summary(cid, f["date_from"], f["date_to"])
    if request.query_params.get("export") == "xlsx":
        rows = [["NCR by item kind", k, v] for k, v in s["by_item_kind"].items()]
        rows += [["NCR by disposition", k, v] for k, v in s["by_disposition"].items()]
        rows += [["NCR by status", k, v] for k, v in s["by_status"].items()]
        rows += [[f"Inspections failed — {b['label']}", b["failed"], b["inspected"]] for b in s["failed_by_kind"].values()]
        return _xlsx("Non-conforming summary", ["Section", "Key", "Count"], rows, f"nc_summary_{date.today().isoformat()}.xlsx")
    chart = {"labels": list(s["by_item_kind"]), "counts": list(s["by_item_kind"].values()),
             "kind_labels": [b["label"] for b in s["failed_by_kind"].values()],
             "kind_failed": [b["failed"] for b in s["failed_by_kind"].values()],
             "kind_inspected": [b["inspected"] for b in s["failed_by_kind"].values()]}
    return templates.TemplateResponse("quality/report_nc_summary.html", _ctx(
        request, "report_nc_summary", summary=s, filters=f, chart_json=json.dumps(chart, default=_json_default)))


@router.get("/reports/audit-summary", name="quality_report_audit_summary")
async def report_audit_summary(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    f = _report_filters(request)
    d_from = to_date(request.query_params.get("from"))  # audits: default to the whole year ahead / behind
    d_to = to_date(request.query_params.get("to"))
    s = quality_store.audit_summary(cid, d_from, d_to)
    f["date_from"], f["date_to"] = d_from or "", d_to or ""
    if request.query_params.get("export") == "xlsx":
        rows = [[a["audit_no"], a["standard"], a["scope"], a["auditee_department"], a["auditor"], a.get("planned_date"),
                 a.get("actual_date"), a["status"], a["item_count"], a["nc_count"]] for a in s["schedule"] + s["history"]]
        return _xlsx("Audit summary", ["Audit no.", "Standard", "Scope", "Auditee", "Auditor", "Planned", "Actual", "Status",
                                       "Items", "Non-conformities"], rows, f"audit_summary_{date.today().isoformat()}.xlsx")
    return templates.TemplateResponse("quality/report_audit_summary.html", _ctx(
        request, "report_audit_summary", summary=s, filters=f,
        chart_json=json.dumps({"labels": list(s["checklist_results"]), "counts": list(s["checklist_results"].values())})))
