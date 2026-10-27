"""
Pure-logic tests — Security Operations (no database).

Covers secops_logic (roster expansion, conflicts, coverage gaps, hours /
night / overtime, overtime multipliers, payroll feed rows, licence status)
plus the side-effect-free parts of secops_jobs and secops_routes.
"""
import sys
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path

import pytest

_REPO_ROOT = Path(__file__).resolve().parents[2]
_WEB_DIR = _REPO_ROOT / "web"
for p in (str(_REPO_ROOT), str(_WEB_DIR)):
    if p not in sys.path:
        sys.path.insert(0, p)

import secops_logic as L  # noqa: E402

DAY = {"id": "d", "name": "Day", "start_time": "06:00", "end_time": "18:00", "is_night": False, "break_minutes": 0}
NIGHT = {"id": "n", "name": "Night", "start_time": time(18, 0), "end_time": time(6, 0), "is_night": True, "break_minutes": 30}
POST_24 = {"id": "p1", "site_id": "s1", "name": "Gate", "coverage": "24h", "guards_required": 2, "is_active": True}
POST_NIGHT = {"id": "p2", "site_id": "s1", "name": "Patrol", "coverage": "night", "guards_required": 1, "is_active": True}
POST_OFF = {"id": "p3", "site_id": "s1", "name": "Old", "coverage": "day", "guards_required": 1, "is_active": False}


# ── time helpers ──────────────────────────────────────────────────
def test_parse_helpers():
    assert L.parse_time("06:30") == time(6, 30)
    assert L.parse_time("0630") == time(6, 30)
    assert L.parse_time(time(1, 2)) == time(1, 2)
    assert L.parse_time("") is None and L.parse_time("bad") is None
    assert L.parse_date("2026-10-26") == date(2026, 10, 26)
    assert L.parse_date(datetime(2026, 1, 1, 5)) == date(2026, 1, 1)
    assert L.parse_date("nope") is None
    assert L.parse_datetime("2026-10-26T07:05") == datetime(2026, 10, 26, 7, 5)
    assert L.parse_datetime("2026-10-26 07:05:09") == datetime(2026, 10, 26, 7, 5, 9)
    assert L.parse_datetime(date(2026, 1, 1)) == datetime(2026, 1, 1)
    assert L.opt("") is None and L.opt("x") == "x"
    assert L.D("") == 0 and L.D("1.5") == Decimal("1.5") and L.D("abc") == 0


def test_parse_month_and_previous_month():
    assert L.parse_month("2026-02") == (date(2026, 2, 1), date(2026, 2, 28))
    assert L.parse_month("2026-12") == (date(2026, 12, 1), date(2026, 12, 31))
    with pytest.raises(ValueError):
        L.parse_month("2026-13")
    assert L.previous_month(date(2026, 1, 15)) == "2025-12"
    assert L.previous_month(date(2026, 10, 1)) == "2026-09"


def test_week_range_is_monday_to_sunday():
    s, e = L.week_range(date(2026, 10, 28))  # a Wednesday
    assert s == date(2026, 10, 26) and e == date(2026, 11, 1)
    assert s.weekday() == 0


def test_shift_bounds_overnight_and_hours():
    s, e = L.shift_bounds(date(2026, 10, 26), DAY)
    assert (s, e) == (datetime(2026, 10, 26, 6), datetime(2026, 10, 26, 18))
    s, e = L.shift_bounds(date(2026, 10, 26), NIGHT)
    assert (s, e) == (datetime(2026, 10, 26, 18), datetime(2026, 10, 27, 6))
    assert L.shift_hours(DAY) == Decimal("12")
    assert L.shift_hours(NIGHT) == Decimal("11.50")


# ── roster expansion / gaps / conflicts ───────────────────────────
def test_slots_for_post_by_coverage():
    assert [t["id"] for t in L.slots_for_post(POST_24, [DAY, NIGHT])] == ["d", "n"]
    assert [t["id"] for t in L.slots_for_post(POST_NIGHT, [DAY, NIGHT])] == ["n"]
    assert [t["id"] for t in L.slots_for_post({"coverage": "day"}, [DAY, NIGHT])] == ["d"]
    assert [t["id"] for t in L.slots_for_post({"coverage": "custom"}, [DAY, NIGHT])] == ["d", "n"]
    assert L.slots_for_post(POST_24, [dict(DAY, is_active=False)]) == []


def test_expand_roster_skips_inactive_posts_and_counts_required():
    slots = L.expand_roster([POST_24, POST_NIGHT, POST_OFF], [DAY, NIGHT], date(2026, 10, 26), date(2026, 10, 27))
    assert len(slots) == 2 * 2 + 1 * 2  # gate day+night ×2 days, patrol night ×2 days
    assert sum(s["guards_required"] for s in slots) == 2 * 2 * 2 + 1 * 2
    assert all("p3" != s["post_id"] for s in slots)
    night = next(s for s in slots if s["post_id"] == "p2" and s["duty_date"] == date(2026, 10, 27))
    assert night["end_at"] == datetime(2026, 10, 28, 6)


def test_coverage_gaps_and_pct():
    slots = L.expand_roster([POST_24], [DAY, NIGHT], date(2026, 10, 26), date(2026, 10, 26))
    assigned = [{"post_id": "p1", "duty_date": "2026-10-26", "shift_template_id": "d", "guard_id": "g1", "status": "scheduled"},
                {"post_id": "p1", "duty_date": date(2026, 10, 26), "shift_template_id": "d", "guard_id": "g2"},
                {"post_id": "p1", "duty_date": date(2026, 10, 26), "shift_template_id": "n", "guard_id": "g3", "status": "cancelled"}]
    gaps = L.coverage_gaps(slots, assigned)
    assert len(gaps) == 1 and gaps[0]["shift_template_id"] == "n" and gaps[0]["gap"] == 2 and gaps[0]["assigned"] == 0
    assert L.coverage_pct(4, 2) == Decimal("50.00")
    assert L.coverage_pct(0, 0) == Decimal(100)
    assert L.coverage_pct(2, 5) == Decimal("100.00")


def test_find_conflicts_same_guard_overlap_any_site():
    cand = {"guard_id": "g1", "start_at": datetime(2026, 10, 26, 18), "end_at": datetime(2026, 10, 27, 6)}
    existing = [
        {"id": "a1", "guard_id": "g1", "start_at": datetime(2026, 10, 26, 6), "end_at": datetime(2026, 10, 26, 18), "site_name": "A"},   # touches, no overlap
        {"id": "a2", "guard_id": "g1", "start_at": datetime(2026, 10, 27, 0), "end_at": datetime(2026, 10, 27, 8), "site_name": "B"},   # overlaps
        {"id": "a3", "guard_id": "g2", "start_at": datetime(2026, 10, 26, 18), "end_at": datetime(2026, 10, 27, 6)},                    # other guard
        {"id": "a4", "guard_id": "g1", "start_at": datetime(2026, 10, 26, 20), "end_at": datetime(2026, 10, 26, 22), "status": "cancelled"},
        {"id": "a5", "guard_id": "g1", "duty_date": date(2026, 10, 26), "start_time": "20:00", "end_time": "23:00"},                    # from template times
    ]
    ids = [c["id"] for c in L.find_conflicts(cand, existing)]
    assert ids == ["a2", "a5"]
    assert L.find_conflicts(dict(cand, id="a2"), existing) == [existing[4]]


def test_auto_fill_shifts_previous_week_forward():
    prev = [{"post_id": "p1", "guard_id": "g1", "shift_template_id": "d", "duty_date": date(2026, 10, 19), "is_relief": False},
            {"post_id": "p1", "guard_id": "g2", "shift_template_id": "n", "duty_date": date(2026, 10, 25), "status": "cancelled"},
            {"post_id": "p1", "guard_id": "g3", "shift_template_id": "n", "duty_date": "2026-10-25"}]
    out = L.auto_fill_from_previous(prev, date(2026, 10, 26), date(2026, 11, 1), date(2026, 10, 19))
    assert [(o["guard_id"], o["duty_date"]) for o in out] == [("g1", date(2026, 10, 26)), ("g3", date(2026, 11, 1))]


# ── hours ─────────────────────────────────────────────────────────
def test_night_overlap_hours_wraps_midnight():
    assert L.night_overlap_hours(datetime(2026, 10, 26, 18), datetime(2026, 10, 27, 6)) == Decimal("8.00")
    assert L.night_overlap_hours(datetime(2026, 10, 26, 6), datetime(2026, 10, 26, 18)) == Decimal("0.00")
    assert L.night_overlap_hours(datetime(2026, 10, 26, 21), datetime(2026, 10, 26, 23)) == Decimal("1.00")
    assert L.night_overlap_hours(datetime(2026, 10, 26, 5), datetime(2026, 10, 26, 7), time(22), time(6)) == Decimal("1.00")
    # custom non-wrapping window
    assert L.night_overlap_hours(datetime(2026, 10, 26, 0), datetime(2026, 10, 27, 0), time(1), time(4)) == Decimal("3.00")
    assert L.night_overlap_hours(datetime(2026, 10, 26, 7), datetime(2026, 10, 26, 6)) == 0


def test_compute_hours_day_shift_with_overtime():
    r = L.compute_hours(datetime(2026, 10, 26, 6, 0), datetime(2026, 10, 26, 18, 0), DAY, work_hours_per_day=8)
    assert r == {"hours_worked": Decimal("12.00"), "night_hours": Decimal("0.00"), "overtime_hours": Decimal("4.00"), "late_minutes": 0}


def test_compute_hours_night_shift_break_and_late():
    r = L.compute_hours("2026-10-26 18:25", "2026-10-27 06:00", NIGHT, work_hours_per_day=8, grace_minutes=15,
                        night_start="22:00", night_end="06:00", duty_date=date(2026, 10, 26))
    assert r["hours_worked"] == Decimal("11.08")      # 11h35 − 30 min break
    assert r["night_hours"] == Decimal("8.00")
    assert r["overtime_hours"] == Decimal("3.08")
    assert r["late_minutes"] == 10                     # 25 min late − 15 grace


def test_compute_hours_missing_checkout_or_bad_order():
    r = L.compute_hours(datetime(2026, 10, 26, 6, 30), None, DAY, grace_minutes=0)
    assert r["hours_worked"] == 0 and r["overtime_hours"] == 0 and r["late_minutes"] == 30
    r = L.compute_hours(datetime(2026, 10, 26, 18), datetime(2026, 10, 26, 6), DAY)
    assert r["hours_worked"] == 0
    r = L.compute_hours(None, None)
    assert r["late_minutes"] == 0


def test_late_minutes_respects_grace():
    s = datetime(2026, 10, 26, 6)
    assert L.late_minutes_for(datetime(2026, 10, 26, 6, 10), s, 15) == 0
    assert L.late_minutes_for(datetime(2026, 10, 26, 6, 20), s, 15) == 5
    assert L.late_minutes_for(datetime(2026, 10, 26, 5, 50), s, 0) == 0
    assert L.late_minutes_for(None, s) == 0


# ── pay ───────────────────────────────────────────────────────────
def test_overtime_multiplier_weekday_weekend_holiday():
    hol = [{"holiday_date": date(2026, 9, 27)}]
    assert L.overtime_multiplier(date(2026, 10, 26), None, hol) == Decimal("1.5")   # Monday
    assert L.overtime_multiplier(date(2026, 11, 1), None, hol) == Decimal("2.0")    # Sunday (default rest day)
    assert L.overtime_multiplier(date(2026, 10, 31), None, hol) == Decimal("1.5")   # Saturday not a rest day by default
    assert L.overtime_multiplier(date(2026, 10, 31), {"weekend_days": "5,6"}, hol) == Decimal("2.0")
    assert L.overtime_multiplier(date(2026, 9, 27), None, hol) == Decimal("2.5")    # Meskel
    assert L.overtime_multiplier(date(2026, 9, 27), {"overtime_holiday": "3"}, ["2026-09-27"]) == Decimal("3")
    assert L.weekend_day_set("") == {6} and L.weekend_day_set("5, 6, 9") == {5, 6}


def test_hourly_rate():
    assert L.hourly_rate(8000, 26, 8) == Decimal("38.46")
    assert L.hourly_rate(8000, None, None) == Decimal("38.46")
    assert L.hourly_rate(8000, 0, 8) == Decimal("38.46")   # 0 days → default 26
    assert L.hourly_rate(0, 26, 8) == 0


def test_payroll_feed_row_aggregates_and_multiplies():
    guard = {"id": "g1", "employee_id": "EMP1", "name": "Abebe", "basic_salary": 8000, "work_days_per_month": 26, "work_hours_per_day": 8}
    settings = {"night_allowance_per_hour": "10", "weekend_days": "6"}
    att = [
        {"duty_date": date(2026, 10, 26), "status": "present", "hours_worked": "12", "night_hours": "0", "overtime_hours": "4"},   # Mon ×1.5
        {"duty_date": date(2026, 11, 1), "status": "late", "hours_worked": "12", "night_hours": "8", "overtime_hours": "4"},      # Sun ×2.0
        {"duty_date": date(2026, 10, 27), "status": "absent"},
        {"duty_date": date(2026, 10, 27), "status": "present", "hours_worked": "8"},                                              # same day → present wins
        {"duty_date": date(2026, 10, 28), "status": "leave"},
    ]
    row = L.payroll_feed_row(att, guard, settings, holidays=[])
    assert row["days_present"] == 3 and row["days_absent"] == 0
    assert row["hours_worked"] == Decimal("32.00") and row["night_hours"] == Decimal("8.00") and row["overtime_hours"] == Decimal("8.00")
    assert row["hourly_rate"] == Decimal("38.46")
    assert row["overtime_pay"] == Decimal("538.44")   # 4×38.46×1.5 + 4×38.46×2.0
    assert row["night_allowance"] == Decimal("80.00")
    assert row["employee_id"] == "EMP1"


def test_payroll_feed_row_empty():
    row = L.payroll_feed_row([], {"id": "g", "employee_id": "E", "basic_salary": 0})
    assert row["days_present"] == 0 and row["overtime_pay"] == 0 and row["hourly_rate"] == 0


def test_billable_amount():
    assert L.billable_amount("720", "85") == Decimal("61200.00")


# ── documents / licences ──────────────────────────────────────────
def test_licence_status_and_expiring_documents():
    today = date(2026, 10, 26)
    assert L.licence_status(None, today) == "missing"
    assert L.licence_status(date(2026, 10, 25), today) == "expired"
    assert L.licence_status(date(2026, 11, 20), today) == "expiring"
    assert L.licence_status(date(2027, 1, 1), today) == "valid"
    g = {"licence_expiry": date(2026, 11, 1), "firearm_permit_expiry": date(2026, 10, 1), "medical_fitness_expiry": None}
    docs = L.expiring_documents(g, today)
    assert [(d["label"], d["status"], d["days"]) for d in docs] == [("Guard licence", "expiring", 6), ("Firearm permit", "expired", -25)]
    assert L.worst_document_status(g, today) == "expired"
    assert L.worst_document_status({"licence_expiry": date(2027, 1, 1)}, today) == "valid"
    assert L.worst_document_status({}, today) == "missing"


def test_reminder_due_horizons():
    today = date(2026, 10, 26)
    assert L.reminder_due(date(2026, 11, 25), today) == 30
    assert L.reminder_due(date(2026, 11, 2), today) == 7
    assert L.reminder_due(date(2026, 11, 3), today) is None
    assert L.reminder_due(date(2026, 10, 26), today) == 0
    assert L.reminder_due(date(2026, 10, 1), today) == 0
    assert L.reminder_due(None, today) is None


def test_numbering_and_grades():
    assert L.format_incident_no(2026, 7) == "INC-2026-000007"
    assert L.guard_no_for(12) == "G-0012"
    assert L.grade_meets("senior_guard", "guard") and not L.grade_meets("trainee", "guard")
    assert L.grade_meets("guard", "") and L.grade_meets(None, "guard")


# ── jobs & routes (side-effect free parts) ────────────────────────
def test_jobs_licence_reminders_and_digest():
    import secops_jobs as J
    today = date(2026, 10, 26)
    guards = [{"guard_no": "G-1", "name": "Abebe", "status": "active", "licence_expiry": date(2026, 11, 25), "firearm_permit_expiry": date(2026, 10, 20)},
              {"guard_no": "G-2", "name": "Gone", "status": "exited", "licence_expiry": date(2026, 11, 2)},
              {"guard_no": "G-3", "name": "Ok", "status": "active", "licence_expiry": date(2027, 5, 1)}]
    rem = J.licence_reminders(guards, today)
    assert [(r["guard_no"], r["document"], r["days"], r["expired"]) for r in rem] == \
        [("G-1", "Guard licence", 30, False), ("G-1", "Firearm permit", 0, True)]
    cov = {"posts_uncovered": 2, "sites": [{"site_name": "Bole", "uncovered": 2, "required": 4, "assigned": 2, "roster_id": "r"},
                                           {"site_name": "HQ", "uncovered": 0, "required": 2, "assigned": 0, "roster_id": None}]}
    text = J._digest("default", today, 3, rem, cov)
    assert "No-shows marked absent" in text and "Bole — 2 missing" in text and "Sites without a roster" in text and "EXPIRED" in text


def test_jobs_register_is_side_effect_free():
    import secops_jobs as J

    class S:
        def __init__(self):
            self.jobs = []

        def add_job(self, fn, trigger=None, **kw):
            self.jobs.append((fn, kw.get("id")))

    s = S()
    J.register_jobs(s)
    assert {j[1] for j in s.jobs} == {J.DAILY_JOB_ID, J.MONTHLY_JOB_ID}
    assert all(callable(j[0]) for j in s.jobs)


def test_routes_static_paths_precede_param_paths_and_names_unique():
    from secops_routes import router
    routes = [(r.path, r.name, tuple(sorted(r.methods))) for r in router.routes]
    names = [n for _p, n, _m in routes]
    assert len(names) == len(set(names))
    assert all(n.startswith("secops_") for n in names)
    assert all(p.startswith("/secops") for p, _n, _m in routes)
    for prefix, param in (("/secops/sites", "{site_id}"), ("/secops/guards", "{guard_id}"), ("/secops/incidents", "{incident_id}"),
                          ("/secops/rosters", "{roster_id}"), ("/secops/reports", "{key}")):
        idx = [i for i, (p, _n, _m) in enumerate(routes) if p.startswith(prefix)]
        first_param = next(i for i in idx if param in routes[i][0])
        statics = [i for i in idx if param not in routes[i][0] and routes[i][0].count("/") == routes[first_param][0].count("/")]
        assert all(i < first_param for i in statics), prefix


def test_routes_manager_only_endpoints():
    from secops_routes import router, manager_required
    guarded = {"secops_roster_publish", "secops_settings", "secops_settings_post", "secops_payroll_feed_generate",
               "secops_payroll_feed_push", "secops_kit_writeoff", "secops_holiday_new", "secops_holiday_delete"}
    for r in router.routes:
        deps = [d.call for d in r.dependant.dependencies]
        if r.name in guarded:
            assert manager_required in deps, r.name
        elif r.name != "secops_api_checkin":
            assert manager_required not in deps, r.name


def test_xlsx_helper_writes_decimals_and_dates(tmp_path):
    from openpyxl import load_workbook
    from secops_routes import _xlsx
    resp = _xlsx("Hours / overtime: [x]?", ["A", "B", "C"], [[Decimal("2.5"), date(2026, 10, 26), datetime(2026, 10, 26, 7, 5)]], "t.xlsx", "sub")
    wb = load_workbook(resp.path)
    ws = wb.active
    assert not set(ws.title) & set("/\\*?:[]")
    assert ws["A5"].value == 2.5 and ws["B5"].value == "2026-10-26" and ws["C5"].value == "2026-10-26 07:05"


def test_report_definitions_are_consistent():
    from secops_routes import REPORTS, SUM_KEYS
    for key, cfg in REPORTS.items():
        assert cfg["title"] and cfg["columns"] and cfg["icon"].startswith("bi-")
        assert len({k for k, _l in cfg["columns"]}) == len(cfg["columns"]), key
    assert "amount" in SUM_KEYS and "hours_worked" in SUM_KEYS
