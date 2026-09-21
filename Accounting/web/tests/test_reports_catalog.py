"""
Pure unit tests for the Report Builder — no database required.

  * reports_catalog: query compiler emits parameterised SQL over whitelisted
    identifiers only, rejects unknown columns/operators, supports group-by +
    aggregates, date presets (incl. Ethiopian fiscal year) and runtime column drop.
  * reports_jobs.compute_next_run: daily / weekly / monthly schedules.
  * reports_engine: CSV / XLSX / built-in PDF renderers work on in-memory results.
"""
import re
import sys
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pytest

_REPO_ROOT = Path(__file__).resolve().parents[2]
_WEB_DIR = _REPO_ROOT / "web"
for p in (str(_REPO_ROOT), str(_WEB_DIR)):
    if p not in sys.path:
        sys.path.insert(0, p)

import reports_catalog as rc  # noqa: E402
from reports_catalog import CatalogError, compile_query, date_range, validate_definition  # noqa: E402
from reports_jobs import compute_next_run, describe_schedule  # noqa: E402

TODAY = date(2026, 9, 12)   # a Saturday
_IDENT = re.compile(r'"([A-Za-z0-9_]+)"')


def _identifiers(sql: str):
    return set(_IDENT.findall(sql))


# ── catalogue sanity ────────────────────────────────────────────────

def test_catalogue_has_required_sources():
    required = {"vat_income", "vat_expenses", "vat_capital", "bid_records", "cpo_records", "employees",
                "inventory_items", "inventory_movements", "journal_entries", "transactions",
                "proc_purchase_orders", "proc_vendors", "pm_projects", "contracts", "hrm_payroll_runs"}
    assert required <= set(rc.SOURCES)
    for s in rc.SOURCES.values():
        assert s.company_column == "company_id"
        assert s.columns


def test_catalog_json_shape():
    cj = rc.catalog_json()
    assert {"sources", "operators", "aggregates", "date_presets", "chart_types"} <= set(cj)
    src = next(s for s in cj["sources"] if s["key"] == "vat_income")
    assert src["date_column"] == "income_date"
    assert {"name": "gross_amount", "label": "Gross amount", "type": "number"} in src["columns"]


# ── compile: plain select ───────────────────────────────────────────

def test_plain_select_is_parameterised_and_whitelisted():
    cq = compile_query({"source_key": "vat_income", "columns": ["income_date", "customer_name", "gross_amount"],
                        "filters": [{"column": "customer_name", "op": "contains", "value": "Ac'me"},
                                    {"column": "gross_amount", "op": "gte", "value": "100"}],
                        "sort": [{"column": "gross_amount", "dir": "desc"}], "date_preset": "all",
                        "limit": 25}, "acme-co", today=TODAY)
    assert cq.sql.startswith('SELECT "income_date" AS "income_date", "customer_name" AS "customer_name", '
                             '"gross_amount" AS "gross_amount" FROM "vat_income" WHERE "company_id" = %s')
    assert "Ac'me" not in cq.sql                     # values never inlined
    assert cq.params == ["acme-co", "%Ac'me%", 100.0, 25]
    assert 'ORDER BY "gross_amount" DESC' in cq.sql and cq.sql.endswith("LIMIT %s")
    allowed = {c.name for c in rc.SOURCES["vat_income"].columns} | {"vat_income", "company_id"}
    assert _identifiers(cq.sql) <= allowed
    assert [c.key for c in cq.columns] == ["income_date", "customer_name", "gross_amount"]


def test_company_scope_is_always_first_predicate():
    cq = compile_query({"source_key": "employees"}, "co-1", today=TODAY)
    assert 'WHERE "company_id" = %s' in cq.sql
    assert cq.params[0] == "co-1"


def test_all_columns_when_none_selected():
    cq = compile_query({"source_key": "contracts"}, "d", today=TODAY)
    assert [c.key for c in cq.columns] == [c.name for c in rc.SOURCES["contracts"].columns]


def test_limit_is_capped_and_floored():
    assert compile_query({"source_key": "employees", "limit": 10 ** 9}, "d").params[-1] == rc.MAX_LIMIT
    assert compile_query({"source_key": "employees", "limit": -5}, "d").params[-1] == 1
    assert compile_query({"source_key": "employees"}, "d", limit=100).params[-1] == 100
    assert compile_query({"source_key": "employees"}, "d").params[-1] == rc.DEFAULT_LIMIT


# ── compile: rejections ─────────────────────────────────────────────

@pytest.mark.parametrize("bad", [
    {"source_key": "users"},                                                             # not in catalogue
    {"source_key": "vat_income", "columns": ["password"]},                              # unknown column
    {"source_key": "vat_income", "columns": ['gross_amount"; DROP TABLE x; --']},      # injection attempt
    {"source_key": "vat_income", "filters": [{"column": "gross_amount", "op": "like", "value": 1}]},
    {"source_key": "vat_income", "filters": [{"column": "nope", "op": "eq", "value": 1}]},
    {"source_key": "vat_income", "filters": [{"column": "gross_amount", "op": "gt", "value": "abc"}]},
    {"source_key": "vat_income", "filters": [{"column": "gross_amount", "op": "between", "value": [1]}]},
    {"source_key": "vat_income", "group_by": ["category"], "aggregates": [{"column": "customer_name", "fn": "sum"}]},
    {"source_key": "vat_income", "group_by": ["category"], "aggregates": [{"column": "gross_amount", "fn": "median"}]},
    {"source_key": "vat_income", "aggregates": [{"column": "gross_amount", "fn": "sum"}]},   # agg w/o group
    {"source_key": "vat_income", "group_by": ["customer_name:month"]},                  # granularity on text
    {"source_key": "vat_income", "group_by": ["category"], "sort": [{"column": "customer_name", "dir": "asc"}]},
    {"source_key": "vat_income", "sort": [{"column": "gross_amount", "dir": "sideways"}]},
    {"source_key": "vat_income", "date_column": "customer_name"},
    {"source_key": "vat_income", "date_preset": "next_year"},
    {"source_key": "vat_income", "chart": {"type": "scatter", "x": "category", "y": "gross_amount"}},
    {"source_key": "vat_income", "chart": {"type": "bar", "x": "category", "y": "not_a_col"}},
    {"source_key": "vat_income", "limit": "lots"},
])
def test_rejects_non_whitelisted(bad):
    with pytest.raises(CatalogError):
        validate_definition(bad)


# ── compile: group by + aggregates ──────────────────────────────────

def test_group_by_month_with_aggregates():
    cq = compile_query({"source_key": "vat_income", "group_by": ["income_date:month", "category"],
                        "aggregates": [{"column": "gross_amount", "fn": "sum"}, {"column": "*", "fn": "count"},
                                       {"column": "vat_amount", "fn": "avg"}],
                        "sort": [{"column": "sum__gross_amount", "dir": "desc"}],
                        "date_preset": "this_year"}, "d", today=TODAY)
    assert [c.key for c in cq.columns] == ["income_date__month", "category", "sum__gross_amount",
                                           "count__all", "avg__vat_amount"]
    assert "to_char(date_trunc('month', \"income_date\"), 'YYYY-MM') AS \"income_date__month\"" in cq.sql
    assert 'SUM("gross_amount") AS "sum__gross_amount"' in cq.sql
    assert 'COUNT(*) AS "count__all"' in cq.sql
    assert "GROUP BY to_char(date_trunc('month', \"income_date\"), 'YYYY-MM'), \"category\"" in cq.sql
    assert 'ORDER BY "sum__gross_amount" DESC' in cq.sql
    assert cq.columns[0].type == "text" and cq.columns[2].type == "number"


def test_group_by_without_aggregates_counts_rows():
    cq = compile_query({"source_key": "bid_records", "group_by": ["status"]}, "d")
    assert 'COUNT(*) AS "count__all"' in cq.sql and "ORDER BY 1" in cq.sql


def test_text_date_columns_are_parsed_defensively():
    cq = compile_query({"source_key": "transactions", "group_by": ["date:quarter"],
                        "aggregates": [{"column": "debit", "fn": "sum"}], "date_preset": "this_month"}, "d", today=TODAY)
    assert "substr(\"date\", 1, 10)::date" in cq.sql
    assert "'^[0-9]{4}-[0-9]{2}-[0-9]{2}'" in cq.sql
    assert cq.params[1:3] == ["2026-09-01", "2026-09-30"]


def test_computed_catalogue_column():
    cq = compile_query({"source_key": "inventory_items", "columns": ["name", "stock_value"]}, "d")
    assert '("current_stock" * "cost_price") AS "stock_value"' in cq.sql


# ── compile: filters ────────────────────────────────────────────────

def test_filter_operators():
    cq = compile_query({"source_key": "vat_expenses", "columns": ["description"], "filters": [
        {"column": "category", "op": "in", "value": "Rent, Utilities"},
        {"column": "gross_amount", "op": "between", "value": ["10", "20"]},
        {"column": "supplier_tin", "op": "is_null", "value": True},
        {"column": "receipt_number", "op": "is_null", "value": False},
        {"column": "is_active", "op": "eq", "value": "true"},
        {"column": "expense_date", "op": "gte", "value": "2026-07-08"},
        {"column": "vat_type", "op": "ne", "value": "EXEMPT"},
    ]}, "d")
    assert '"category" IN (%s, %s)' in cq.sql
    assert '"gross_amount" BETWEEN %s AND %s' in cq.sql
    assert '("supplier_tin" IS NULL)' in cq.sql and '("receipt_number" IS NOT NULL)' in cq.sql
    assert '"expense_date" >= %s' in cq.sql and '"vat_type" <> %s' in cq.sql
    assert cq.params == ["d", "Rent", "Utilities", 10.0, 20.0, True, "2026-07-08", "EXEMPT", rc.DEFAULT_LIMIT]
    assert "Category in (Rent, Utilities)" in cq.filters_summary


def test_invalid_date_filter_rejected():
    with pytest.raises(CatalogError):
        validate_definition({"source_key": "vat_income", "filters": [{"column": "income_date", "op": "eq", "value": "yesterday"}]})


# ── date presets ────────────────────────────────────────────────────

@pytest.mark.parametrize("preset,expected", [
    ("today", (date(2026, 9, 12), date(2026, 9, 12))),
    ("yesterday", (date(2026, 9, 11), date(2026, 9, 11))),
    ("this_week", (date(2026, 9, 7), date(2026, 9, 13))),
    ("last_week", (date(2026, 8, 31), date(2026, 9, 6))),
    ("this_month", (date(2026, 9, 1), date(2026, 9, 30))),
    ("last_month", (date(2026, 8, 1), date(2026, 8, 31))),
    ("this_quarter", (date(2026, 7, 1), date(2026, 9, 30))),
    ("this_year", (date(2026, 1, 1), date(2026, 12, 31))),
    ("last_year", (date(2025, 1, 1), date(2025, 12, 31))),
    ("fiscal_year", (date(2026, 7, 8), date(2027, 7, 7))),
    ("last_fiscal_year", (date(2025, 7, 8), date(2026, 7, 7))),
    ("last_7_days", (date(2026, 9, 6), date(2026, 9, 12))),
    ("last_30_days", (date(2026, 8, 14), date(2026, 9, 12))),
    ("last_45_days", (date(2026, 7, 30), date(2026, 9, 12))),
])
def test_date_presets(preset, expected):
    assert date_range(preset, TODAY) == expected


def test_fiscal_year_boundaries():
    assert date_range("fiscal_year", date(2026, 7, 7)) == (date(2025, 7, 8), date(2026, 7, 7))
    assert date_range("fiscal_year", date(2026, 7, 8)) == (date(2026, 7, 8), date(2027, 7, 7))
    assert date_range("this_quarter", date(2026, 12, 31)) == (date(2026, 10, 1), date(2026, 12, 31))
    assert date_range("all", TODAY) is None and date_range("", TODAY) is None
    with pytest.raises(CatalogError):
        date_range("last_0_days", TODAY)


def test_date_preset_applied_to_datetime_column_uses_half_open_range():
    cq = compile_query({"source_key": "contracts", "columns": ["title"], "date_column": "created_at",
                        "date_preset": "today"}, "d", today=TODAY)
    assert '"created_at" >= %s AND "created_at" < %s' in cq.sql
    assert cq.params[1:3] == ["2026-09-12", "2026-09-13"]
    assert cq.date_range == (TODAY, TODAY)


# ── runtime column availability ─────────────────────────────────────

def test_missing_runtime_columns_are_dropped():
    available = {"company_id", "income_date", "gross_amount", "customer_name"}
    cq = compile_query({"source_key": "vat_income", "columns": ["income_date", "brand", "gross_amount"],
                        "filters": [{"column": "penalty", "op": "eq", "value": "yes"}],
                        "sort": [{"column": "brand", "dir": "asc"}], "date_preset": "today"},
                       "d", today=TODAY, available=available)
    assert [c.key for c in cq.columns] == ["income_date", "gross_amount"]
    assert set(cq.dropped) == {"brand", "penalty"}
    assert "penalty" not in cq.sql and "brand" not in cq.sql


def test_missing_table_raises():
    with pytest.raises(CatalogError):
        compile_query({"source_key": "payroll_data"}, "d", available=set())


def test_prebuilt_templates_validate_and_fit():
    for tpl in rc.PREBUILT_TEMPLATES:
        d = validate_definition(tpl)
        assert d["name"] == tpl["name"]
        physical = {c.name for c in rc.SOURCES[tpl["source_key"]].columns if c.physical} | {"company_id"}
        assert rc.template_fits(tpl, physical), tpl["name"]
        compile_query(d, "d", today=TODAY, available=physical)
    assert not rc.template_fits(rc.PREBUILT_TEMPLATES[0], {"company_id", "description"})


# ── next-run computation ────────────────────────────────────────────

def test_next_run_daily():
    after = datetime(2026, 9, 12, 6, 30)
    assert compute_next_run("daily", 7, 0, after=after) == datetime(2026, 9, 12, 7, 0)
    assert compute_next_run("daily", 7, 0, after=datetime(2026, 9, 12, 7, 0)) == datetime(2026, 9, 13, 7, 0)
    assert compute_next_run("daily", 7, 0, after=datetime(2026, 9, 12, 7, 0, 30)) == datetime(2026, 9, 13, 7, 0)
    assert compute_next_run("daily", 23, 59, after=datetime(2026, 12, 31, 23, 59)) == datetime(2027, 1, 1, 23, 59)


def test_next_run_weekly():
    sat = datetime(2026, 9, 12, 10, 0)             # Saturday
    assert compute_next_run("weekly", 8, 0, weekday=0, after=sat) == datetime(2026, 9, 14, 8, 0)   # Monday
    assert compute_next_run("weekly", 8, 0, weekday=5, after=sat) == datetime(2026, 9, 19, 8, 0)   # next Sat (time passed)
    assert compute_next_run("weekly", 11, 0, weekday=5, after=sat) == datetime(2026, 9, 12, 11, 0)  # later today
    assert compute_next_run("weekly", 8, 0, weekday=6, after=sat) == datetime(2026, 9, 13, 8, 0)   # Sunday


def test_next_run_monthly_clamps_day():
    assert compute_next_run("monthly", 6, 0, day_of_month=31, after=datetime(2026, 2, 1)) == datetime(2026, 2, 28, 6, 0)
    assert compute_next_run("monthly", 6, 0, day_of_month=31, after=datetime(2028, 2, 1)) == datetime(2028, 2, 29, 6, 0)
    assert compute_next_run("monthly", 6, 0, day_of_month=1, after=datetime(2026, 9, 1, 6, 0)) == datetime(2026, 10, 1, 6, 0)
    assert compute_next_run("monthly", 6, 0, day_of_month=15, after=datetime(2026, 12, 20)) == datetime(2027, 1, 15, 6, 0)
    with pytest.raises(ValueError):
        compute_next_run("hourly", 1, 0)


def test_describe_schedule():
    assert describe_schedule({"frequency": "weekly", "weekday": 0, "hour": 7, "minute": 5}) == "Weekly on Monday at 07:05"
    assert describe_schedule({"frequency": "monthly", "day_of_month": 3, "hour": 6, "minute": 0}) == "Monthly on day 3 at 06:00"
    assert describe_schedule({"frequency": "daily", "hour": 18, "minute": 30}) == "Daily at 18:30"


# ── renderers (in-memory) ───────────────────────────────────────────

def _result():
    from reports_engine import ReportResult
    cols = [rc.OutCol("income_date__month", "Month of Income date", "text"),
            rc.OutCol("customer_name", "Customer", "text"),
            rc.OutCol("sum__gross_amount", "Sum of Gross amount", "number", "gross_amount", "sum"),
            rc.OutCol("count__all", "Count", "number", None, "count")]
    rows = [{"income_date__month": "2026-07", "customer_name": "Acme (Pty) Ltd", "sum__gross_amount": Decimal("1234.5"), "count__all": 3},
            {"income_date__month": "2026-08", "customer_name": "Ünïcode & <Co>", "sum__gross_amount": 99.0, "count__all": 1}]
    return ReportResult(columns=cols, rows=rows, totals={"sum__gross_amount": 1333.5, "count__all": 4}, row_count=2)


META = {"report_name": "Income by month", "description": "desc", "company_name": "Test Co", "source_label": "VAT Income",
        "generated_at": "2026-09-12 10:00", "filters_summary": "Income date: fiscal year", "date_preset": "FY",
        "row_count": 2, "truncated": False}


def test_render_csv():
    from reports_engine import render_csv
    text = render_csv(_result()).decode("utf-8-sig")
    lines = text.splitlines()
    assert lines[0] == "Month of Income date,Customer,Sum of Gross amount,Count"
    assert lines[1] == "2026-07,Acme (Pty) Ltd,1234.5,3"
    assert lines[-1].startswith("Total,")


def test_render_xlsx():
    import io
    from openpyxl import load_workbook
    from reports_engine import render_xlsx
    wb = load_workbook(io.BytesIO(render_xlsx(_result(), META)))
    ws = wb.active
    assert ws["A1"].value == "Test Co" and ws["A5"].value == "Month of Income date"
    assert ws["C6"].value == 1234.5 and ws["C6"].number_format == "#,##0.00"
    assert ws["A8"].value == "Total" and ws["C8"].value == 1333.5
    assert ws.auto_filter.ref.startswith("A5:D")


def test_render_pdf_builtin_multi_page():
    from reports_engine import ReportResult, render_pdf
    r = _result()
    big = ReportResult(columns=r.columns, rows=r.rows * 60, totals=r.totals, row_count=120)
    pdf = render_pdf(big, META, force_builtin=True)
    assert pdf.startswith(b"%PDF-1.4") and pdf.rstrip().endswith(b"%%EOF")
    pages = pdf.count(b"/Type /Page ")
    assert pages >= 3                                    # ~45 rows per landscape page
    assert f"(Page 1 of {pages})".encode() in pdf
    assert b"(Test Co)" in pdf and b"(Income by month)" in pdf
    assert b"/Helvetica-Bold" in pdf


def test_render_pdf_default_path_produces_pdf():
    from reports_engine import render_pdf
    pdf = render_pdf(_result(), META)          # reportlab if installed, else builtin
    assert pdf.startswith(b"%PDF")


def test_render_pdf_empty_result():
    from reports_engine import ReportResult, render_pdf
    pdf = render_pdf(ReportResult(columns=[], rows=[]), META, force_builtin=True)
    assert pdf.startswith(b"%PDF") and b"No data" in pdf


def test_result_to_json_and_chart_data():
    from reports_engine import chart_data
    r = _result()
    j = r.to_json(1)
    assert len(j["rows"]) == 1 and j["rows"][0]["sum__gross_amount"] == 1234.5
    assert j["totals"] == {"sum__gross_amount": 1333.5, "count__all": 4}
    c = chart_data(r, {"type": "bar", "x": "income_date__month", "y": "sum__gross_amount"})
    assert c == {"type": "bar", "labels": ["2026-07", "2026-08"], "values": [1234.5, 99.0], "label": "Sum of Gross amount"}
    assert chart_data(r, {"type": "bar", "x": "nope", "y": "sum__gross_amount"}) is None


def test_format_cell():
    from reports_engine import format_cell
    assert format_cell(1234567.891, "number") == "1,234,567.89"
    assert format_cell(42, "number") == "42"
    assert format_cell(Decimal("10"), "number") == "10.00"
    assert format_cell(True, "bool") == "Yes" and format_cell(None, "text") == ""
    assert format_cell(datetime(2026, 1, 2, 3, 4), "datetime") == "2026-01-02 03:04"


def test_parse_recipients():
    from reports_mailer import parse_recipients
    assert parse_recipients(" a@x.com, b@y.org;a@x.com\nnot-an-email ") == ["a@x.com", "b@y.org"]
    assert parse_recipients(None) == [] and parse_recipients(["c@z.io"]) == ["c@z.io"]


def test_mailer_noop_without_api_key(monkeypatch):
    from reports_mailer import send_report_email
    monkeypatch.delenv("RESEND_API_KEY", raising=False)
    assert send_report_email("a@x.com", "s", "<p>x</p>", [("r.pdf", b"%PDF", "application/pdf")]) is False


def test_mailer_posts_base64_attachment(monkeypatch):
    import reports_mailer as rm
    monkeypatch.setenv("RESEND_API_KEY", "re_test")
    monkeypatch.setenv("EMAIL_FROM", "EBMS <noreply@example.com>")
    sent = {}

    def fake_post(api_key, payload):
        sent.update(payload); sent["_key"] = api_key
        return 200, "{}"
    monkeypatch.setattr(rm, "_post", fake_post)
    assert rm.send_report_email("a@x.com,b@y.com", "Subj", "<p>hi</p>", [("r.csv", b"a,b\n1,2\n", "text/csv")]) is True
    assert sent["_key"] == "re_test" and sent["from"] == "EBMS <noreply@example.com>"
    assert sent["to"] == ["a@x.com", "b@y.com"]
    assert sent["attachments"] == [{"filename": "r.csv", "content": "YSxiCjEsMgo="}]
