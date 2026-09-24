"""
Template render tests — Commercial module.

Renders every commercial/ template with route-accurate contexts, both EMPTY and
POPULATED, so undefined variables / wrong field names fail the build instead
of 500ing in production. No database required.
"""
import json
import sys
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import jinja2
import pytest

_REPO_ROOT = Path(__file__).resolve().parents[2]
_WEB_DIR = _REPO_ROOT / "web"
for p in (str(_REPO_ROOT), str(_WEB_DIR)):
    if p not in sys.path:
        sys.path.insert(0, p)

import commercial_logic as L  # noqa: E402
from commercial_data_store import CommercialDataStore  # noqa: E402

env = jinja2.Environment(loader=jinja2.FileSystemLoader(str(_WEB_DIR / "templates")))


def _base_ctx(path="/commercial/x"):
    request = SimpleNamespace(
        url=SimpleNamespace(path=path, query="date_from=2026-01-01"),
        query_params=SimpleNamespace(get=lambda k, d=None: d),
        session={"username": "abebe"}, form=SimpleNamespace(),
    )
    return dict(
        request=request, session={"username": "abebe"}, url_for=lambda *a, **k: "#",
        csrf_token=lambda: "x", get_flashed_messages=lambda **k: [],
        static_url=lambda p: p, static_cdn_url="", app_version="1.0",
        current_company_id="default", current_tenant=None,
        active_page="dashboard", today=date(2026, 9, 24), L=L,
        so_statuses=L.SO_STATUSES, payment_methods=L.PAYMENT_METHODS, segments=L.SEGMENTS,
        lead_statuses=L.LEAD_STATUSES, campaign_statuses=L.CAMPAIGN_STATUSES,
    )


D2 = Decimal


def _settings(**kw):
    s = dict(company_id="default", so_approvers=L.normalize_approvers(None), default_price_list_id=None,
             header_info="Belayab Cable Manufacturing PLC\nAddis Ababa", footer_info="Bank: CBE 1000…",
             iso_doc_no="BCM/SAL/F-01", invoice_series_code="INV", allow_over_credit=False,
             default_vat_rate=D2("0.15"), default_credit_days=30)
    s.update(kw)
    return s


def _customer(**kw):
    c = dict(id="c1", company_id="default", code="CUS-0001", name="Ethio Electric Utility", tin="000123456",
             contact_person="Ato Kebede", phone="+251911000000", email="k@eeu.et", address="Bole", region="Addis Ababa",
             territory_id="t1", sales_rep_id="r1", segment="utility", credit_limit=D2("500000"), credit_terms_days=30,
             portal_user_id=None, notes="", is_active=True, created_by="abebe", created_at=datetime(2026, 1, 5),
             updated_at=datetime(2026, 1, 5), territory_name="Central", sales_rep_name="Sara T.",
             invoiced=D2("120000"), paid=D2("20000"), credited=D2(0), balance=D2("100000"), overdue=D2("15000"),
             available_credit=D2("400000"), over_limit=False)
    c.update(kw)
    return c


def _product(**kw):
    p = dict(id="p1", company_id="default", item_code="NYA-2.5-BLK", product_code="NYA25B", description="NYA 2.5 mm² single core",
             size_mm2="2.5", color="black", unit="m", packaging="coil 100 m", list_price=D2("52"), currency="ETB",
             vat_rate=D2("0.15"), category="Building wire", is_active=True, created_at=datetime(2026, 1, 1), updated_at=datetime(2026, 1, 1))
    p.update(kw)
    return p


def _rep(**kw):
    r = dict(id="r1", company_id="default", name="Sara T.", username="sara", territory_id="t1", commission_pct=D2("2.5"),
             phone="", email="", is_active=True, created_at=datetime(2026, 1, 1), territory_name="Central")
    r.update(kw)
    return r


def _territory(**kw):
    t = dict(id="t1", company_id="default", name="Central", region="Addis Ababa", created_at=datetime(2026, 1, 1), customers=3)
    t.update(kw)
    return t


def _so_line(**kw):
    l = dict(id="l1", so_id="so1", line_no=1, item_code="NYA-2.5-BLK", product_code="NYA25B", size="2.5", description="NYA 2.5 mm²",
             color="black", unit="m", packaging="coil", quantity=D2("1000.000"), unit_price=D2("52.00"), discount=D2("0"),
             line_total=D2("52000.00"), reserved_qty=D2("500"), delivered_qty=D2("400"), invoiced_qty=D2("0"), mo_no="MO-0001",
             mo_request_id="mr1")
    l.update(kw)
    return l


def _approvals(*decisions):
    return [dict(id=f"a{i}", so_id="so1", seq=i + 1, role_label=lab, approver="manager", decided_by="mgr" if d else "",
                 decision=d, decided_at=datetime(2026, 9, 2) if d else None, comment="ok" if d else "")
            for i, (lab, d) in enumerate(zip(("Sales Manager", "Finance Manager", "General Manager"), decisions))]


def _order(**kw):
    lines = kw.pop("lines", [_so_line(), _so_line(id="l2", line_no=2, item_code="NYA-4-RED", mo_request_id=None, mo_no="",
                                                  delivered_qty=D2(0), reserved_qty=D2(0))])
    o = dict(id="so1", company_id="default", so_no="SO-2026-000001", customer_id="c1", customer_name="Ethio Electric Utility",
             proforma_id=None, order_date=date(2026, 9, 1), required_date=date(2026, 9, 20), status="approved", currency="ETB",
             subtotal=D2("104000"), order_discount=D2("0"), discount_total=D2("0"), vat_rate=D2("0.15"), vat_total=D2("15600"),
             grand_total=D2("119600"), credit_sale=True, payment_terms="30 days", sales_rep_id="r1", territory_id="t1",
             prepared_by="abebe", checked_by="", approved_by="mgr", verification_note="", starting_reading=D2("100"),
             ending_reading=D2("1100"), difference_reading=D2("1000"), copy_finance=True, copy_sales=True, copy_first=False,
             header_info="", footer_info="", iso_doc_no="", approval_request_id=None, credit_override_by="", notes="",
             submitted_at=datetime(2026, 9, 1), approved_at=datetime(2026, 9, 2), closed_at=None, created_by="abebe",
             created_at=datetime(2026, 9, 1), updated_at=datetime(2026, 9, 2), sales_rep_name="Sara T.", territory_name="Central",
             lines=lines, approvals=_approvals("approved", "approved", "approved"), mo_requests=[], deliveries=[], invoices=[],
             fulfilment=L.fulfilment(2000, 400, 0), approval_state="approved", next_seq=None)
    o.update(kw)
    return o


def _availability(order):
    return [{"line": l, "on_hand": D2("800") if i == 0 else None, "reserved_total": D2("500"), "available": D2("800") if i == 0 else None,
             "remaining": D2("600"), "can_fulfil": True, "shortfall": D2(0) if i == 0 else None} for i, l in enumerate(order["lines"])]


def _proforma(**kw):
    p = dict(id="pf1", proforma_no="PFI-2026-000001", customer_id="c1", customer_name="Ethio Electric Utility", customer_tin="000123456",
             proforma_date=date(2026, 8, 20), valid_until=date(2026, 9, 19), payment_method="bank", payment_info="CBE 1000…",
             currency="ETB", subtotal=D2("52000"), discount_total=D2(0), vat_total=D2("7800"), grand_total=D2("59800"), status="sent",
             sales_order_id=None, notes="", prepared_by="abebe", checked_by="", approved_by="", header_info="", footer_info="",
             iso_doc_no="", created_by="abebe", created_at=datetime(2026, 8, 20), updated_at=datetime(2026, 8, 20),
             lines=[dict(id="pl1", line_no=1, ref_no="R1", item_code="NYA-2.5-BLK", size="2.5", description="NYA 2.5", color="black",
                         unit="m", quantity=D2("1000"), unit_price=D2("52"), discount=D2(0), line_total=D2("52000"))])
    p.update(kw)
    return p


def _invoice(**kw):
    i = dict(id="i1", invoice_no="INV-2026-000001", so_id="so1", so_no="SO-2026-000001", customer_id="c1", customer_name="Ethio Electric Utility",
             customer_tin="000123456", invoice_date=date(2026, 9, 10), due_date=date(2026, 10, 10), currency="ETB", subtotal=D2("20800"),
             discount_total=D2(0), vat_rate=D2("0.15"), vat_total=D2("3120"), grand_total=D2("23920"), paid_total=D2("10000"),
             credited_total=D2(0), credit_sale=True, status="partially_paid", sales_rep_id="r1", territory_id="t1", vat_income_id="v1",
             erca_invoice_id=None, erca_number="", prepared_by="abebe", checked_by="", approved_by="", header_info="", footer_info="",
             iso_doc_no="", notes="", created_by="abebe", created_at=datetime(2026, 9, 10), updated_at=datetime(2026, 9, 10),
             sales_rep_name="Sara T.", outstanding=D2("13920"), days_overdue=0,
             lines=[dict(id="il1", line_no=1, so_line_id="l1", item_code="NYA-2.5-BLK", description="NYA 2.5", size="2.5", color="black",
                         unit="m", quantity=D2("400"), unit_price=D2("52"), discount=D2(0), line_total=D2("20800"))],
             receipts=[dict(id="rc1", receipt_no="RCT-2026-000001", invoice_id="i1", amount=D2("10000"), method="telebirr", reference="TB123",
                            received_at=datetime(2026, 9, 12), received_by="abebe", source="manual", payment_id=None)],
             returns=[])
    i.update(kw)
    return i


def _delivery(**kw):
    d = dict(id="d1", di_no="DI-2026-000001", so_id="so1", so_no="SO-2026-000001", customer_id="c1", customer_name="Ethio Electric Utility",
             di_date=date(2026, 9, 5), deliver_to="Site A", warehouse="Main store", status="dispatched", notes="", prepared_by="abebe",
             checked_by="", approved_by="", header_info="", footer_info="", iso_doc_no="", created_by="abebe", created_at=datetime(2026, 9, 5),
             dispatches=1, total_qty=D2("400"),
             lines=[dict(id="dl1", di_id="d1", so_line_id="l1", line_no=1, product_code="NYA25B", item_code="NYA-2.5-BLK", description="NYA 2.5",
                         size="2.5", color="black", unit="m", packaging="coil", quantity=D2("400"))])
    d["dispatches"] = [dict(id="x1", dispatch_no="DN-2026-000001", di_id="d1", so_id="so1", vehicle="3-A12345", driver="Dawit", driver_phone="",
                            dispatched_at=datetime(2026, 9, 6), received_by="", delivered_at=None, accepted_at=None, acceptance_note="",
                            status="dispatched", notes="", created_by="abebe", created_at=datetime(2026, 9, 6), di_no="DI-2026-000001",
                            customer_name="Ethio Electric Utility", so_no="SO-2026-000001", required_date=date(2026, 9, 20))]
    d.update(kw)
    return d


def _return(**kw):
    r = dict(id="rn1", rn_no="RN-2026-000001", invoice_id="i1", invoice_no="INV-2026-000001", so_id="so1", customer_id="c1",
             customer_name="Ethio Electric Utility", return_date=date(2026, 9, 15), reason="Damaged in transit", status="draft",
             credit_note_no="", subtotal=D2("520"), vat_total=D2("78"), amount=D2("598"), notes="", prepared_by="abebe", approved_by="",
             created_at=datetime(2026, 9, 15),
             lines=[dict(id="rl1", line_no=1, item_code="NYA-2.5-BLK", description="NYA 2.5", unit="m", quantity=D2("10"), unit_price=D2("52"), line_total=D2("520"))])
    r.update(kw)
    return r


def _mo_request(**kw):
    r = dict(id="mr1", request_no="MOR-2026-000001", so_id="so1", so_line_id="l1", so_no="SO-2026-000001", customer_name="Ethio Electric Utility",
             product_code="NYA25B", item_code="NYA-2.5-BLK", description="NYA 2.5", size="2.5", color="black", unit="m", quantity=D2("1000"),
             packing="coil", cutting_length="100 m", delivery_date=date(2026, 9, 18), comment="", to_factory="Plant 1", status="sent",
             mo_no="MO-0001", mfg_order_id="mfg1", prepared_by="abebe", checked_by="", approved_by="", manager="", header_info="",
             footer_info="", iso_doc_no="", created_at=datetime(2026, 9, 2), mfg_status={"status": "in_production", "produced": 300})
    r.update(kw)
    return r


def _campaign(**kw):
    c = dict(id="cp1", name="Expo 2026", channel="trade fair", segment="contractor", start_date=date(2026, 3, 1), end_date=date(2026, 3, 5),
             budget=D2("200000"), spent=D2("150000"), leads=40, conversions=8, revenue_attributed=D2("900000"), status="completed", owner="abebe",
             notes="Good turnout", created_at=datetime(2026, 2, 1), lead_count=12, won_count=3, event_count=1, roi_pct=D2("500.00"),
             conversion_pct=D2("20.00"), budget_used_pct=D2("75.00"),
             events=[dict(id="e1", name="Booth", event_date=date(2026, 3, 2), location="Millennium Hall", attendees=300, leads=20, cost=D2("50000"), notes="")],
             lead_rows=[dict(id="ld1", name="Hana", company="Build Co", phone="0911", status="qualified", est_value=D2("50000"))])
    c.update(kw)
    return c


def _lead(**kw):
    l = dict(id="ld1", campaign_id="cp1", campaign_name="Expo 2026", name="Hana", company="Build Co", phone="0911", email="h@b.et", segment="contractor",
             source="expo", status="qualified", est_value=D2("50000"), owner="abebe", notes="", converted_customer_id=None,
             created_at=datetime(2026, 3, 3), updated_at=datetime(2026, 3, 3))
    l.update(kw)
    return l


def _dashboard(populated):
    s = CommercialDataStore.empty_dashboard()
    if populated:
        s.update(orders_open=4, orders_pending_approval=1, orders_month=3, revenue_mtd=D2("119600"), revenue_ytd=D2("2000000"),
                 receivables=D2("300000"), overdue=D2("15000"), credit_breaches=1,
                 pipeline={"draft": {"count": 1, "value": D2("1000")}, "approved": {"count": 2, "value": D2("200000")}},
                 top_customers=[{"customer_name": "EEU", "revenue": D2("1000000"), "invoices": 5}],
                 top_products=[{"item_code": "NYA-2.5-BLK", "description": "NYA 2.5", "qty": D2("10000"), "revenue": D2("520000")}],
                 monthly_revenue=[{"period": "2026-08", "total": D2("100000"), "invoices": 2}, {"period": "2026-09", "total": D2("119600"), "invoices": 1}],
                 on_time_pct=D2("80.00"), recent_orders=[_order()], pending_my_approval=[dict(_approvals("")[0], so_no="SO-2026-000002")],
                 proformas_open=2, leads_open=5, backlog_value=D2("200000"))
    chart = {"monthly_labels": [m["period"] for m in s["monthly_revenue"]], "monthly_totals": [float(m["total"]) for m in s["monthly_revenue"]],
             "pipeline_labels": list(s["pipeline"]), "pipeline_counts": [v["count"] for v in s["pipeline"].values()]}
    return dict(stats=s, chart_json=json.dumps(chart), open_summary={"count": 4, "value": D2("200000"), "by_status": {}, "overdue_deliveries": 1 if populated else 0})


def _marketing(populated):
    m = CommercialDataStore.empty_marketing_dashboard()
    if populated:
        m.update(campaigns=[_campaign()], budget=D2("200000"), spent=D2("150000"), revenue=D2("900000"),
                 leads_by_status={"new": 3, "qualified": 2, "won": 1}, leads_total=6, conversion_pct=D2("16.67"), events=1, attendees=300,
                 by_segment=[{"segment": "contractor", "customers": 5}], acquisition=[{"period": "2026-08", "new": 1, "returning": 2, "active": 3}],
                 competitors=2, revenue_by_segment=[{"segment": "contractor", "revenue": D2("500000"), "customers": 3}], cost_per_lead=D2("25000"))
    chart = {"seg_labels": [], "seg_values": [], "acq_labels": [], "acq_new": [], "acq_ret": []}
    return dict(m=m, chart_json=json.dumps(chart), leads=[_lead()] if populated else [])


def _report(populated, key="invoices", extra=False):
    cols = ["Invoice", "Date", "Customer", "Total"]
    rows = [["INV-1", date(2026, 9, 1), "EEU", D2("1000.5")], ["INV-2", None, "X", D2("2")]] if populated else []
    r = {"key": key, "title": "Sales invoice report", "columns": cols, "rows": rows, "totals": ["TOTAL", "", "", D2("1002.5")] if populated else None,
         "date_cols": [1], "period": "2026-10"}
    if extra:
        r["extra"] = {"title": "Extra", "columns": ["A", "B"], "rows": [["x", D2("1")]] if populated else []}
    return dict(report=r, filters={"date_from": "", "date_to": "", "customer_id": "", "item_code": "", "sales_rep_id": "", "territory_id": "",
                                   "status": "", "segment": "", "period": "", "q": ""}, periods=["2026-10"] if populated else [],
                customers=[_customer()], products=[_product()], reps=[_rep()], territories=[_territory()])


def _doc(kind):
    import commercial_routes as R
    obj = {"proforma": _proforma(), "sales_order": _order(), "mo_request": _mo_request(), "delivery": _delivery(),
           "dispatch": dict(_delivery()["dispatches"][0], lines=_delivery()["lines"]), "invoice": _invoice(), "credit_note": _return()}[kind]
    return dict(doc=R._doc_dict.__wrapped__(kind, obj, "default") if hasattr(R._doc_dict, "__wrapped__") else _doc_dict_no_db(R, kind, obj), company_id="default")


def _doc_dict_no_db(R, kind, obj):
    """_doc_dict reads settings through the store; stub that call."""
    import commercial_routes as routes
    store = routes.store
    orig = store.get_settings
    store.get_settings = lambda cid: _settings()
    try:
        return R._doc_dict(kind, obj, "default")
    finally:
        store.get_settings = orig


LOOK = dict(customers=[_customer()], products=[_product()], reps=[_rep()], territories=[_territory()])
EMPTY_LOOK = dict(customers=[], products=[], reps=[], territories=[])
FILTERS = {"date_from": "", "date_to": "", "customer_id": "", "item_code": "", "sales_rep_id": "", "territory_id": "", "status": "", "segment": "", "period": "", "q": ""}

CASES = [
    ("commercial/dashboard.html", lambda: _dashboard(True)),
    ("commercial/dashboard.html", lambda: _dashboard(False)),
    ("commercial/settings.html", lambda: dict(settings=_settings(), price_lists=[dict(id="pl1", name="2026 list", valid_from=None, valid_to=None, currency="ETB", items=3)])),
    ("commercial/settings.html", lambda: dict(settings=_settings(header_info="", iso_doc_no=""), price_lists=[])),
    ("commercial/customers.html", lambda: dict(LOOK, filters=FILTERS, customers=[_customer(), _customer(id="c2", over_limit=True, is_active=False)])),
    ("commercial/customers.html", lambda: dict(EMPTY_LOOK, filters=FILTERS, customers=[])),
    ("commercial/customer_form.html", lambda: dict(customer=_customer(), is_edit=True, portal_users=[{"id": "u1", "name": "Portal user"}], **LOOK)),
    ("commercial/customer_form.html", lambda: dict(customer={"code": "CUS-0002", "credit_terms_days": 30, "is_active": True}, is_edit=False, portal_users=[], **EMPTY_LOOK)),
    ("commercial/customer_detail.html", lambda: dict(customer=_customer(), history={"orders": [_order()], "invoices": [_invoice()], "proformas": [_proforma()]},
                                                     credit=L.credit_check(500000, 100000, 0, allow_over=True), events=[dict(event="created", note="x", actor="a", created_at=datetime(2026, 1, 1))])),
    ("commercial/customer_detail.html", lambda: dict(customer=_customer(balance=D2("600000"), available_credit=D2("-100000")), history={"orders": [], "invoices": [], "proformas": []},
                                                     credit=L.credit_check(500000, 600000, 0, allow_over=True), events=[])),
    ("commercial/products.html", lambda: dict(products=[_product(), _product(id="p2", item_code="X", is_active=False)], q="")),
    ("commercial/products.html", lambda: dict(products=[], q="nya")),
    ("commercial/product_form.html", lambda: dict(product=_product(), is_edit=True)),
    ("commercial/product_form.html", lambda: dict(product={"unit": "m", "currency": "ETB", "vat_rate": "0.15", "is_active": True}, is_edit=False)),
    ("commercial/pricing.html", lambda: dict(price_lists=[dict(id="pl1", name="2026 list", valid_from=date(2026, 1, 1), valid_to=None, currency="ETB", items=1)],
                                             selected=dict(id="pl1", name="2026 list", items=[dict(id="i1", item_code="NYA-2.5-BLK", description="NYA", unit="m", min_qty=D2("0"), unit_price=D2("52"))]),
                                             discounts=[dict(id="d1", name="Gov 5%", kind="percent", value=D2("5"), applies_to="segment", target="government", min_order_value=D2(0), valid_from=None, valid_to=None, is_active=True)],
                                             products=[_product()], settings=_settings(default_price_list_id="pl1"))),
    ("commercial/pricing.html", lambda: dict(price_lists=[], selected=None, discounts=[], products=[], settings=_settings())),
    ("commercial/territories.html", lambda: dict(territories=[_territory()], reps=[_rep(), _rep(id="r2", is_active=False, username=None, territory_name=None)])),
    ("commercial/territories.html", lambda: dict(territories=[], reps=[])),
    ("commercial/proformas.html", lambda: dict(filters=FILTERS, proformas=[_proforma(), _proforma(id="pf2", status="draft", valid_until=date(2026, 1, 1))], statuses=L.PROFORMA_STATUSES)),
    ("commercial/proformas.html", lambda: dict(filters=FILTERS, proformas=[], statuses=L.PROFORMA_STATUSES)),
    ("commercial/proforma_form.html", lambda: dict(proforma={"proforma_date": "2026-09-24", "prepared_by": "abebe"}, lines=[], settings=_settings(), **LOOK)),
    ("commercial/proforma_form.html", lambda: dict(proforma={}, lines=[], settings=_settings(iso_doc_no=""), **EMPTY_LOOK)),
    ("commercial/proforma_detail.html", lambda: dict(proforma=_proforma(), events=[], statuses=L.PROFORMA_STATUSES)),
    ("commercial/proforma_detail.html", lambda: dict(proforma=_proforma(status="converted", sales_order_id="so1", lines=[]), events=[], statuses=L.PROFORMA_STATUSES)),
    ("commercial/sales_orders.html", lambda: dict(filters=dict(FILTERS, status="open"), orders=[_order(), _order(id="so2", status="in_production", required_date=date(2026, 1, 1))], **LOOK)),
    ("commercial/sales_orders.html", lambda: dict(filters=FILTERS, orders=[], **EMPTY_LOOK)),
    ("commercial/sales_order_form.html", lambda: dict(order=_order(status="draft"), lines=_order()["lines"], is_edit=True, settings=_settings(), **LOOK)),
    ("commercial/sales_order_form.html", lambda: dict(order={"order_date": "2026-09-24", "copy_finance": True, "copy_sales": True, "copy_first": True}, lines=[], is_edit=False, settings=_settings(), **EMPTY_LOOK)),
    ("commercial/sales_order_detail.html", lambda: dict(order=_order(), credit=L.credit_check(500000, 100000, 0), availability=_availability(_order()), can_step=None,
                                                        events=[], mfg_status={"status": "in_production"}, is_manager=True, settings=_settings())),
    ("commercial/sales_order_detail.html", lambda: dict(order=_order(status="pending_approval", approvals=_approvals("approved", "", ""), next_seq=2, approval_state="pending_approval"),
                                                        credit=L.credit_check(500000, 100000, 119600), availability=_availability(_order()), can_step=_approvals("approved", "", "")[1],
                                                        events=[], mfg_status=None, is_manager=True, settings=_settings())),
    ("commercial/sales_order_detail.html", lambda: dict(order=_order(status="draft", approvals=[], lines=[], approval_state=None, next_seq=None, fulfilment=L.fulfilment(0, 0)),
                                                        credit=L.credit_check(500000, 480000, 119600), availability=[], can_step=None, events=[], mfg_status=None, is_manager=False, settings=_settings())),
    ("commercial/approvals.html", lambda: dict(pending=[dict(_approvals("")[0], so_no="SO-2026-000002", customer_name="X", grand_total=D2("1"), order_date=date(2026, 9, 1), created_by="z")],
                                               all_pending=[_order(status="pending_approval")], approvers=L.normalize_approvers(None))),
    ("commercial/approvals.html", lambda: dict(pending=[], all_pending=[], approvers=L.normalize_approvers(None))),
    ("commercial/mo_requests.html", lambda: dict(requests=[_mo_request()], status_filter="")),
    ("commercial/mo_requests.html", lambda: dict(requests=[], status_filter="sent")),
    ("commercial/mo_request_detail.html", lambda: dict(req=_mo_request())),
    ("commercial/mo_request_detail.html", lambda: dict(req=_mo_request(so_id=None, mfg_order_id=None, mfg_status=None, delivery_date=None))),
    ("commercial/deliveries.html", lambda: dict(deliveries=[_delivery()], dispatches=_delivery()["dispatches"], status_filter="", dispatch_statuses=L.DISPATCH_STATUSES)),
    ("commercial/deliveries.html", lambda: dict(deliveries=[], dispatches=[], status_filter="", dispatch_statuses=L.DISPATCH_STATUSES)),
    ("commercial/delivery_form.html", lambda: dict(order=_order(), availability=_availability(_order()), actor="abebe")),
    ("commercial/delivery_form.html", lambda: dict(order=_order(lines=[]), availability=[], actor="")),
    ("commercial/delivery_detail.html", lambda: dict(delivery=_delivery(), events=[], dispatch_statuses=L.DISPATCH_STATUSES)),
    ("commercial/delivery_detail.html", lambda: dict(delivery=_delivery(dispatches=[], lines=[], status="open"), events=[], dispatch_statuses=L.DISPATCH_STATUSES)),
    ("commercial/invoices.html", lambda: dict(invoices=[_invoice(), _invoice(id="i2", status="overdue", erca_number="ERCA-1")], filters=FILTERS,
                                              totals={"total": D2("47840"), "paid": D2("20000"), "outstanding": D2("27840")}, credit_only=True, overdue_only=False, statuses=L.INVOICE_STATUSES, **LOOK)),
    ("commercial/invoices.html", lambda: dict(invoices=[], filters=FILTERS, totals={"total": D2(0), "paid": D2(0), "outstanding": D2(0)}, credit_only=False, overdue_only=False, statuses=L.INVOICE_STATUSES, **EMPTY_LOOK)),
    ("commercial/invoice_form.html", lambda: dict(order=_order(), customer=_customer(), invoice={"invoice_date": "2026-09-24", "credit_sale": True}, lines=[], is_manager=True, settings=_settings(), **LOOK)),
    ("commercial/invoice_form.html", lambda: dict(order=None, customer=None, invoice={"invoice_date": "2026-09-24"}, lines=[], is_manager=False, settings=_settings(), **EMPTY_LOOK)),
    ("commercial/invoice_detail.html", lambda: dict(invoice=_invoice(), events=[], is_manager=True)),
    ("commercial/invoice_detail.html", lambda: dict(invoice=_invoice(status="overdue", days_overdue=12, receipts=[], lines=[], paid_total=D2(0), credited_total=D2("598"),
                                                                     returns=[_return(status="credited", credit_note_no="CN-2026-000001")]), events=[], is_manager=False)),
    ("commercial/receipts.html", lambda: dict(receipts=[dict(_invoice()["receipts"][0], invoice_no="INV-2026-000001", customer_name="EEU")], filters=FILTERS, method="", total=D2("10000"))),
    ("commercial/receipts.html", lambda: dict(receipts=[], filters=FILTERS, method="cash", total=D2(0))),
    ("commercial/returns.html", lambda: dict(returns=[_return()], status_filter="", statuses=L.RETURN_STATUSES)),
    ("commercial/returns.html", lambda: dict(returns=[], status_filter="", statuses=L.RETURN_STATUSES)),
    ("commercial/return_form.html", lambda: dict(invoice=_invoice())),
    ("commercial/return_form.html", lambda: dict(invoice=_invoice(lines=[]))),
    ("commercial/return_detail.html", lambda: dict(ret=_return(), is_manager=True)),
    ("commercial/return_detail.html", lambda: dict(ret=_return(status="credited", credit_note_no="CN-2026-000001", approved_by="mgr", lines=[]), is_manager=False)),
    ("commercial/forecasts.html", lambda: dict(forecasts=[dict(id="f1", period="2026-10", item_code="NYA-2.5-BLK", description="NYA", territory_name=None, method="moving_avg",
                                                              detail={"ma3": "180", "ma6": "150", "linear": "220"}, forecast_qty=D2("183.33"), forecast_value=D2("9533"))],
                                               periods=["2026-10"], period="2026-10", total_qty=D2("183.33"), total_value=D2("9533"), next_period="2026-11", **LOOK)),
    ("commercial/forecasts.html", lambda: dict(forecasts=[], periods=[], period="2026-10", total_qty=D2(0), total_value=D2(0), next_period="2026-11", **EMPTY_LOOK)),
    ("commercial/commissions.html", lambda: dict(commissions=[dict(id="cm1", period="2026-09", sales_rep_id="r1", sales_rep_name="Sara T.", invoice_id="i1", invoice_no="INV-1",
                                                                   customer_name="EEU", base_amount=D2("20800"), pct=D2("2.5"), commission_amount=D2("520"), status="accrued")],
                                                 summary=[dict(sales_rep_id="r1", sales_rep_name="Sara T.", period="2026-09", base=D2("20800"), commission=D2("520"), accrued=D2("520"), count=1)],
                                                 filters={"period": "", "sales_rep_id": "", "status": ""}, is_manager=True, **LOOK)),
    ("commercial/commissions.html", lambda: dict(commissions=[], summary=[], filters={"period": "2026-09", "sales_rep_id": "", "status": "paid"}, is_manager=False, **EMPTY_LOOK)),
    ("commercial/marketing_dashboard.html", lambda: _marketing(True)),
    ("commercial/marketing_dashboard.html", lambda: _marketing(False)),
    ("commercial/campaigns.html", lambda: dict(campaigns=[_campaign(), _campaign(id="cp2", roi_pct=None, budget_used_pct=None, budget=D2(0), spent=D2(0))], status_filter="")),
    ("commercial/campaigns.html", lambda: dict(campaigns=[], status_filter="active")),
    ("commercial/campaign_form.html", lambda: dict(campaign=_campaign(), is_edit=True)),
    ("commercial/campaign_form.html", lambda: dict(campaign={"status": "planned"}, is_edit=False)),
    ("commercial/campaign_detail.html", lambda: dict(campaign=_campaign())),
    ("commercial/campaign_detail.html", lambda: dict(campaign=_campaign(events=[], lead_rows=[], notes="", roi_pct=None, budget_used_pct=None))),
    ("commercial/events.html", lambda: dict(events=[dict(_campaign()["events"][0], campaign_name="Expo 2026")], campaigns=[_campaign()])),
    ("commercial/events.html", lambda: dict(events=[], campaigns=[])),
    ("commercial/leads.html", lambda: dict(leads=[_lead(), _lead(id="ld2", converted_customer_id="c1", status="won", segment="")], campaigns=[_campaign()], status_filter="", q="")),
    ("commercial/leads.html", lambda: dict(leads=[], campaigns=[], status_filter="new", q="x")),
    ("commercial/competitors.html", lambda: dict(competitors=[dict(id="k1", name="Rival Cables", product="NYA 2.5", price_observed=D2("49"), observed_on=date(2026, 9, 1), region="Addis", notes="", created_by="a")], products=[_product()])),
    ("commercial/competitors.html", lambda: dict(competitors=[], products=[])),
    ("commercial/reports_index.html", lambda: dict(reports=[(k, v[0]) for k, v in CommercialDataStore.REPORTS.items()])),
    ("commercial/report.html", lambda: _report(True)),
    ("commercial/report.html", lambda: _report(False)),
    ("commercial/report.html", lambda: _report(True, key="forecast", extra=True)),
    ("commercial/report.html", lambda: _report(True, key="customers")),
]


@pytest.mark.parametrize("template,ctx_fn", CASES, ids=[f"{t}:{i}" for i, (t, _) in enumerate(CASES)])
def test_commercial_template_renders(template, ctx_fn):
    html = env.get_template(template).render(**_base_ctx(), **ctx_fn())
    assert len(html) > 1000
    assert "module-sidebar" in html


@pytest.mark.parametrize("kind", ["proforma", "sales_order", "mo_request", "delivery", "dispatch", "invoice", "credit_note"])
def test_doc_print_renders_every_document_kind(kind):
    ctx = _doc(kind)
    html = env.get_template("commercial/doc_print.html").render(**_base_ctx(), **ctx)
    assert ctx["doc"]["number"] in html
    assert "Signature" in html and ctx["doc"]["title"].split()[0] in html
    assert "BCM/SAL/F-01" in html  # ISO doc number from settings


def test_doc_print_renders_with_empty_lines():
    import commercial_routes as R
    obj = _proforma(lines=[], header_info="", footer_info="", iso_doc_no="")
    doc = _doc_dict_no_db(R, "proforma", obj)
    html = env.get_template("commercial/doc_print.html").render(**_base_ctx(), doc=doc, company_id="default")
    assert "No lines" in html


@pytest.mark.parametrize("kind", ["proforma", "sales_order", "invoice", "delivery", "mo_request", "dispatch", "credit_note"])
def test_pdf_builds_for_every_document_kind(kind):
    import commercial_routes as R
    pdf = R._doc_pdf(_doc(kind)["doc"])
    assert pdf.startswith(b"%PDF") and len(pdf) > 800


def test_plain_pdf_fallback_is_valid():
    import commercial_routes as R
    pdf = R._plain_pdf(_doc("invoice")["doc"])
    assert pdf.startswith(b"%PDF") and b"%%EOF" in pdf


def test_xlsx_export_writes_file(tmp_path):
    import os
    import commercial_routes as R
    r = _report(True, extra=True)["report"]
    path = R._xlsx_file(r["title"], r["columns"], r["rows"], r["totals"], [("From", "2026-01-01")], r["extra"])
    try:
        from openpyxl import load_workbook
        ws = load_workbook(path).active
        assert ws["A1"].value == r["title"]
        assert any(c.value == "TOTAL" for row in ws.iter_rows() for c in row)
    finally:
        os.unlink(path)


def test_templates_do_not_use_top_level_jquery_ready():
    for tpl in (_WEB_DIR / "templates" / "commercial").glob("*.html"):
        assert "$(document).ready" not in tpl.read_text(encoding="utf-8"), tpl.name


def test_every_commercial_template_is_covered():
    covered = {t.split("/")[1] for t, _ in CASES} | {"doc_print.html", "_sidebar.html", "_macros.html", "_lines_editor.html"}
    on_disk = {p.name for p in (_WEB_DIR / "templates" / "commercial").glob("*.html")}
    assert on_disk <= covered, f"templates without a render test: {sorted(on_disk - covered)}"


def test_sidebar_links_resolve_to_real_route_names():
    import re
    import commercial_routes as R
    names = {r.name for r in R.router.routes}
    src = (_WEB_DIR / "templates" / "commercial" / "_sidebar.html").read_text(encoding="utf-8")
    for ep in re.findall(r"url_for\('commercial\.([a-z_]+)'", src):
        assert f"commercial_{ep}" in names, ep


def test_all_template_url_for_endpoints_exist():
    import re
    import commercial_routes as R
    names = {r.name for r in R.router.routes}
    for tpl in (_WEB_DIR / "templates" / "commercial").glob("*.html"):
        for ep in re.findall(r"url_for\('commercial\.([a-z_]+)'", tpl.read_text(encoding="utf-8")):
            assert f"commercial_{ep}" in names, f"{tpl.name}: commercial.{ep}"
