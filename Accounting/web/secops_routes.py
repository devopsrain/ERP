"""
Security Operations (manned guarding) routes — /secops

Screens: deployment dashboard, clients, sites & posts, shift templates, guard
profiles, roster builder, attendance (supervisor sheet / self-service /
Telegram API), kit, incidents, payroll feed, reports, settings.
"""
from __future__ import annotations

import json
import logging
import os
import re
import tempfile
from datetime import date, datetime, timedelta
from decimal import Decimal

from fastapi import APIRouter, Depends, Request
from fastapi.responses import FileResponse, JSONResponse, RedirectResponse, Response

from deps import current_company, flash, get_current_user, login_required, require_auth, template_context, validate_upload
from template_engine import templates
import secops_logic as L
from secops_data_store import secops_store

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/secops", tags=["secops"])

_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOC_MODULE = "secops"
UPLOAD_EXTS = {"pdf", "doc", "docx", "xls", "xlsx", "jpg", "jpeg", "png", "txt", "zip"}
manager_required = require_auth("manager")


# ── helpers ───────────────────────────────────────────────────────
def _actor(request: Request) -> str:
    return request.session.get("username", "") or ""


def _ctx(request: Request, active_page: str = "", **extra) -> dict:
    ctx = template_context(request)
    ctx.update(active_page=active_page, grades=L.GRADES, site_types=L.SITE_TYPES, site_statuses=L.SITE_STATUSES,
               coverages=L.COVERAGES, guard_statuses=L.GUARD_STATUSES, roster_statuses=L.ROSTER_STATUSES,
               attendance_statuses=L.ATTENDANCE_STATUSES, kit_kinds=L.KIT_KINDS, kit_statuses=L.KIT_STATUSES,
               incident_categories=L.INCIDENT_CATEGORIES, severities=L.SEVERITIES, incident_statuses=L.INCIDENT_STATUSES,
               training_levels=L.TRAINING_LEVELS, today=date.today())
    ctx.update(extra)
    return ctx


async def _form(request: Request) -> dict:
    form = await request.form()
    return {k: v for k, v in form.items() if not hasattr(v, "filename")}


def _json_default(o):
    if isinstance(o, Decimal):
        return float(o)
    if isinstance(o, (date, datetime)):
        return o.isoformat()
    return str(o)


def _xlsx(title: str, headers: list, rows: list, filename: str, subtitle: str = "") -> FileResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\\/*?:\[\]]", "-", title)[:30].strip() or "Report"
    ws.append([title]); ws["A1"].font = Font(bold=True, size=13)
    if subtitle:
        ws.append([subtitle])
    ws.append([])
    ws.append(headers)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([(float(v) if isinstance(v, Decimal) else v.strftime("%Y-%m-%d %H:%M") if isinstance(v, datetime)
                    else v.isoformat() if isinstance(v, date) else v) for v in r])
    for i, h in enumerate(headers):
        col = chr(65 + i) if i < 26 else "A" + chr(65 + i - 26)
        ws.column_dimensions[col].width = max(12, min(40, len(str(h)) + 4))
    fd, path = tempfile.mkstemp(suffix=".xlsx"); os.close(fd)
    wb.save(path)
    return FileResponse(path, filename=filename, media_type=_XLSX)


def _notify_topic(cid: str, text: str) -> int:
    try:
        from telegram_bot import notify_topic
        return notify_topic(cid, "daily_digest", text)
    except Exception as e:
        logger.debug("secops telegram unavailable: %s", e)
        return 0


def _notify_guard(guard: dict, text: str) -> int:
    """Best effort: a guard's Telegram link is keyed by EBMS username; try e-mail and name."""
    try:
        from telegram_bot import notify_user
        n = 0
        for cand in (guard.get("email"), guard.get("name"), guard.get("employee_id")):
            if cand:
                n += notify_user(str(cand), text)
        return n
    except Exception as e:
        logger.debug("secops telegram unavailable: %s", e)
        return 0


def _documents(cid: str, entity_id: str) -> list:
    try:
        import document_storage as storage
        return storage.documents_for(cid, DOC_MODULE, entity_id)
    except Exception as e:
        logger.warning("secops documents unavailable for %s: %s", entity_id, e)
        return []


async def _store_upload(request: Request, cid: str, entity_id: str, upload, kind: str):
    if not upload or not getattr(upload, "filename", ""):
        return None, "No file selected"
    content = await upload.read()
    ok, err = validate_upload(upload.filename, content, allowed_exts=UPLOAD_EXTS)
    if not ok:
        return None, err
    try:
        import document_storage as storage
        doc = storage.store_document(cid, DOC_MODULE, entity_id, upload.filename, content,
                                     content_type=getattr(upload, "content_type", None), uploaded_by=_actor(request),
                                     tags=[kind], ip=(request.client.host if request.client else ""))
        return doc, None
    except Exception as e:
        logger.error("secops upload failed: %s", e)
        return None, f"Upload failed: {e}"


def _serve_document(request: Request, cid: str, entity_id: str, doc_id: str, back: str):
    try:
        import document_storage as storage
        doc = storage.get_document(doc_id, cid)
        if doc.get("module") != DOC_MODULE or doc.get("entity_id") != entity_id or doc.get("deleted_at"):
            raise storage.DocumentNotFound("Document not found")
        data, filename, ctype = storage.open_document(doc_id, cid, by=_actor(request), ip=(request.client.host if request.client else ""))
        inline = request.query_params.get("download") != "1"
        return Response(content=data, media_type=ctype, headers={
            "Content-Disposition": storage.content_disposition(filename, inline=inline),
            "Content-Length": str(len(data)), "X-Content-Type-Options": "nosniff", "Cache-Control": "private, no-store"})
    except Exception as e:
        flash(request, f"Could not open file: {e}", "error")
        return RedirectResponse(back, status_code=303)


def _date_filters(request: Request, default_days: int = 30) -> dict:
    qp = request.query_params
    today = date.today()
    df = L.parse_date(qp.get("from")) or today - timedelta(days=default_days)
    dt = L.parse_date(qp.get("to")) or today
    return {"date_from": df, "date_to": dt, "site_id": qp.get("site_id") or "", "client_id": qp.get("client_id") or "",
            "guard_id": qp.get("guard_id") or "", "category": qp.get("category") or "", "severity": qp.get("severity") or "",
            "client_owned": qp.get("client_owned") or ""}


# ══════════════════════════════════════════════════════════════════
#  Dashboard
# ══════════════════════════════════════════════════════════════════
@router.get("/", name="secops_dashboard")
async def dashboard(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    secops_store.seed_default_shifts(cid)
    day = L.parse_date(request.query_params.get("day")) or date.today()
    stats = secops_store.dashboard_stats(cid, day)
    chart = {"labels": [s["site_name"] for s in stats["coverage"]["sites"]],
             "required": [s["required"] for s in stats["coverage"]["sites"]],
             "assigned": [s["assigned"] for s in stats["coverage"]["sites"]],
             "present": [s["present"] for s in stats["coverage"]["sites"]]}
    return templates.TemplateResponse("secops/dashboard.html", _ctx(
        request, "dashboard", stats=stats, cov=stats["coverage"], day=day,
        chart_json=json.dumps(chart, default=_json_default)))


@router.get("/coverage", name="secops_coverage")
async def coverage(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    f = _date_filters(request, default_days=0)
    if f["date_to"] == f["date_from"]:
        f["date_to"] = f["date_from"] + timedelta(days=6)
    gaps = secops_store.coverage_gap_report(cid, f["date_from"], f["date_to"], f["site_id"] or None)
    if request.query_params.get("export") == "xlsx":
        return _xlsx("Coverage gaps", ["Date", "Site", "Client", "Post", "Shift", "Required", "Assigned", "Gap"],
                     [[g["duty_date"], g["site_name"], g["client_name"], g["post_name"], g["shift_name"], g["guards_required"], g["assigned"], g["gap"]] for g in gaps],
                     f"coverage_gaps_{f['date_from']}_{f['date_to']}.xlsx", f"{f['date_from']} → {f['date_to']}")
    return templates.TemplateResponse("secops/coverage.html", _ctx(
        request, "coverage", gaps=gaps, filters=f, sites=secops_store.list_sites(cid, status="active"),
        total_gap=sum(g["gap"] for g in gaps)))


# ══════════════════════════════════════════════════════════════════
#  Clients
# ══════════════════════════════════════════════════════════════════
@router.get("/clients", name="secops_clients")
async def clients(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    q = (request.query_params.get("q") or "").strip()
    return templates.TemplateResponse("secops/clients.html", _ctx(
        request, "clients", clients=secops_store.list_clients(cid, q=q or None), search=q))


@router.get("/clients/new", name="secops_client_new_get")
async def client_new_get(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    return templates.TemplateResponse("secops/client_form.html", _ctx(
        request, "clients", client={"is_active": True}, is_edit=False, lookups=secops_store.lookups(cid)))


@router.post("/clients/new", name="secops_client_new_post")
async def client_new_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    c = secops_store.create_client(cid, await _form(request))
    if c:
        flash(request, "Client created", "success")
        return RedirectResponse("/secops/clients", status_code=303)
    flash(request, "Failed to create client — name is required", "error")
    return RedirectResponse("/secops/clients/new", status_code=303)


@router.get("/clients/{client_id}/edit", name="secops_client_edit_get")
async def client_edit_get(client_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    c = secops_store.get_client(cid, client_id)
    if not c:
        flash(request, "Client not found", "error")
        return RedirectResponse("/secops/clients", status_code=303)
    return templates.TemplateResponse("secops/client_form.html", _ctx(
        request, "clients", client=c, is_edit=True, lookups=secops_store.lookups(cid),
        sites=secops_store.list_sites(cid, client_id=client_id)))


@router.post("/clients/{client_id}/edit", name="secops_client_edit_post")
async def client_edit_post(client_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    ok = secops_store.update_client(cid, client_id, await _form(request))
    flash(request, "Client updated" if ok else "Failed to update client", "success" if ok else "error")
    return RedirectResponse("/secops/clients", status_code=303)


# ══════════════════════════════════════════════════════════════════
#  Sites & posts
# ══════════════════════════════════════════════════════════════════
@router.get("/sites", name="secops_sites")
async def sites(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    qp = request.query_params
    return templates.TemplateResponse("secops/sites.html", _ctx(
        request, "sites", sites=secops_store.list_sites(cid, qp.get("status") or None, qp.get("client_id") or None, (qp.get("q") or "").strip() or None),
        clients=secops_store.list_clients(cid), status_filter=qp.get("status") or "", client_filter=qp.get("client_id") or "",
        search=(qp.get("q") or "").strip()))


@router.get("/sites/new", name="secops_site_new_get")
async def site_new_get(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    return templates.TemplateResponse("secops/site_form.html", _ctx(
        request, "sites", site={"status": "active", "site_type": "office", "required_guard_grade": "guard", "client_id": request.query_params.get("client_id") or ""},
        is_edit=False, clients=secops_store.list_clients(cid, active_only=True), lookups=secops_store.lookups(cid)))


@router.post("/sites/new", name="secops_site_new_post")
async def site_new_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    s = secops_store.create_site(cid, await _form(request))
    if s:
        flash(request, "Site created — now add its posts", "success")
        return RedirectResponse(f"/secops/sites/{s['id']}", status_code=303)
    flash(request, "Failed to create site — name is required", "error")
    return RedirectResponse("/secops/sites/new", status_code=303)


@router.get("/sites/{site_id}", name="secops_site_detail")
async def site_detail(site_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    s = secops_store.get_site(cid, site_id)
    if not s:
        flash(request, "Site not found", "error")
        return RedirectResponse("/secops/sites", status_code=303)
    today = date.today()
    return templates.TemplateResponse("secops/site_detail.html", _ctx(
        request, "sites", site=s, posts=secops_store.list_posts(cid, site_id), shifts=secops_store.list_shifts(cid, active_only=True),
        rosters=secops_store.list_rosters(cid, site_id=site_id, limit=12),
        guards=[g for g in secops_store.list_guards(cid, site_id=site_id) if g["status"] != "exited"],
        kit=secops_store.outstanding_kit(cid, site_id=site_id), incidents=secops_store.list_incidents(cid, site_id=site_id, limit=10),
        documents=_documents(cid, site_id), coverage=next((c for c in secops_store.coverage_summary(cid, today)["sites"] if c["site_id"] == site_id), None),
        week_start=L.week_range(today)[0]))


@router.get("/sites/{site_id}/edit", name="secops_site_edit_get")
async def site_edit_get(site_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    s = secops_store.get_site(cid, site_id)
    if not s:
        flash(request, "Site not found", "error")
        return RedirectResponse("/secops/sites", status_code=303)
    return templates.TemplateResponse("secops/site_form.html", _ctx(
        request, "sites", site=s, is_edit=True, clients=secops_store.list_clients(cid), lookups=secops_store.lookups(cid)))


@router.post("/sites/{site_id}/edit", name="secops_site_edit_post")
async def site_edit_post(site_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    ok = secops_store.update_site(cid, site_id, await _form(request))
    flash(request, "Site updated" if ok else "Failed to update site", "success" if ok else "error")
    return RedirectResponse(f"/secops/sites/{site_id}", status_code=303)


@router.post("/sites/{site_id}/posts/new", name="secops_post_new")
async def post_new(site_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    if not secops_store.get_site(cid, site_id):
        flash(request, "Site not found", "error")
        return RedirectResponse("/secops/sites", status_code=303)
    p = secops_store.create_post(cid, site_id, await _form(request))
    flash(request, "Post added" if p else "Failed to add post — name is required", "success" if p else "error")
    return RedirectResponse(f"/secops/sites/{site_id}#posts", status_code=303)


@router.post("/sites/{site_id}/posts/{post_id}/edit", name="secops_post_edit")
async def post_edit(site_id: str, post_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    ok = secops_store.update_post(cid, post_id, await _form(request))
    flash(request, "Post updated" if ok else "Failed to update post", "success" if ok else "error")
    return RedirectResponse(f"/secops/sites/{site_id}#posts", status_code=303)


@router.post("/sites/{site_id}/posts/{post_id}/delete", name="secops_post_delete")
async def post_delete(site_id: str, post_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    ok, note = secops_store.delete_post(cid, post_id)
    flash(request, note or ("Post removed" if ok else "Failed to remove post"), "success" if ok else "error")
    return RedirectResponse(f"/secops/sites/{site_id}#posts", status_code=303)


@router.post("/sites/{site_id}/upload", name="secops_site_upload")
async def site_upload(site_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    if not secops_store.get_site(cid, site_id):
        flash(request, "Site not found", "error")
        return RedirectResponse("/secops/sites", status_code=303)
    form = await request.form()
    doc, err = await _store_upload(request, cid, site_id, form.get("file"), (form.get("kind") or "site_document")[:40])
    flash(request, "File uploaded" if doc else (err or "Upload failed"), "success" if doc else "error")
    return RedirectResponse(f"/secops/sites/{site_id}#documents", status_code=303)


@router.get("/sites/{site_id}/file/{doc_id}", name="secops_site_file")
async def site_file(site_id: str, doc_id: str, request: Request, user=Depends(login_required)):
    return _serve_document(request, current_company(request), site_id, doc_id, f"/secops/sites/{site_id}#documents")


# ══════════════════════════════════════════════════════════════════
#  Shift templates
# ══════════════════════════════════════════════════════════════════
@router.get("/shifts", name="secops_shifts")
async def shifts(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    secops_store.seed_default_shifts(cid)
    return templates.TemplateResponse("secops/shifts.html", _ctx(request, "shifts", shifts=secops_store.list_shifts(cid)))


@router.post("/shifts/new", name="secops_shift_new")
async def shift_new(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    s = secops_store.create_shift(cid, await _form(request))
    flash(request, "Shift template created" if s else "Failed to create shift template", "success" if s else "error")
    return RedirectResponse("/secops/shifts", status_code=303)


@router.post("/shifts/{shift_id}/edit", name="secops_shift_edit")
async def shift_edit(shift_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    ok = secops_store.update_shift(cid, shift_id, await _form(request))
    flash(request, "Shift template updated" if ok else "Failed to update shift template", "success" if ok else "error")
    return RedirectResponse("/secops/shifts", status_code=303)


@router.post("/shifts/{shift_id}/delete", name="secops_shift_delete")
async def shift_delete(shift_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    ok = secops_store.delete_shift(cid, shift_id)
    flash(request, "Shift template removed" if ok else "Failed to remove shift template", "success" if ok else "error")
    return RedirectResponse("/secops/shifts", status_code=303)


# ══════════════════════════════════════════════════════════════════
#  Guards
# ══════════════════════════════════════════════════════════════════
@router.get("/guards", name="secops_guards")
async def guards(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    qp = request.query_params
    rows = secops_store.list_guards(cid, qp.get("status") or None, (qp.get("q") or "").strip() or None,
                                    qp.get("site_id") or None, qp.get("grade") or None)
    if qp.get("docs") == "expiring":
        rows = [g for g in rows if g["expiring"]]
    if qp.get("export") == "xlsx":
        return _xlsx("Guards", ["Guard no", "Name", "Employee ID", "Grade", "Status", "Home site", "Licence no", "Licence expiry",
                                "Firearm permit", "Permit expiry", "Medical expiry", "Training"],
                     [[g["guard_no"], g.get("name"), g["employee_id"], g["grade"], g["status"], g.get("home_site_name"), g["licence_no"],
                       g.get("licence_expiry"), g.get("firearm_permit_no"), g.get("firearm_permit_expiry"), g.get("medical_fitness_expiry"),
                       g["training_level"]] for g in rows], "guards.xlsx")
    counts = {s: sum(1 for g in rows if g["status"] == s) for s in L.GUARD_STATUSES}
    return templates.TemplateResponse("secops/guards.html", _ctx(
        request, "guards", guards=rows, counts=counts, sites=secops_store.list_sites(cid), status_filter=qp.get("status") or "",
        site_filter=qp.get("site_id") or "", grade_filter=qp.get("grade") or "", search=(qp.get("q") or "").strip(),
        docs_filter=qp.get("docs") or ""))


@router.get("/guards/new", name="secops_guard_new_get")
async def guard_new_get(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    return templates.TemplateResponse("secops/guard_form.html", _ctx(
        request, "guards", guard={"grade": "guard", "status": "active", "training_level": "basic"}, is_edit=False,
        employees=secops_store.employees_without_profile(cid), sites=secops_store.list_sites(cid, status="active")))


@router.post("/guards/new", name="secops_guard_new_post")
async def guard_new_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    g, err = secops_store.create_guard(cid, await _form(request))
    if g:
        flash(request, f"Guard profile {g['guard_no']} created", "success")
        return RedirectResponse(f"/secops/guards/{g['id']}", status_code=303)
    flash(request, err or "Failed to create guard profile", "error")
    return RedirectResponse("/secops/guards/new", status_code=303)


@router.get("/guards/{guard_id}", name="secops_guard_detail")
async def guard_detail(guard_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    g = secops_store.get_guard(cid, guard_id)
    if not g:
        flash(request, "Guard not found", "error")
        return RedirectResponse("/secops/guards", status_code=303)
    today = date.today()
    return templates.TemplateResponse("secops/guard_detail.html", _ctx(
        request, "guards", guard=g, kit=secops_store.list_kit_issues(cid, guard_id=guard_id),
        duties=secops_store.list_assignments(cid, guard_id=guard_id, date_from=today, date_to=today + timedelta(days=14)),
        attendance=secops_store.list_attendance(cid, guard_id=guard_id, limit=30),
        incidents=[i for i in secops_store.list_incidents(cid, limit=500) if i.get("guard_id") == guard_id][:10],
        documents=_documents(cid, guard_id)))


@router.get("/guards/{guard_id}/edit", name="secops_guard_edit_get")
async def guard_edit_get(guard_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    g = secops_store.get_guard(cid, guard_id)
    if not g:
        flash(request, "Guard not found", "error")
        return RedirectResponse("/secops/guards", status_code=303)
    return templates.TemplateResponse("secops/guard_form.html", _ctx(
        request, "guards", guard=g, is_edit=True, employees=[], sites=secops_store.list_sites(cid, status="active")))


@router.post("/guards/{guard_id}/edit", name="secops_guard_edit_post")
async def guard_edit_post(guard_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    ok = secops_store.update_guard(cid, guard_id, await _form(request))
    flash(request, "Guard profile updated" if ok else "Failed to update guard profile", "success" if ok else "error")
    return RedirectResponse(f"/secops/guards/{guard_id}", status_code=303)


@router.post("/guards/{guard_id}/status", name="secops_guard_status")
async def guard_status(guard_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    ok, outstanding = secops_store.set_guard_status(cid, guard_id, data.get("status", ""))
    if ok:
        flash(request, f"Guard marked {data.get('status', '').replace('_', ' ')}", "success")
        if outstanding:
            flash(request, f"{len(outstanding)} kit item(s) still issued to this guard — collect before final settlement", "error")
    else:
        flash(request, "Failed to update status", "error")
    return RedirectResponse(f"/secops/guards/{guard_id}", status_code=303)


@router.get("/guards/{guard_id}/kit-card", name="secops_guard_kit_card")
async def guard_kit_card(guard_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    g = secops_store.get_guard(cid, guard_id)
    if not g:
        flash(request, "Guard not found", "error")
        return RedirectResponse("/secops/guards", status_code=303)
    return templates.TemplateResponse("secops/kit_card.html", _ctx(
        request, "kit", guard=g, kit=secops_store.list_kit_issues(cid, guard_id=guard_id), printed_at=datetime.now()))


@router.post("/guards/{guard_id}/upload", name="secops_guard_upload")
async def guard_upload(guard_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    if not secops_store.get_guard(cid, guard_id):
        flash(request, "Guard not found", "error")
        return RedirectResponse("/secops/guards", status_code=303)
    form = await request.form()
    doc, err = await _store_upload(request, cid, guard_id, form.get("file"), (form.get("kind") or "licence")[:40])
    flash(request, "File uploaded" if doc else (err or "Upload failed"), "success" if doc else "error")
    return RedirectResponse(f"/secops/guards/{guard_id}#documents", status_code=303)


@router.get("/guards/{guard_id}/file/{doc_id}", name="secops_guard_file")
async def guard_file(guard_id: str, doc_id: str, request: Request, user=Depends(login_required)):
    return _serve_document(request, current_company(request), guard_id, doc_id, f"/secops/guards/{guard_id}#documents")


# ══════════════════════════════════════════════════════════════════
#  Rosters
# ══════════════════════════════════════════════════════════════════
@router.get("/rosters", name="secops_rosters")
async def rosters(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    qp = request.query_params
    return templates.TemplateResponse("secops/rosters.html", _ctx(
        request, "rosters", rosters=secops_store.list_rosters(cid, qp.get("site_id") or None, qp.get("status") or None),
        sites=secops_store.list_sites(cid, status="active"), site_filter=qp.get("site_id") or "", status_filter=qp.get("status") or "",
        week_start=L.week_range()[0]))


@router.post("/rosters/new", name="secops_roster_new")
async def roster_new(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    start = L.parse_date(data.get("period_start"))
    if start:
        start = L.week_range(start)[0] if data.get("snap_week", "1") == "1" else start
    r = secops_store.create_roster(cid, data.get("site_id"), start, data.get("period_end"), _actor(request))
    if r:
        return RedirectResponse(f"/secops/rosters/{r['id']}", status_code=303)
    flash(request, "Choose a site and a week start", "error")
    return RedirectResponse("/secops/rosters", status_code=303)


@router.get("/rosters/{roster_id}", name="secops_roster_builder")
async def roster_builder(roster_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    r = secops_store.get_roster(cid, roster_id)
    if not r:
        flash(request, "Roster not found", "error")
        return RedirectResponse("/secops/rosters", status_code=303)
    grid = secops_store.roster_grid(cid, r)
    guards = [g for g in secops_store.list_guards(cid, status="active")]
    home = [g for g in guards if g.get("home_site_id") == r["site_id"]]
    others = [g for g in guards if g.get("home_site_id") != r["site_id"]]
    prev_start = r["period_start"] - (r["period_end"] - r["period_start"] + timedelta(days=1))
    return templates.TemplateResponse("secops/roster_builder.html", _ctx(
        request, "rosters", roster=r, grid=grid, guards=home + others, home_guard_ids={g["id"] for g in home},
        prev_start=prev_start, next_start=r["period_end"] + timedelta(days=1)))


@router.post("/rosters/{roster_id}/assign", name="secops_roster_assign")
async def roster_assign(roster_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    row, err = secops_store.add_assignment(cid, roster_id, data)
    if row:
        flash(request, "Guard assigned", "success")
    else:
        flash(request, err or "Failed to assign guard", "error")
    return RedirectResponse(f"/secops/rosters/{roster_id}#d{data.get('duty_date', '')}", status_code=303)


@router.post("/rosters/{roster_id}/assignments/{assignment_id}/delete", name="secops_assignment_delete")
async def assignment_delete(roster_id: str, assignment_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    ok = secops_store.delete_assignment(cid, assignment_id)
    flash(request, "Assignment removed" if ok else "Failed to remove assignment", "success" if ok else "error")
    return RedirectResponse(f"/secops/rosters/{roster_id}", status_code=303)


@router.post("/rosters/{roster_id}/autofill", name="secops_roster_autofill")
async def roster_autofill(roster_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    res = secops_store.auto_fill(cid, roster_id)
    if res.get("error"):
        flash(request, res["error"], "error")
    else:
        flash(request, f"Auto-fill: {res['added']} assignment(s) copied, {res['skipped']} skipped (conflicts or already filled)", "success")
    return RedirectResponse(f"/secops/rosters/{roster_id}", status_code=303)


@router.post("/rosters/{roster_id}/publish", name="secops_roster_publish")
async def roster_publish(roster_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    r = secops_store.get_roster(cid, roster_id)
    if not r:
        flash(request, "Roster not found", "error")
        return RedirectResponse("/secops/rosters", status_code=303)
    if secops_store.set_roster_status(cid, roster_id, "published"):
        assignments = secops_store.list_assignments(cid, roster_id=roster_id)
        by_guard: dict = {}
        for a in assignments:
            by_guard.setdefault(a["guard_id"], []).append(a)
        guards = {g["id"]: g for g in secops_store.list_guards(cid)}
        sent = 0
        for gid, duties in by_guard.items():
            g = guards.get(gid)
            if not g:
                continue
            lines = [f"🛡 <b>Duty roster {r['site_name']}</b> {r['period_start']} → {r['period_end']}"]
            lines += [f" • {d['duty_date']} {d.get('shift_name') or ''} — {d.get('post_name')}" for d in duties[:14]]
            sent += _notify_guard(g, "\n".join(lines))
        _notify_topic(cid, f"🛡 Roster published: <b>{r['site_name']}</b> {r['period_start']} → {r['period_end']} "
                           f"({len(assignments)} duties, {len(by_guard)} guards)")
        flash(request, f"Roster published — {len(by_guard)} guard(s) can now see their duties" + (f", {sent} Telegram notice(s) queued" if sent else ""), "success")
    else:
        flash(request, "Failed to publish roster", "error")
    return RedirectResponse(f"/secops/rosters/{roster_id}", status_code=303)


@router.post("/rosters/{roster_id}/status", name="secops_roster_status")
async def roster_status(roster_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    status = data.get("status", "")
    if status == "published":
        flash(request, "Use the publish button", "error")
    elif secops_store.set_roster_status(cid, roster_id, status):
        flash(request, f"Roster marked {status}", "success")
    else:
        flash(request, "Failed to update roster", "error")
    return RedirectResponse(f"/secops/rosters/{roster_id}", status_code=303)


# ══════════════════════════════════════════════════════════════════
#  Attendance
# ══════════════════════════════════════════════════════════════════
@router.get("/attendance", name="secops_attendance")
async def attendance(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    f = _date_filters(request, default_days=7)
    status = request.query_params.get("status") or ""
    rows = secops_store.list_attendance(cid, f["date_from"], f["date_to"], f["site_id"] or None, f["guard_id"] or None, status or None)
    if request.query_params.get("export") == "xlsx":
        return _xlsx("Attendance", ["Date", "Site", "Guard no", "Guard", "Post", "Shift", "Status", "Check-in", "Check-out", "Hours", "Night", "Overtime", "Late (min)", "Source"],
                     [[r["duty_date"], r.get("site_name"), r.get("guard_no"), r.get("guard_name"), r.get("post_name"), r.get("shift_name"), r["status"],
                       r.get("check_in_at"), r.get("check_out_at"), r["hours_worked"], r["night_hours"], r["overtime_hours"], r.get("late_minutes"), r.get("source")] for r in rows],
                     f"attendance_{f['date_from']}_{f['date_to']}.xlsx")
    totals = {"hours": sum((Decimal(str(r["hours_worked"] or 0)) for r in rows), Decimal(0)),
              "night": sum((Decimal(str(r["night_hours"] or 0)) for r in rows), Decimal(0)),
              "overtime": sum((Decimal(str(r["overtime_hours"] or 0)) for r in rows), Decimal(0)),
              "absent": sum(1 for r in rows if r["status"] == "absent"), "late": sum(1 for r in rows if r["status"] == "late")}
    return templates.TemplateResponse("secops/attendance.html", _ctx(
        request, "attendance", rows=rows, filters=f, status_filter=status, totals=totals,
        sites=secops_store.list_sites(cid), guards=[g for g in secops_store.list_guards(cid) if g["status"] != "exited"]))


@router.get("/attendance/sheet", name="secops_attendance_sheet")
async def attendance_sheet(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    qp = request.query_params
    day = L.parse_date(qp.get("day")) or date.today()
    site_id = qp.get("site_id") or ""
    sites = secops_store.list_sites(cid, status="active")
    rows = secops_store.attendance_sheet(cid, site_id, day) if site_id else []
    site = secops_store.get_site(cid, site_id) if site_id else None
    return templates.TemplateResponse("secops/attendance_sheet.html", _ctx(
        request, "attendance", rows=rows, sites=sites, site=site, site_id=site_id, day=day,
        guards=[g for g in secops_store.list_guards(cid, status="active")], posts=secops_store.list_posts(cid, site_id, active_only=True) if site_id else [],
        prev_day=day - timedelta(days=1), next_day=day + timedelta(days=1)))


@router.post("/attendance/sheet", name="secops_attendance_sheet_post")
async def attendance_sheet_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    site_id = data.get("site_id") or ""
    day = L.parse_date(data.get("day")) or date.today()
    settings = secops_store.get_settings(cid)
    keys = [k[len("status_"):] for k in data if k.startswith("status_")]
    saved = 0
    for key in keys:
        status = data.get(f"status_{key}") or ""
        if not status:
            continue
        row = secops_store.save_attendance(cid, {
            "id": data.get(f"att_{key}") or None, "assignment_id": data.get(f"assignment_{key}") or None,
            "guard_id": data.get(f"guard_{key}"), "site_id": site_id, "post_id": data.get(f"post_{key}") or None,
            "shift_template_id": data.get(f"shift_{key}") or None, "duty_date": day, "status": status,
            "check_in_at": data.get(f"in_{key}") or None, "check_out_at": data.get(f"out_{key}") or None,
            "note": data.get(f"note_{key}") or "", "source": "supervisor"}, _actor(request), settings)
        if row:
            saved += 1
    flash(request, f"Attendance saved for {saved} guard(s)", "success" if saved else "error")
    return RedirectResponse(f"/secops/attendance/sheet?site_id={site_id}&day={day}", status_code=303)


@router.post("/attendance/new", name="secops_attendance_new")
async def attendance_new(request: Request, user=Depends(login_required)):
    """Ad-hoc row (walk-in relief, correction)."""
    cid = current_company(request)
    data = await _form(request)
    data.setdefault("source", "manual")
    row = secops_store.save_attendance(cid, data, _actor(request))
    flash(request, "Attendance recorded" if row else "Failed to record attendance — guard and date are required", "success" if row else "error")
    back = data.get("back") or f"/secops/attendance/sheet?site_id={data.get('site_id', '')}&day={data.get('duty_date', '')}"
    return RedirectResponse(back, status_code=303)


# ── Guard self-service ────────────────────────────────────────────
def _guard_for_request(request: Request):
    cid = current_company(request)
    return secops_store.guard_for_user(cid, request.session.get("username", ""), request.session.get("email", ""),
                                       request.session.get("full_name", ""))


@router.get("/me", name="secops_me")
async def me(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    g = _guard_for_request(request)
    ctx = _ctx(request, "me", guard=g, open_row=None, duties=[], recent=[], site=None)
    if g:
        ctx.update(open_row=secops_store.open_attendance(cid, g["id"]), duties=secops_store.guard_duties(cid, g["id"], 7),
                   recent=secops_store.list_attendance(cid, guard_id=g["id"], limit=10),
                   kit=secops_store.outstanding_kit(cid, guard_id=g["id"]))
    else:
        ctx.update(kit=[])
    return templates.TemplateResponse("secops/me.html", ctx)


@router.post("/me/checkin", name="secops_me_checkin")
async def me_checkin(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    g = _guard_for_request(request)
    if not g:
        flash(request, "Your user is not linked to a guard profile", "error")
        return RedirectResponse("/secops/me", status_code=303)
    data = await _form(request)
    res = secops_store.record_checkin(cid, guard_id=g["id"], action=data.get("action") or "in", source="self",
                                      lat=data.get("lat") or None, lng=data.get("lng") or None, note=data.get("note") or "",
                                      recorded_by=_actor(request))
    if res["ok"]:
        if (data.get("action") or "in") == "in":
            flash(request, "Checked in" + (" (late)" if (res.get("attendance") or {}).get("status") == "late" else "")
                  + (" — no rostered duty found, recorded at your home site" if res.get("unscheduled") else ""), "success")
        else:
            att = res.get("attendance") or {}
            flash(request, f"Checked out — {att.get('hours_worked', 0)} h worked", "success")
    else:
        flash(request, res.get("error") or "Failed", "error")
    return RedirectResponse("/secops/me", status_code=303)


@router.post("/api/checkin", name="secops_api_checkin")
async def api_checkin(request: Request):
    """JSON check-in/out for the Telegram bot: {guard_no|employee_id|guard_id, action: in|out, lat, lng, company_id?}.
    Auth: logged-in session / Bearer token, or header X-SecOps-Token == SECOPS_API_TOKEN."""
    token = os.environ.get("SECOPS_API_TOKEN", "").strip()
    header = request.headers.get("X-SecOps-Token", "").strip()
    if not get_current_user(request) and not (token and header == token):
        return JSONResponse({"ok": False, "error": "Authentication required"}, status_code=401)
    try:
        body = await request.json()
    except Exception:
        body = {}
    cid = body.get("company_id") or current_company(request)
    res = secops_store.record_checkin(cid, guard_id=body.get("guard_id"), guard_no=body.get("guard_no"),
                                      employee_id=body.get("employee_id"), action=body.get("action") or "in",
                                      source="telegram", lat=body.get("lat"), lng=body.get("lng"), note=body.get("note") or "",
                                      recorded_by=body.get("recorded_by") or "telegram")
    status = 200 if res["ok"] else 400
    return JSONResponse(json.loads(json.dumps(res, default=_json_default)), status_code=status)


# ══════════════════════════════════════════════════════════════════
#  Kit
# ══════════════════════════════════════════════════════════════════
@router.get("/kit", name="secops_kit")
async def kit(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    return templates.TemplateResponse("secops/kit_items.html", _ctx(
        request, "kit", items=secops_store.list_kit_items(cid), lookups=secops_store.lookups(cid)))


@router.post("/kit/items/new", name="secops_kit_item_new")
async def kit_item_new(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    it = secops_store.create_kit_item(cid, await _form(request))
    flash(request, "Kit item created" if it else "Failed to create kit item", "success" if it else "error")
    return RedirectResponse("/secops/kit", status_code=303)


@router.post("/kit/items/{item_id}/edit", name="secops_kit_item_edit")
async def kit_item_edit(item_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    ok = secops_store.update_kit_item(cid, item_id, await _form(request))
    flash(request, "Kit item updated" if ok else "Failed to update kit item", "success" if ok else "error")
    return RedirectResponse("/secops/kit", status_code=303)


@router.get("/kit/issues", name="secops_kit_issues")
async def kit_issues(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    qp = request.query_params
    co = qp.get("client_owned")
    rows = secops_store.list_kit_issues(cid, qp.get("guard_id") or None, qp.get("site_id") or None, qp.get("status") or None,
                                        True if co == "1" else False if co == "0" else None)
    if qp.get("export") == "xlsx":
        return _xlsx("Kit issues", ["Item", "Kind", "Guard", "Site", "Serial", "Size", "Qty", "Issued", "Expected return", "Returned", "Status", "Client-owned"],
                     [[r["item_name"], r["kind"], r.get("guard_name"), r.get("site_name"), r["serial_no"], r["size"], r["qty"], r["issued_at"],
                       r.get("expected_return"), r.get("returned_at"), r["status"], "yes" if r["client_owned"] else "no"] for r in rows], "kit_issues.xlsx")
    return templates.TemplateResponse("secops/kit_issues.html", _ctx(
        request, "kit_issues", issues=rows, guards=[g for g in secops_store.list_guards(cid) if g["status"] != "exited"],
        sites=secops_store.list_sites(cid), guard_filter=qp.get("guard_id") or "", site_filter=qp.get("site_id") or "",
        status_filter=qp.get("status") or "", client_owned_filter=co or ""))


@router.get("/kit/issue", name="secops_kit_issue_get")
async def kit_issue_get(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    qp = request.query_params
    return templates.TemplateResponse("secops/kit_issue_form.html", _ctx(
        request, "kit_issues", items=secops_store.list_kit_items(cid, active_only=True),
        guards=[g for g in secops_store.list_guards(cid, status="active")], sites=secops_store.list_sites(cid, status="active"),
        preset={"guard_id": qp.get("guard_id") or "", "site_id": qp.get("site_id") or ""}))


@router.post("/kit/issue", name="secops_kit_issue_post")
async def kit_issue_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    row, err = secops_store.issue_kit(cid, data, _actor(request))
    if row:
        flash(request, "Kit issued" + (" — inventory movement posted" if row.get("inventory_posted") else ""), "success")
        back = f"/secops/guards/{data['guard_id']}" if data.get("guard_id") else f"/secops/sites/{data['site_id']}#kit" if data.get("site_id") else "/secops/kit/issues"
        return RedirectResponse(back, status_code=303)
    flash(request, err or "Failed to issue kit", "error")
    return RedirectResponse("/secops/kit/issue", status_code=303)


@router.post("/kit/issues/{issue_id}/return", name="secops_kit_return")
async def kit_return(issue_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    ok = secops_store.return_kit(cid, issue_id, data.get("condition_in") or "good", "returned", _actor(request), data.get("notes") or "")
    flash(request, "Kit returned" if ok else "Failed to return kit", "success" if ok else "error")
    return RedirectResponse(data.get("back") or "/secops/kit/issues", status_code=303)


@router.post("/kit/issues/{issue_id}/writeoff", name="secops_kit_writeoff")
async def kit_writeoff(issue_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    data = await _form(request)
    status = data.get("status") if data.get("status") in ("lost", "damaged") else "lost"
    ok = secops_store.return_kit(cid, issue_id, data.get("condition_in") or status, status, _actor(request), data.get("notes") or "")
    flash(request, f"Kit written off as {status}" if ok else "Failed to write off kit", "success" if ok else "error")
    return RedirectResponse(data.get("back") or "/secops/kit/issues", status_code=303)


# ══════════════════════════════════════════════════════════════════
#  Incidents
# ══════════════════════════════════════════════════════════════════
@router.get("/incidents", name="secops_incidents")
async def incidents(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    qp = request.query_params
    rows = secops_store.list_incidents(cid, qp.get("status") or None, qp.get("site_id") or None, qp.get("severity") or None,
                                       qp.get("category") or None, qp.get("from") or None, qp.get("to") or None)
    if qp.get("export") == "xlsx":
        return _xlsx("Incidents", ["No", "Occurred", "Site", "Client", "Category", "Severity", "Status", "Guard", "Police", "Police ref", "Description"],
                     [[r["incident_no"], r["occurred_at"], r.get("site_name"), r.get("client_name"), r["category"], r["severity"], r["status"],
                       r.get("guard_name"), "yes" if r["reported_to_police"] else "no", r["police_ref"], r["description"]] for r in rows], "incidents.xlsx")
    counts = {s: sum(1 for r in rows if r["status"] == s) for s in L.INCIDENT_STATUSES}
    return templates.TemplateResponse("secops/incidents.html", _ctx(
        request, "incidents", incidents=rows, counts=counts, sites=secops_store.list_sites(cid), status_filter=qp.get("status") or "",
        site_filter=qp.get("site_id") or "", severity_filter=qp.get("severity") or "", category_filter=qp.get("category") or "",
        date_from=qp.get("from") or "", date_to=qp.get("to") or ""))


@router.get("/incidents/new", name="secops_incident_new_get")
async def incident_new_get(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    qp = request.query_params
    return templates.TemplateResponse("secops/incident_form.html", _ctx(
        request, "incidents", incident={"site_id": qp.get("site_id") or "", "guard_id": qp.get("guard_id") or "", "severity": "low",
                                        "category": "other", "occurred_at": datetime.now().strftime("%Y-%m-%dT%H:%M")},
        is_edit=False, sites=secops_store.list_sites(cid), posts=secops_store.list_posts(cid),
        guards=[g for g in secops_store.list_guards(cid) if g["status"] != "exited"]))


@router.post("/incidents/new", name="secops_incident_new_post")
async def incident_new_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    form = await request.form()
    data = {k: v for k, v in form.items() if not hasattr(v, "filename")}
    inc = secops_store.create_incident(cid, data, _actor(request))
    if not inc:
        flash(request, "Failed to record incident — site and description are required", "error")
        return RedirectResponse("/secops/incidents/new", status_code=303)
    upload = form.get("file")
    if upload is not None and getattr(upload, "filename", ""):
        doc, err = await _store_upload(request, cid, inc["id"], upload, "incident_photo")
        if err:
            flash(request, f"Incident saved, but the file was not stored: {err}", "error")
    if inc["severity"] in ("high", "critical"):
        _notify_topic(cid, f"🚨 <b>{inc['severity'].upper()} incident {inc['incident_no']}</b> — {inc['category'].replace('_', ' ')}\n{inc['description'][:300]}")
    flash(request, f"Incident {inc['incident_no']} recorded", "success")
    return RedirectResponse(f"/secops/incidents/{inc['id']}", status_code=303)


@router.get("/incidents/{incident_id}", name="secops_incident_detail")
async def incident_detail(incident_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    inc = secops_store.get_incident(cid, incident_id)
    if not inc:
        flash(request, "Incident not found", "error")
        return RedirectResponse("/secops/incidents", status_code=303)
    return templates.TemplateResponse("secops/incident_detail.html", _ctx(
        request, "incidents", incident=inc, documents=_documents(cid, incident_id), sites=secops_store.list_sites(cid),
        posts=secops_store.list_posts(cid, inc["site_id"]), guards=[g for g in secops_store.list_guards(cid) if g["status"] != "exited"]))


@router.post("/incidents/{incident_id}/edit", name="secops_incident_edit")
async def incident_edit(incident_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    ok = secops_store.update_incident(cid, incident_id, await _form(request))
    flash(request, "Incident updated" if ok else "Failed to update incident", "success" if ok else "error")
    return RedirectResponse(f"/secops/incidents/{incident_id}", status_code=303)


@router.post("/incidents/{incident_id}/status", name="secops_incident_status")
async def incident_status(incident_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    ok = secops_store.set_incident_status(cid, incident_id, data.get("status", ""))
    flash(request, f"Incident marked {data.get('status', '')}" if ok else "Failed to update status", "success" if ok else "error")
    return RedirectResponse(f"/secops/incidents/{incident_id}", status_code=303)


@router.post("/incidents/{incident_id}/notify-client", name="secops_incident_notify_client")
async def incident_notify_client(incident_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    inc = secops_store.get_incident(cid, incident_id)
    if not inc:
        flash(request, "Incident not found", "error")
        return RedirectResponse("/secops/incidents", status_code=303)
    data = await _form(request)
    note = (data.get("note") or "").strip()
    emailed = False
    if data.get("send_email") and inc.get("client_email"):
        try:
            from email_service import send_email
            html = (f"<p>Incident <b>{inc['incident_no']}</b> at <b>{inc['site_name']}</b> on {inc['occurred_at']:%Y-%m-%d %H:%M}</p>"
                    f"<p>Category: {inc['category'].replace('_', ' ')} — Severity: {inc['severity']}</p>"
                    f"<p>{inc['description']}</p><p>Actions taken: {inc['actions_taken'] or '-'}</p>"
                    + (f"<p>{note}</p>" if note else ""))
            emailed = bool(send_email(inc["client_email"], f"Security incident {inc['incident_no']} — {inc['site_name']}", html, category="secops"))
        except Exception as e:
            logger.warning("secops client email failed: %s", e)
    secops_store.note_client_notified(cid, incident_id, note + (" (e-mail sent)" if emailed else ""))
    flash(request, "Client notification recorded" + (" and e-mailed" if emailed else ""), "success")
    return RedirectResponse(f"/secops/incidents/{incident_id}", status_code=303)


@router.post("/incidents/{incident_id}/upload", name="secops_incident_upload")
async def incident_upload(incident_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    if not secops_store.get_incident(cid, incident_id):
        flash(request, "Incident not found", "error")
        return RedirectResponse("/secops/incidents", status_code=303)
    form = await request.form()
    doc, err = await _store_upload(request, cid, incident_id, form.get("file"), (form.get("kind") or "incident_photo")[:40])
    flash(request, "File uploaded" if doc else (err or "Upload failed"), "success" if doc else "error")
    return RedirectResponse(f"/secops/incidents/{incident_id}#documents", status_code=303)


@router.get("/incidents/{incident_id}/file/{doc_id}", name="secops_incident_file")
async def incident_file(incident_id: str, doc_id: str, request: Request, user=Depends(login_required)):
    return _serve_document(request, current_company(request), incident_id, doc_id, f"/secops/incidents/{incident_id}#documents")


# ══════════════════════════════════════════════════════════════════
#  Payroll feed
# ══════════════════════════════════════════════════════════════════
FEED_HEADERS = ["Employee ID", "Name", "Guard No", "Period", "Days Present", "Days Absent", "Hours Worked", "Night Hours",
                "Overtime Hours", "Hourly Rate", "Overtime Pay", "Night Allowance"]


def _feed_rows(rows):
    return [[r["employee_id"], r.get("name"), r.get("guard_no"), r["period"], r["days_present"], r["days_absent"], r["hours_worked"],
             r["night_hours"], r["overtime_hours"], r["hourly_rate"], r["overtime_pay"], r["night_allowance"]] for r in rows]


@router.get("/payroll-feed", name="secops_payroll_feed")
async def payroll_feed(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    period = (request.query_params.get("period") or L.previous_month()).strip()
    rows = secops_store.list_feed(cid, period)
    totals = {k: sum((Decimal(str(r[k] or 0)) for r in rows), Decimal(0)) for k in ("hours_worked", "night_hours", "overtime_hours", "overtime_pay", "night_allowance")}
    totals["days_present"] = sum(r["days_present"] for r in rows)
    totals["days_absent"] = sum(r["days_absent"] for r in rows)
    return templates.TemplateResponse("secops/payroll_feed.html", _ctx(
        request, "payroll", period=period, rows=rows, totals=totals, periods=secops_store.feed_periods(cid),
        settings=secops_store.get_settings(cid), exported=any(r["status"] == "exported" for r in rows)))


@router.post("/payroll-feed/generate", name="secops_payroll_feed_generate")
async def payroll_feed_generate(request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    data = await _form(request)
    period = (data.get("period") or "").strip()
    res = secops_store.generate_feed(cid, period, force=data.get("force") == "1")
    if res.get("error"):
        flash(request, res["error"], "error")
    else:
        flash(request, f"Payroll feed for {period}: {res['rows']} guard(s)", "success")
    return RedirectResponse(f"/secops/payroll-feed?period={period}", status_code=303)


@router.get("/payroll-feed/export", name="secops_payroll_feed_export")
async def payroll_feed_export(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    period = (request.query_params.get("period") or L.previous_month()).strip()
    rows = secops_store.list_feed(cid, period)
    secops_store.mark_feed_exported(cid, period)
    return _xlsx(f"Payroll feed {period}", FEED_HEADERS, _feed_rows(rows), f"secops_payroll_feed_{period}.xlsx",
                 "Overtime pay and night allowance per guard — import as payroll allowances")


@router.post("/payroll-feed/push", name="secops_payroll_feed_push")
async def payroll_feed_push(request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    data = await _form(request)
    period = (data.get("period") or "").strip()
    res = secops_store.push_to_payroll(cid, period)
    if res.get("error"):
        flash(request, f"Exported file only — {res['error']}", "error")
    else:
        flash(request, f"Pushed {res['pushed']} allowance line(s) to payroll", "success")
    return RedirectResponse(f"/secops/payroll-feed?period={period}", status_code=303)


# ══════════════════════════════════════════════════════════════════
#  Settings
# ══════════════════════════════════════════════════════════════════
@router.get("/settings", name="secops_settings")
async def settings(request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    return templates.TemplateResponse("secops/settings.html", _ctx(
        request, "settings", settings=secops_store.get_settings(cid), holidays=secops_store.list_holidays(cid),
        weekend_days=L.weekend_day_set(secops_store.get_settings(cid).get("weekend_days"))))


@router.post("/settings", name="secops_settings_post")
async def settings_post(request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    form = await request.form()
    data = {k: v for k, v in form.items()}
    data["weekend_days"] = ",".join(form.getlist("weekend_days")) if hasattr(form, "getlist") else data.get("weekend_days", "6")
    ok = secops_store.save_settings(cid, data)
    flash(request, "Settings saved" if ok else "Failed to save settings", "success" if ok else "error")
    return RedirectResponse("/secops/settings", status_code=303)


@router.post("/settings/holidays/new", name="secops_holiday_new")
async def holiday_new(request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    data = await _form(request)
    ok = secops_store.add_holiday(cid, data.get("holiday_date"), data.get("name") or "")
    flash(request, "Holiday added" if ok else "Failed to add holiday — date is required", "success" if ok else "error")
    return RedirectResponse("/secops/settings#holidays", status_code=303)


@router.post("/settings/holidays/{holiday_id}/delete", name="secops_holiday_delete")
async def holiday_delete(holiday_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    ok = secops_store.delete_holiday(cid, holiday_id)
    flash(request, "Holiday removed" if ok else "Failed to remove holiday", "success" if ok else "error")
    return RedirectResponse("/secops/settings#holidays", status_code=303)


# ══════════════════════════════════════════════════════════════════
#  Reports
# ══════════════════════════════════════════════════════════════════
REPORTS = {
    "coverage": {"title": "Deployment coverage per site and day", "icon": "bi-grid-3x3-gap", "filters": ("date", "site", "client"),
                 "columns": [("duty_date", "Date"), ("site_name", "Site"), ("client_name", "Client"), ("required", "Required"),
                             ("assigned", "Assigned"), ("present", "Present"), ("absent", "Absent"), ("uncovered", "Uncovered"), ("coverage_pct", "Coverage %")]},
    "attendance": {"title": "Attendance register per site", "icon": "bi-calendar-check", "filters": ("date", "site", "guard"),
                   "columns": [("site_name", "Site"), ("duty_date", "Date"), ("guard_no", "Guard no"), ("guard_name", "Guard"), ("post_name", "Post"),
                               ("shift_name", "Shift"), ("status", "Status"), ("check_in_at", "Check-in"), ("check_out_at", "Check-out"),
                               ("hours_worked", "Hours"), ("night_hours", "Night hours"), ("overtime_hours", "Overtime"), ("late_minutes", "Late (min)"), ("source", "Source")]},
    "absenteeism": {"title": "Absenteeism and lateness per guard", "icon": "bi-person-x", "filters": ("date", "site", "guard"),
                    "columns": [("guard_no", "Guard no"), ("guard_name", "Guard"), ("scheduled", "Recorded days"), ("present", "Present"), ("late", "Late"),
                                ("absent", "Absent"), ("leave", "Leave"), ("late_minutes", "Late minutes"), ("absence_pct", "Absence %")]},
    "hours": {"title": "Hours, overtime and night work per guard", "icon": "bi-clock-history", "filters": ("date", "site", "guard"),
              "columns": [("guard_no", "Guard no"), ("name", "Guard"), ("days_present", "Days present"), ("days_absent", "Days absent"),
                          ("hours_worked", "Hours"), ("night_hours", "Night hours"), ("overtime_hours", "Overtime hours"), ("hourly_rate", "Hourly rate"),
                          ("overtime_pay", "Overtime pay"), ("night_allowance", "Night allowance")]},
    "kit_guards": {"title": "Kit outstanding per guard", "icon": "bi-person-badge", "filters": ("guard",),
                   "columns": [("holder", "Guard"), ("guard_no", "Guard no"), ("item_name", "Item"), ("kind", "Kind"), ("serial_no", "Serial"), ("size", "Size"),
                               ("qty", "Qty"), ("issued_at", "Issued"), ("expected_return", "Expected return"), ("overdue", "Overdue"), ("condition_out", "Condition")]},
    "kit_sites": {"title": "Client-owned equipment per site", "icon": "bi-building", "filters": ("site",),
                  "columns": [("site_name", "Site"), ("item_name", "Item"), ("kind", "Kind"), ("serial_no", "Serial"), ("qty", "Qty"), ("issued_at", "Issued"),
                              ("client_owned", "Client-owned"), ("condition_out", "Condition")]},
    "licences": {"title": "Licence and permit expiry", "icon": "bi-card-checklist", "filters": (),
                 "columns": [("guard_no", "Guard no"), ("guard_name", "Guard"), ("grade", "Grade"), ("home_site", "Home site"), ("document", "Document"),
                             ("expiry", "Expiry"), ("days", "Days left"), ("status", "Status")]},
    "incidents": {"title": "Incidents by site, category and severity", "icon": "bi-exclamation-triangle", "filters": ("date", "site"),
                  "columns": [("site_name", "Site"), ("category", "Category"), ("total", "Total"), ("low", "Low"), ("medium", "Medium"), ("high", "High"),
                              ("critical", "Critical"), ("open", "Open")]},
    "incident_list": {"title": "Incident register", "icon": "bi-list-ul", "filters": ("date", "site", "category", "severity"),
                      "columns": [("incident_no", "No"), ("occurred_at", "Occurred"), ("site_name", "Site"), ("client_name", "Client"), ("category", "Category"),
                                  ("severity", "Severity"), ("status", "Status"), ("guard_name", "Guard"), ("reported_to_police", "Police"),
                                  ("client_notified", "Client notified"), ("description", "Description")]},
    "billing": {"title": "Client billing summary", "icon": "bi-cash-coin", "filters": ("date", "client", "site"),
                "columns": [("client_name", "Client"), ("site_name", "Site"), ("period", "Month"), ("guard_days", "Guard-days"), ("hours", "Billable hours"),
                            ("night_hours", "Night hours"), ("overtime_hours", "Overtime hours"), ("rate", "Rate / guard-hour"), ("amount", "Amount")]},
}
SUM_KEYS = {"required", "assigned", "present", "absent", "uncovered", "hours_worked", "night_hours", "overtime_hours", "overtime_pay",
            "night_allowance", "scheduled", "late", "leave", "late_minutes", "days_present", "days_absent", "qty", "total", "low", "medium",
            "high", "critical", "open", "guard_days", "hours", "amount"}


def _run_report(cid: str, key: str, f: dict) -> list:
    s = secops_store
    co = True if f["client_owned"] == "1" else False if f["client_owned"] == "0" else None
    if key == "coverage":
        return s.report_coverage(cid, f["date_from"], f["date_to"], f["site_id"] or None, f["client_id"] or None)
    if key == "attendance":
        return s.report_attendance_register(cid, f["date_from"], f["date_to"], f["site_id"] or None, f["guard_id"] or None)
    if key == "absenteeism":
        return s.report_absenteeism(cid, f["date_from"], f["date_to"], f["site_id"] or None, f["guard_id"] or None)
    if key == "hours":
        return s.report_hours(cid, f["date_from"], f["date_to"], f["site_id"] or None, f["guard_id"] or None)
    if key == "kit_guards":
        return [r for r in s.report_kit_outstanding(cid, guard_id=f["guard_id"] or None, client_owned=co) if r["holder_type"] == "guard"]
    if key == "kit_sites":
        return [r for r in s.report_kit_outstanding(cid, site_id=f["site_id"] or None, client_owned=True if co is None else co) if r["holder_type"] == "site"]
    if key == "licences":
        return s.report_licence_expiry(cid)
    if key == "incidents":
        return s.report_incident_summary(cid, f["date_from"], f["date_to"], f["site_id"] or None)
    if key == "incident_list":
        return s.report_incidents(cid, f["date_from"], f["date_to"], f["site_id"] or None, f["category"] or None, f["severity"] or None)
    if key == "billing":
        return s.report_billing(cid, f["date_from"], f["date_to"], f["client_id"] or None, f["site_id"] or None)
    return []


@router.get("/reports", name="secops_reports")
async def reports_index(request: Request, user=Depends(login_required)):
    return templates.TemplateResponse("secops/reports_index.html", _ctx(request, "reports", reports=REPORTS))


@router.get("/reports/{key}", name="secops_report")
async def report(key: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    cfg = REPORTS.get(key)
    if not cfg:
        flash(request, "Unknown report", "error")
        return RedirectResponse("/secops/reports", status_code=303)
    f = _date_filters(request)
    rows = _run_report(cid, key, f)
    cols = cfg["columns"]
    totals = {}
    if rows:
        for k, _label in cols:
            if k in SUM_KEYS:
                totals[k] = sum((Decimal(str(r.get(k) or 0)) for r in rows), Decimal(0))
    if request.query_params.get("export") == "xlsx":
        return _xlsx(cfg["title"], [lbl for _k, lbl in cols], [[r.get(k) for k, _l in cols] for r in rows],
                     f"secops_{key}_{f['date_from']}_{f['date_to']}.xlsx", f"{f['date_from']} → {f['date_to']}")
    return templates.TemplateResponse("secops/report_table.html", _ctx(
        request, f"report_{key}", report_key=key, report_title=cfg["title"], columns=cols, rows=rows, totals=totals, filters=f,
        filter_fields=cfg["filters"], sites=secops_store.list_sites(cid), clients=secops_store.list_clients(cid),
        guards=[g for g in secops_store.list_guards(cid) if g["status"] != "exited"]))
