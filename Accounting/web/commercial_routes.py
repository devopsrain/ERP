"""
Commercial — Sales Order & Marketing management routes (/commercial).

Route names are ``commercial_*``; templates live in templates/commercial/.
Every printable document (proforma, sales order, MO request, delivery
instruction, dispatch note, invoice, credit note) is rendered from one
normalised ``doc`` dict by ``_doc_dict`` → HTML print view (doc_print.html),
PDF (reportlab via erca_forms.PdfDoc, built-in fallback) and Excel (openpyxl).
"""
from __future__ import annotations

import io
import json
import logging
import os
import tempfile
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, Depends, Request
from fastapi.responses import FileResponse, RedirectResponse, Response

import commercial_logic as L
from commercial_data_store import commercial_store as store
from commercial_logic import D, q2, q3
from deps import current_company, flash, login_required, require_auth, template_context
from template_engine import templates

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/commercial", tags=["commercial"])
manager_required = require_auth("manager")

_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
P = "/commercial"


@router.on_event("startup")
async def _startup():
    store.ensure_schema()


# ── helpers ───────────────────────────────────────────────────────
def _actor(request: Request) -> str:
    return request.session.get("username", "") or ""


def _priv(request: Request) -> str:
    return request.session.get("privilege_level", "viewer") or "viewer"


def _ctx(request: Request, **extra) -> dict:
    ctx = template_context(request)
    ctx.update(active_page=extra.pop("active_page", ""), today=date.today(), L=L,
               so_statuses=L.SO_STATUSES, payment_methods=L.PAYMENT_METHODS, segments=L.SEGMENTS,
               lead_statuses=L.LEAD_STATUSES, campaign_statuses=L.CAMPAIGN_STATUSES)
    ctx.update(extra)
    return ctx


async def _form(request: Request) -> dict:
    form = await request.form()
    return {k: v for k, v in form.items()}


async def _form_with_lines(request: Request, fields=("item_code", "description", "size", "color", "unit", "packaging",
                                                      "quantity", "unit_price", "discount", "ref_no")) -> tuple:
    """Header dict + list of line dicts from parallel ``line_<field>`` inputs."""
    form = await request.form()
    data = {k: v for k, v in form.items() if not k.startswith("line_")}
    cols = {f: form.getlist(f"line_{f}") for f in fields}
    n = max((len(v) for v in cols.values()), default=0)
    lines = [{f: (cols[f][i] if i < len(cols[f]) else "") for f in fields} for i in range(n)]
    return data, lines


def _qty_map(form: dict, prefix: str = "qty_") -> Dict[str, Any]:
    return {k[len(prefix):]: v for k, v in form.items() if k.startswith(prefix) and str(v).strip() not in ("", "0")}


def _json_default(o):
    if isinstance(o, Decimal):
        return float(o)
    if isinstance(o, (date, datetime)):
        return o.isoformat()
    return str(o)


def _filters(request: Request) -> dict:
    qp = request.query_params
    return {k: (qp.get(k) or "") for k in ("date_from", "date_to", "customer_id", "item_code", "sales_rep_id",
                                            "territory_id", "status", "segment", "period", "q")}


def _lookups(cid: str) -> dict:
    return dict(customers=store.list_customers(cid, active_only=True, with_balance=False),
                products=store.list_products(cid, active_only=True), reps=store.list_reps(cid, active_only=True),
                territories=store.list_territories(cid))


def _money(v) -> str:
    return f"{q2(v):,.2f}"


def _d(v) -> str:
    if v in (None, ""):
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, date):
        return v.isoformat()
    return str(v)[:16]


# ── normalised printable document ────────────────────────────────
def _doc_dict(kind: str, obj: dict, cid: str) -> dict:
    """One shape for every print/PDF/Excel output."""
    s = store.get_settings(cid)
    doc = {"kind": kind, "header_info": obj.get("header_info") or s["header_info"], "footer_info": obj.get("footer_info") or s["footer_info"],
           "iso_doc_no": obj.get("iso_doc_no") or s["iso_doc_no"], "meta": [], "columns": [], "lines": [], "totals": [], "notes": obj.get("notes") or "",
           "signatures": [("Prepared by", obj.get("prepared_by") or ""), ("Checked by", obj.get("checked_by") or ""),
                          ("Approved by", obj.get("approved_by") or "")], "copies": [], "extra_rows": []}
    money_cols = ()
    if kind == "proforma":
        doc.update(title="PROFORMA INVOICE", number=obj["proforma_no"], back=f"{P}/proformas/{obj['id']}")
        doc["meta"] = [("Customer", obj["customer_name"]), ("TIN", obj.get("customer_tin") or ""), ("Date", _d(obj["proforma_date"])),
                       ("Valid until", _d(obj.get("valid_until"))), ("Payment method", obj.get("payment_method") or ""),
                       ("Payment info", obj.get("payment_info") or ""), ("Status", obj["status"])]
        doc["columns"] = ["#", "Ref", "Item code", "Size", "Description", "Qty", "Colour", "Unit", "Unit price", "Total"]
        doc["lines"] = [[l["line_no"], l["ref_no"], l["item_code"], l["size"], l["description"], q3(l["quantity"]), l["color"], l["unit"],
                         q2(l["unit_price"]), q2(l["line_total"])] for l in obj["lines"]]
        money_cols = (8, 9)
    elif kind == "sales_order":
        doc.update(title="SALES ORDER", number=obj["so_no"], back=f"{P}/orders/{obj['id']}")
        doc["meta"] = [("Customer", obj["customer_name"]), ("Order date", _d(obj["order_date"])), ("Required date", _d(obj.get("required_date"))),
                       ("Payment terms", obj.get("payment_terms") or ("Credit" if obj.get("credit_sale") else "Cash")),
                       ("Sales rep", obj.get("sales_rep_name") or ""), ("Status", obj["status"]), ("Proforma ref", obj.get("proforma_id") and "yes" or "")]
        doc["columns"] = ["#", "Item code", "Size", "Description", "Colour", "Unit", "Packaging", "Qty", "Unit price", "Discount", "Total"]
        doc["lines"] = [[l["line_no"], l["item_code"], l["size"], l["description"], l["color"], l["unit"], l["packaging"], q3(l["quantity"]),
                         q2(l["unit_price"]), q2(l["discount"]), q2(l["line_total"])] for l in obj["lines"]]
        money_cols = (8, 9, 10)
        doc["extra_rows"] = [("Starting", _d(obj.get("starting_reading"))), ("Ending", _d(obj.get("ending_reading"))),
                             ("Difference", _d(obj.get("difference_reading"))), ("Verification", obj.get("verification_note") or "")]
        doc["copies"] = [("Finance copy", obj.get("copy_finance")), ("Sales copy", obj.get("copy_sales")), ("First copy (customer)", obj.get("copy_first"))]
        doc["signatures"] += [(f"Approver {a['seq']} — {a['role_label']}", (a.get("decided_by") or "") + (f" ({a['decision']})" if a.get("decision") else ""))
                              for a in obj.get("approvals", [])]
    elif kind == "mo_request":
        doc.update(title="MANUFACTURING ORDER REQUEST", number=obj["request_no"], back=f"{P}/mo-requests/{obj['id']}")
        doc["meta"] = [("Sales order", obj.get("so_no") or ""), ("Customer", obj.get("customer_name") or ""), ("To factory", obj.get("to_factory") or ""),
                       ("Delivery date", _d(obj.get("delivery_date"))), ("Status", obj["status"]), ("MO no.", obj.get("mo_no") or "")]
        doc["columns"] = ["Product code", "Item description", "Size", "Colour", "Unit", "Quantity", "Packing", "Cutting length"]
        doc["lines"] = [[obj["product_code"], obj["description"], obj["size"], obj["color"], obj["unit"], q3(obj["quantity"]), obj["packing"], obj["cutting_length"]]]
        doc["notes"] = obj.get("comment") or ""
        doc["signatures"] += [("Manager", obj.get("manager") or "")]
    elif kind == "delivery":
        doc.update(title="DELIVERY INSTRUCTION FOR INVENTORY", number=obj["di_no"], back=f"{P}/deliveries/{obj['id']}")
        doc["meta"] = [("Sales order", obj.get("so_no") or ""), ("Customer", obj["customer_name"]), ("Date", _d(obj["di_date"])),
                       ("Deliver to", obj.get("deliver_to") or ""), ("Warehouse", obj.get("warehouse") or ""), ("Status", obj["status"])]
        doc["columns"] = ["#", "Product code", "Item description", "Size", "Colour", "Unit", "Packaging", "Quantity"]
        doc["lines"] = [[l["line_no"], l["product_code"], l["description"], l["size"], l["color"], l["unit"], l["packaging"], q3(l["quantity"])] for l in obj["lines"]]
    elif kind == "dispatch":
        doc.update(title="DISPATCH NOTE", number=obj["dispatch_no"], back=f"{P}/deliveries/{obj['di_id']}")
        doc["meta"] = [("Delivery instruction", obj.get("di_no") or ""), ("Sales order", obj.get("so_no") or ""), ("Customer", obj.get("customer_name") or ""),
                       ("Vehicle", obj.get("vehicle") or ""), ("Driver", f"{obj.get('driver') or ''} {obj.get('driver_phone') or ''}".strip()),
                       ("Dispatched", _d(obj.get("dispatched_at"))), ("Delivered", _d(obj.get("delivered_at"))), ("Received by", obj.get("received_by") or ""),
                       ("Status", obj["status"])]
        doc["columns"] = ["#", "Product code", "Item description", "Size", "Colour", "Unit", "Packaging", "Quantity"]
        doc["lines"] = [[l["line_no"], l["product_code"], l["description"], l["size"], l["color"], l["unit"], l["packaging"], q3(l["quantity"])] for l in obj.get("lines", [])]
        doc["signatures"] = [("Dispatched by", obj.get("created_by") or ""), ("Driver", obj.get("driver") or ""), ("Received by (customer)", obj.get("received_by") or "")]
    elif kind == "invoice":
        doc.update(title="SALES INVOICE" + (" (CREDIT)" if obj.get("credit_sale") else ""), number=obj["invoice_no"], back=f"{P}/invoices/{obj['id']}")
        doc["meta"] = [("Customer", obj["customer_name"]), ("TIN", obj.get("customer_tin") or ""), ("Invoice date", _d(obj["invoice_date"])),
                       ("Due date", _d(obj.get("due_date"))), ("Sales order", obj.get("so_no") or ""), ("Status", obj["status"]),
                       ("ERCA no.", obj.get("erca_number") or "")]
        doc["columns"] = ["#", "Item code", "Description", "Size", "Colour", "Unit", "Qty", "Unit price", "Discount", "Total"]
        doc["lines"] = [[l["line_no"], l["item_code"], l["description"], l["size"], l["color"], l["unit"], q3(l["quantity"]), q2(l["unit_price"]),
                         q2(l["discount"]), q2(l["line_total"])] for l in obj["lines"]]
        money_cols = (7, 8, 9)
    elif kind == "credit_note":
        doc.update(title="CREDIT NOTE / SALES RETURN", number=obj.get("credit_note_no") or obj["rn_no"], back=f"{P}/returns/{obj['id']}")
        doc["meta"] = [("Return no.", obj["rn_no"]), ("Invoice", obj.get("invoice_no") or ""), ("Customer", obj["customer_name"]),
                       ("Date", _d(obj["return_date"])), ("Reason", obj.get("reason") or ""), ("Status", obj["status"])]
        doc["columns"] = ["#", "Item code", "Description", "Unit", "Qty", "Unit price", "Total"]
        doc["lines"] = [[l["line_no"], l["item_code"], l["description"], l["unit"], q3(l["quantity"]), q2(l["unit_price"]), q2(l["line_total"])] for l in obj["lines"]]
        money_cols = (5, 6)
        doc["signatures"] = [("Prepared by", obj.get("prepared_by") or ""), ("Approved by", obj.get("approved_by") or "")]
    if kind in ("proforma", "sales_order", "invoice"):
        doc["totals"] = [("Subtotal", q2(obj["subtotal"])), ("Discount", q2(obj["discount_total"])),
                         (f"VAT {D(obj.get('vat_rate') or L.VAT_RATE) * 100:.0f}%", q2(obj["vat_total"])), ("GRAND TOTAL", q2(obj["grand_total"]))]
        if kind == "invoice":
            doc["totals"] += [("Paid", q2(obj["paid_total"])), ("Credited", q2(obj.get("credited_total"))), ("Outstanding", q2(obj.get("outstanding")))]
    elif kind == "credit_note":
        doc["totals"] = [("Net", q2(obj["subtotal"])), ("VAT", q2(obj["vat_total"])), ("CREDIT AMOUNT", q2(obj["amount"]))]
    doc["money_cols"] = list(money_cols)
    doc["currency"] = obj.get("currency") or "ETB"
    return doc


def _doc_pdf(doc: dict) -> bytes:
    try:
        from erca_forms import PdfDoc, _PAGE_W
    except Exception:  # pragma: no cover — erca module missing: plain-text PDF
        return _plain_pdf(doc)
    pdf = PdfDoc()
    y = 40
    for line in (doc["header_info"] or "").splitlines()[:4]:
        pdf.text(_PAGE_W / 2, y + 10, line, 9, y == 40, "center"); y += 12
    pdf.text(_PAGE_W / 2, y + 22, doc["title"], 14, True, "center")
    pdf.text(_PAGE_W - 40, y + 22, f"No. {doc['number']}", 10, True, "right")
    if doc["iso_doc_no"]:
        pdf.text(40, y + 22, f"Doc. No. {doc['iso_doc_no']}", 8)
    pdf.y = y + 36
    for i in range(0, len(doc["meta"]), 2):
        pair = doc["meta"][i:i + 2]
        pdf.text(40, pdf.y + 9, f"{pair[0][0]}: {pair[0][1]}", 8.5)
        if len(pair) > 1:
            pdf.text(_PAGE_W / 2, pdf.y + 9, f"{pair[1][0]}: {pair[1][1]}", 8.5)
        pdf.y += 12
    pdf.y += 6
    total_w = _PAGE_W - 80
    widths = []
    for i, c in enumerate(doc["columns"]):
        w = 1.0
        if c in ("#",):
            w = 0.35
        elif c in ("Description", "Item description"):
            w = 2.4
        elif c in ("Qty", "Quantity", "Unit", "Size", "Colour"):
            w = 0.7
        widths.append(w)
    scale = total_w / sum(widths)
    cols = [(c, widths[i] * scale, "right" if i in doc["money_cols"] or c in ("Qty", "Quantity") else "left") for i, c in enumerate(doc["columns"])]
    pdf.table(cols, doc["lines"], size=8)
    pdf.y += 6
    for lab, val in doc["totals"]:
        bold = lab.isupper()
        pdf.text(_PAGE_W - 180, pdf.y + 10, lab, 9, bold)
        pdf.text(_PAGE_W - 40, pdf.y + 10, f"{doc['currency']} {_money(val)}", 9, bold, "right")
        pdf.y += 13
    for lab, val in doc["extra_rows"]:
        if val:
            pdf.para(f"{lab}: {val}", 8.5)
    if doc["notes"]:
        pdf.para(f"Notes: {doc['notes']}", 8.5)
    if doc["copies"]:
        pdf.para("Distribution: " + ", ".join(f"[{'x' if on else ' '}] {lab}" for lab, on in doc["copies"]), 8)
    pdf.ensure_space(70)
    pdf.y += 16
    sigs = doc["signatures"]
    per_row = 3
    for i in range(0, len(sigs), per_row):
        for j, (lab, name) in enumerate(sigs[i:i + per_row]):
            x = 40 + j * (total_w / per_row)
            pdf.text(x, pdf.y + 10, name, 8.5)
            pdf.line(x, pdf.y + 14, x + total_w / per_row - 20, pdf.y + 14)
            pdf.text(x, pdf.y + 24, f"{lab} — signature / date", 7.5)
        pdf.y += 40
    pdf.footer((doc["footer_info"] or "Generated by EBMS — Commercial module")[:160])
    return pdf.render()


def _plain_pdf(doc: dict) -> bytes:
    """Dependency-free single-page PDF (used only if erca_forms is absent)."""
    lines = [doc["title"], f"No. {doc['number']}", ""] + [f"{k}: {v}" for k, v in doc["meta"]] + [""]
    lines += [" | ".join(doc["columns"])] + [" | ".join(str(c) for c in row) for row in doc["lines"]] + [""]
    lines += [f"{k}: {_money(v)}" for k, v in doc["totals"]]
    ops = ["BT /F1 9 Tf 40 800 Td 12 TL"]
    for ln in lines[:60]:
        ln = ln.encode("latin-1", "replace").decode("latin-1").replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        ops.append(f"({ln}) Tj T*")
    ops.append("ET")
    stream = "\n".join(ops)
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream", "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out = io.BytesIO(); out.write(b"%PDF-1.4\n"); offs = []
    for i, o in enumerate(objs, 1):
        offs.append(out.tell()); out.write(f"{i} 0 obj\n{o}\nendobj\n".encode("latin-1"))
    x = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for o in offs:
        out.write(f"{o:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{x}\n%%EOF\n".encode())
    return out.getvalue()


def _xlsx_file(title: str, headers: List[str], rows: List[list], totals: Optional[list] = None, meta: Optional[list] = None,
               extra: Optional[dict] = None) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = title[:30]
    ws.append([title]); ws["A1"].font = Font(bold=True, size=13)
    for k, v in (meta or []):
        ws.append([k, _d(v) if isinstance(v, (date, datetime)) else v])
    ws.append([])
    ws.append(headers)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([float(v) if isinstance(v, Decimal) else (v.isoformat() if isinstance(v, (date, datetime)) else v) for v in r])
    if totals:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in totals])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
    if extra:
        ws.append([]); ws.append([extra.get("title", "")]); ws[ws.max_row][0].font = Font(bold=True)
        ws.append(extra["columns"])
        for r in extra["rows"]:
            ws.append([float(v) if isinstance(v, Decimal) else v for v in r])
    for i, _h in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 16
    fd, path = tempfile.mkstemp(suffix=".xlsx"); os.close(fd)
    wb.save(path)
    return path


def _doc_response(request: Request, fmt: str, doc: dict, cid: str):
    if fmt == "pdf":
        return Response(content=_doc_pdf(doc), media_type="application/pdf",
                        headers={"Content-Disposition": f'inline; filename="{doc["number"]}.pdf"'})
    if fmt == "xlsx":
        path = _xlsx_file(f"{doc['title']} {doc['number']}", doc["columns"], doc["lines"], meta=doc["meta"],
                          totals=None, extra={"title": "Totals", "columns": ["", ""], "rows": [[k, v] for k, v in doc["totals"]]} if doc["totals"] else None)
        return FileResponse(path, filename=f"{doc['number']}.xlsx", media_type=_XLSX)
    return templates.TemplateResponse("commercial/doc_print.html", _ctx(request, doc=doc, company_id=cid))


# ── dashboard & settings ─────────────────────────────────────────
@router.get("/", name="commercial_dashboard")
async def dashboard(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    stats = store.dashboard(cid, _actor(request), _priv(request))
    chart = {"monthly_labels": [m["period"] for m in stats["monthly_revenue"]],
             "monthly_totals": [m["total"] for m in stats["monthly_revenue"]],
             "pipeline_labels": list(stats["pipeline"].keys()),
             "pipeline_counts": [v["count"] for v in stats["pipeline"].values()]}
    return templates.TemplateResponse("commercial/dashboard.html", _ctx(
        request, active_page="dashboard", stats=stats, chart_json=json.dumps(chart, default=_json_default),
        open_summary=store.open_orders_summary(cid)))


@router.get("/settings", name="commercial_settings")
async def settings_get(request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    return templates.TemplateResponse("commercial/settings.html", _ctx(
        request, active_page="settings", settings=store.get_settings(cid), price_lists=store.list_price_lists(cid)))


@router.post("/settings", name="commercial_settings_post")
async def settings_post(request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    ok = store.save_settings(cid, await _form(request))
    flash(request, "Settings saved" if ok else "Failed to save settings", "success" if ok else "error")
    return RedirectResponse(f"{P}/settings", status_code=303)


# ── customers ────────────────────────────────────────────────────
@router.get("/customers", name="commercial_customers")
async def customers(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    f = _filters(request)
    look = _lookups(cid)
    look["customers"] = store.list_customers(cid, q=f["q"] or None, segment=f["segment"] or None,
                                             territory_id=f["territory_id"] or None, sales_rep_id=f["sales_rep_id"] or None)
    return templates.TemplateResponse("commercial/customers.html", _ctx(request, active_page="customers", filters=f, **look))


@router.get("/customers/new", name="commercial_customer_new_get")
async def customer_new_get(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    s = store.get_settings(cid)
    return templates.TemplateResponse("commercial/customer_form.html", _ctx(
        request, active_page="customers", customer={"code": store.next_customer_code(cid), "credit_terms_days": s["default_credit_days"], "is_active": True},
        is_edit=False, portal_users=store.portal_customers(cid), **_lookups(cid)))


@router.post("/customers/new", name="commercial_customer_new_post")
async def customer_new_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    data["created_by"] = _actor(request)
    c = store.create_customer(cid, data)
    if c:
        flash(request, f"Customer {c['name']} created", "success")
        return RedirectResponse(f"{P}/customers/{c['id']}", status_code=303)
    flash(request, "Failed to create customer — name is required and the code must be unique", "error")
    return RedirectResponse(f"{P}/customers/new", status_code=303)


@router.get("/customers/{customer_id}", name="commercial_customer_detail")
async def customer_detail(customer_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    c = store.get_customer(cid, customer_id)
    if not c:
        flash(request, "Customer not found", "error")
        return RedirectResponse(f"{P}/customers", status_code=303)
    return templates.TemplateResponse("commercial/customer_detail.html", _ctx(
        request, active_page="customers", customer=c, history=store.customer_history(cid, customer_id),
        credit=L.credit_check(c["credit_limit"], c["balance"], 0, allow_over=True),
        events=store.events_for(cid, "customer", customer_id)))


@router.get("/customers/{customer_id}/edit", name="commercial_customer_edit_get")
async def customer_edit_get(customer_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    c = store.get_customer(cid, customer_id)
    if not c:
        flash(request, "Customer not found", "error")
        return RedirectResponse(f"{P}/customers", status_code=303)
    return templates.TemplateResponse("commercial/customer_form.html", _ctx(
        request, active_page="customers", customer=c, is_edit=True, portal_users=store.portal_customers(cid), **_lookups(cid)))


@router.post("/customers/{customer_id}/edit", name="commercial_customer_edit_post")
async def customer_edit_post(customer_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    if D(data.get("credit_limit")) != D((store.get_customer(cid, customer_id) or {}).get("credit_limit")) and _priv(request) not in ("manager", "admin", "super_admin"):
        flash(request, "Only a manager can change a credit limit", "error")
        return RedirectResponse(f"{P}/customers/{customer_id}/edit", status_code=303)
    ok = store.update_customer(cid, customer_id, data)
    flash(request, "Customer updated" if ok else "Failed to update customer", "success" if ok else "error")
    return RedirectResponse(f"{P}/customers/{customer_id}", status_code=303)


# ── products ─────────────────────────────────────────────────────
@router.get("/products", name="commercial_products")
async def products(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    q = (request.query_params.get("q") or "").strip()
    return templates.TemplateResponse("commercial/products.html", _ctx(
        request, active_page="products", products=store.list_products(cid, q or None), q=q))


@router.get("/products/new", name="commercial_product_new_get")
async def product_new_get(request: Request, user=Depends(login_required)):
    return templates.TemplateResponse("commercial/product_form.html", _ctx(
        request, active_page="products", product={"unit": "m", "currency": "ETB", "vat_rate": "0.15", "is_active": True}, is_edit=False))


@router.post("/products/new", name="commercial_product_new_post")
async def product_new_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    p = store.create_product(cid, await _form(request))
    flash(request, f"Product {p['item_code']} created" if p else "Failed to create product — item code is required and must be unique",
          "success" if p else "error")
    return RedirectResponse(f"{P}/products", status_code=303)


@router.post("/products/import", name="commercial_products_import")
async def products_import(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    n, msg = store.import_from_manufacturing(cid)
    flash(request, msg, "success" if n else "info")
    return RedirectResponse(f"{P}/products", status_code=303)


@router.get("/products/{product_id}/edit", name="commercial_product_edit_get")
async def product_edit_get(product_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    p = store.get_product(cid, product_id)
    if not p:
        flash(request, "Product not found", "error")
        return RedirectResponse(f"{P}/products", status_code=303)
    return templates.TemplateResponse("commercial/product_form.html", _ctx(request, active_page="products", product=p, is_edit=True))


@router.post("/products/{product_id}/edit", name="commercial_product_edit_post")
async def product_edit_post(product_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    ok = store.update_product(cid, product_id, await _form(request))
    flash(request, "Product updated" if ok else "Failed to update product", "success" if ok else "error")
    return RedirectResponse(f"{P}/products", status_code=303)


# ── pricing (price lists + discounts) ────────────────────────────
@router.get("/pricing", name="commercial_pricing")
async def pricing(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    pl_id = request.query_params.get("price_list")
    price_lists = store.list_price_lists(cid)
    selected = store.get_price_list(cid, pl_id) if pl_id else (store.get_price_list(cid, price_lists[0]["id"]) if price_lists else None)
    return templates.TemplateResponse("commercial/pricing.html", _ctx(
        request, active_page="pricing", price_lists=price_lists, selected=selected, discounts=store.list_discounts(cid),
        products=store.list_products(cid, active_only=True), settings=store.get_settings(cid)))


@router.post("/pricing/price-lists/new", name="commercial_price_list_new")
async def price_list_new(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    pl = store.create_price_list(cid, await _form(request))
    flash(request, "Price list created" if pl else "Failed to create price list", "success" if pl else "error")
    return RedirectResponse(f"{P}/pricing" + (f"?price_list={pl['id']}" if pl else ""), status_code=303)


@router.post("/pricing/price-lists/{pl_id}/items", name="commercial_price_list_item_add")
async def price_list_item_add(pl_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    ok = store.add_price_list_item(cid, pl_id, await _form(request))
    flash(request, "Price added" if ok else "Failed to add price", "success" if ok else "error")
    return RedirectResponse(f"{P}/pricing?price_list={pl_id}", status_code=303)


@router.post("/pricing/items/{item_id}/delete", name="commercial_price_list_item_delete")
async def price_list_item_delete(item_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    store.delete_price_list_item(cid, item_id)
    flash(request, "Price removed", "success")
    return RedirectResponse(f"{P}/pricing?price_list={data.get('price_list_id', '')}", status_code=303)


@router.post("/pricing/discounts/new", name="commercial_discount_new")
async def discount_new(request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    d = store.create_discount(cid, await _form(request))
    flash(request, "Discount rule created" if d else "Failed to create discount", "success" if d else "error")
    return RedirectResponse(f"{P}/pricing#discounts", status_code=303)


@router.post("/pricing/discounts/{did}/toggle", name="commercial_discount_toggle")
async def discount_toggle(did: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    store.toggle_discount(cid, did)
    return RedirectResponse(f"{P}/pricing#discounts", status_code=303)


# ── territories & sales reps ─────────────────────────────────────
@router.get("/territories", name="commercial_territories")
async def territories(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    return templates.TemplateResponse("commercial/territories.html", _ctx(
        request, active_page="territories", territories=store.list_territories(cid), reps=store.list_reps(cid)))


@router.post("/territories/new", name="commercial_territory_new")
async def territory_new(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    t = store.create_territory(cid, await _form(request))
    flash(request, "Territory created" if t else "Territory name is required", "success" if t else "error")
    return RedirectResponse(f"{P}/territories", status_code=303)


@router.post("/territories/{tid}/delete", name="commercial_territory_delete")
async def territory_delete(tid: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    store.delete_territory(cid, tid)
    flash(request, "Territory removed", "success")
    return RedirectResponse(f"{P}/territories", status_code=303)


@router.post("/territories/reps/new", name="commercial_rep_new")
async def rep_new(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    r = store.create_rep(cid, await _form(request))
    flash(request, "Sales representative added" if r else "Name is required", "success" if r else "error")
    return RedirectResponse(f"{P}/territories", status_code=303)


@router.post("/territories/reps/{rep_id}/edit", name="commercial_rep_edit")
async def rep_edit(rep_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    ok = store.update_rep(cid, rep_id, await _form(request))
    flash(request, "Sales representative updated" if ok else "Failed to update", "success" if ok else "error")
    return RedirectResponse(f"{P}/territories", status_code=303)


@router.post("/territories/reps/{rep_id}/toggle", name="commercial_rep_toggle")
async def rep_toggle(rep_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    store.toggle_rep(cid, rep_id)
    return RedirectResponse(f"{P}/territories", status_code=303)


# ── proforma invoices ────────────────────────────────────────────
@router.get("/proformas", name="commercial_proformas")
async def proformas(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    f = _filters(request)
    return templates.TemplateResponse("commercial/proformas.html", _ctx(
        request, active_page="proformas", filters=f,
        proformas=store.list_proformas(cid, f["status"] or None, f["customer_id"] or None, f["q"] or None), statuses=L.PROFORMA_STATUSES))


@router.get("/proformas/new", name="commercial_proforma_new_get")
async def proforma_new_get(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    return templates.TemplateResponse("commercial/proforma_form.html", _ctx(
        request, active_page="proformas", proforma={"proforma_date": date.today().isoformat(), "prepared_by": _actor(request)},
        lines=[], settings=store.get_settings(cid), **_lookups(cid)))


@router.post("/proformas/new", name="commercial_proforma_new_post")
async def proforma_new_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data, lines = await _form_with_lines(request)
    p = store.create_proforma(cid, data, lines, _actor(request))
    if p:
        flash(request, f"Proforma {p['proforma_no']} created", "success")
        return RedirectResponse(f"{P}/proformas/{p['id']}", status_code=303)
    flash(request, "Failed to create proforma — a customer and at least one line with quantity are required", "error")
    return RedirectResponse(f"{P}/proformas/new", status_code=303)


@router.get("/proformas/{pid}", name="commercial_proforma_detail")
async def proforma_detail(pid: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    p = store.get_proforma(cid, pid)
    if not p:
        flash(request, "Proforma not found", "error")
        return RedirectResponse(f"{P}/proformas", status_code=303)
    return templates.TemplateResponse("commercial/proforma_detail.html", _ctx(
        request, active_page="proformas", proforma=p, events=store.events_for(cid, "proforma", pid), statuses=L.PROFORMA_STATUSES))


@router.get("/proformas/{pid}/print", name="commercial_proforma_print")
@router.get("/proformas/{pid}/pdf", name="commercial_proforma_pdf")
@router.get("/proformas/{pid}/xlsx", name="commercial_proforma_xlsx")
async def proforma_print(pid: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    p = store.get_proforma(cid, pid)
    if not p:
        return RedirectResponse(f"{P}/proformas", status_code=303)
    return _doc_response(request, request.url.path.rsplit("/", 1)[-1], _doc_dict("proforma", p, cid), cid)


@router.post("/proformas/{pid}/status", name="commercial_proforma_status")
async def proforma_status(pid: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    ok = store.set_proforma_status(cid, pid, data.get("status", ""), _actor(request))
    flash(request, f"Proforma marked {data.get('status')}" if ok else "Invalid status", "success" if ok else "error")
    return RedirectResponse(f"{P}/proformas/{pid}", status_code=303)


@router.post("/proformas/{pid}/convert", name="commercial_proforma_convert")
async def proforma_convert(pid: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    so = store.convert_proforma(cid, pid, _actor(request))
    if so:
        flash(request, f"Sales order {so['so_no']} created from proforma — review and submit for approval", "success")
        return RedirectResponse(f"{P}/orders/{so['id']}", status_code=303)
    flash(request, "Could not convert proforma (already converted?)", "error")
    return RedirectResponse(f"{P}/proformas/{pid}", status_code=303)


# ── sales orders ─────────────────────────────────────────────────
@router.get("/orders", name="commercial_orders")
async def orders(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    f = _filters(request)
    return templates.TemplateResponse("commercial/sales_orders.html", _ctx(
        request, active_page="orders", filters=f,
        orders=store.list_sales_orders(cid, f["status"] or None, f["customer_id"] or None, f["q"] or None, f["date_from"] or None,
                                       f["date_to"] or None, f["sales_rep_id"] or None, f["territory_id"] or None), **_lookups(cid)))


@router.get("/orders/new", name="commercial_order_new_get")
async def order_new_get(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    return templates.TemplateResponse("commercial/sales_order_form.html", _ctx(
        request, active_page="orders", order={"order_date": date.today().isoformat(), "prepared_by": _actor(request), "copy_finance": True,
                                              "copy_sales": True, "copy_first": True, "customer_id": request.query_params.get("customer_id") or ""},
        lines=[], is_edit=False, settings=store.get_settings(cid), **_lookups(cid)))


@router.post("/orders/new", name="commercial_order_new_post")
async def order_new_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data, lines = await _form_with_lines(request)
    so = store.create_sales_order(cid, data, lines, _actor(request))
    if so:
        flash(request, f"Sales order {so['so_no']} created as draft", "success")
        return RedirectResponse(f"{P}/orders/{so['id']}", status_code=303)
    flash(request, "Failed to create order — a customer and at least one line with quantity are required", "error")
    return RedirectResponse(f"{P}/orders/new", status_code=303)


@router.get("/orders/approvals", name="commercial_order_approvals")
async def order_approvals(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    return templates.TemplateResponse("commercial/approvals.html", _ctx(
        request, active_page="approvals", pending=store.pending_approvals_for(cid, _actor(request), _priv(request)),
        all_pending=store.list_sales_orders(cid, status="pending_approval"), approvers=store.get_settings(cid)["so_approvers"]))


@router.get("/orders/{so_id}", name="commercial_order_detail")
async def order_detail(so_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    so = store.get_sales_order(cid, so_id)
    if not so:
        flash(request, "Sales order not found", "error")
        return RedirectResponse(f"{P}/orders", status_code=303)
    credit = store.credit_status(cid, so.get("customer_id"), so["grand_total"] if so["status"] in ("draft", "rejected") else 0, so["credit_sale"])
    can_step = None
    if so["status"] == "pending_approval" and so["next_seq"]:
        step = next((a for a in so["approvals"] if a["seq"] == so["next_seq"]), None)
        if step and L.can_approve(step, _actor(request), _priv(request), requested_by=so.get("created_by") or ""):
            can_step = step
    return templates.TemplateResponse("commercial/sales_order_detail.html", _ctx(
        request, active_page="orders", order=so, credit=credit, availability=store.availability(cid, so), can_step=can_step,
        events=store.events_for(cid, "sales_order", so_id), mfg_status=store.manufacturing_status(cid, so["so_no"]),
        is_manager=_priv(request) in ("manager", "admin", "super_admin"), settings=store.get_settings(cid)))


@router.get("/orders/{so_id}/edit", name="commercial_order_edit_get")
async def order_edit_get(so_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    so = store.get_sales_order(cid, so_id)
    if not so or so["status"] not in ("draft", "rejected"):
        flash(request, "Only draft or rejected orders can be edited", "error")
        return RedirectResponse(f"{P}/orders/{so_id}", status_code=303)
    return templates.TemplateResponse("commercial/sales_order_form.html", _ctx(
        request, active_page="orders", order=so, lines=so["lines"], is_edit=True, settings=store.get_settings(cid), **_lookups(cid)))


@router.post("/orders/{so_id}/edit", name="commercial_order_edit_post")
async def order_edit_post(so_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data, lines = await _form_with_lines(request)
    ok = store.update_sales_order(cid, so_id, data, lines, _actor(request))
    flash(request, "Order updated" if ok else "Failed to update order", "success" if ok else "error")
    return RedirectResponse(f"{P}/orders/{so_id}", status_code=303)


@router.post("/orders/{so_id}/submit", name="commercial_order_submit")
async def order_submit(so_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    override = L.as_bool(data.get("credit_override")) and _priv(request) in ("manager", "admin", "super_admin")
    ok, msg = store.submit_for_approval(cid, so_id, _actor(request), credit_override=override)
    if ok:
        so = store.get_sales_order(cid, so_id)
        try:
            from approval_hooks import request_approval
            rid = request_approval(cid, "sales_order", so_id, f"Sales order {so['so_no']} — {so['customer_name']}", so["grand_total"],
                                   _actor(request), {"so_no": so["so_no"], "customer": so["customer_name"], "credit_sale": so["credit_sale"]})
            if rid:
                store.set_approval_request(so_id, rid)
                msg += " — also routed through the approval workflow engine"
        except Exception as e:
            logger.debug("approval engine unavailable: %s", e)
    flash(request, msg, "success" if ok else "error")
    return RedirectResponse(f"{P}/orders/{so_id}", status_code=303)


@router.post("/orders/{so_id}/approve", name="commercial_order_approve")
async def order_approve(so_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    data = await _form(request)
    ok, msg = store.decide_approval(cid, so_id, int(data.get("seq") or 0), _actor(request), data.get("decision", ""),
                                    data.get("comment", ""), _priv(request))
    flash(request, msg, "success" if ok else "error")
    return RedirectResponse(f"{P}/orders/{so_id}", status_code=303)


@router.post("/orders/{so_id}/status", name="commercial_order_status")
async def order_status(so_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    status = data.get("status", "")
    if status in ("closed", "cancelled") and _priv(request) not in ("manager", "admin", "super_admin"):
        flash(request, "Only a manager can close or cancel an order", "error")
        return RedirectResponse(f"{P}/orders/{so_id}", status_code=303)
    ok = store.set_so_status(cid, so_id, status, _actor(request), data.get("note", ""))
    flash(request, f"Order marked {status.replace('_', ' ')}" if ok else f"Cannot move this order to {status}", "success" if ok else "error")
    return RedirectResponse(f"{P}/orders/{so_id}", status_code=303)


@router.post("/orders/{so_id}/reserve", name="commercial_order_reserve")
async def order_reserve(so_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    qtys = {k[len("reserve_"):]: v for k, v in data.items() if k.startswith("reserve_")}
    n = store.reserve(cid, so_id, qtys, _actor(request))
    flash(request, f"Reservation updated on {n} line(s)" if n else "Nothing reserved", "success" if n else "info")
    return RedirectResponse(f"{P}/orders/{so_id}", status_code=303)


@router.post("/orders/{so_id}/mo", name="commercial_order_mo")
async def order_mo(so_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    line_id = data.get("line_id") or None
    created, linked, msg = store.create_mo_requests(cid, so_id, [line_id] if line_id else None, data, _actor(request))
    flash(request, msg, "success" if created else "error")
    return RedirectResponse(f"{P}/orders/{so_id}", status_code=303)


@router.get("/orders/{so_id}/print", name="commercial_order_print")
@router.get("/orders/{so_id}/pdf", name="commercial_order_pdf")
@router.get("/orders/{so_id}/xlsx", name="commercial_order_xlsx")
async def order_print(so_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    so = store.get_sales_order(cid, so_id)
    if not so:
        return RedirectResponse(f"{P}/orders", status_code=303)
    return _doc_response(request, request.url.path.rsplit("/", 1)[-1], _doc_dict("sales_order", so, cid), cid)


@router.get("/orders/{so_id}/deliver", name="commercial_order_deliver_get")
async def order_deliver_get(so_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    so = store.get_sales_order(cid, so_id)
    if not so:
        return RedirectResponse(f"{P}/orders", status_code=303)
    return templates.TemplateResponse("commercial/delivery_form.html", _ctx(
        request, active_page="deliveries", order=so, availability=store.availability(cid, so), actor=_actor(request)))


@router.post("/orders/{so_id}/deliver", name="commercial_order_deliver_post")
async def order_deliver_post(so_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    d = store.create_delivery(cid, so_id, _qty_map(data), data, _actor(request))
    if d:
        flash(request, f"Delivery instruction {d['di_no']} issued to inventory", "success")
        return RedirectResponse(f"{P}/deliveries/{d['id']}", status_code=303)
    flash(request, "Nothing to deliver — enter quantities (order must be approved)", "error")
    return RedirectResponse(f"{P}/orders/{so_id}/deliver", status_code=303)


@router.get("/orders/{so_id}/invoice", name="commercial_order_invoice_get")
async def order_invoice_get(so_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    so = store.get_sales_order(cid, so_id)
    if not so:
        return RedirectResponse(f"{P}/orders", status_code=303)
    cust = store.get_customer(cid, so["customer_id"]) if so.get("customer_id") else None
    return templates.TemplateResponse("commercial/invoice_form.html", _ctx(
        request, active_page="invoices", order=so, customer=cust, invoice={"invoice_date": date.today().isoformat(), "credit_sale": so["credit_sale"]},
        lines=[], is_manager=_priv(request) in ("manager", "admin", "super_admin"), settings=store.get_settings(cid), **_lookups(cid)))


@router.post("/orders/{so_id}/invoice", name="commercial_order_invoice_post")
async def order_invoice_post(so_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    if L.as_bool(data.get("credit_override")) and _priv(request) not in ("manager", "admin", "super_admin"):
        data["credit_override"] = ""
    inv, msg = store.create_invoice(cid, so_id, _qty_map(data), data, actor=_actor(request))
    if inv:
        flash(request, f"Invoice {inv['invoice_no']} issued", "success")
        return RedirectResponse(f"{P}/invoices/{inv['id']}", status_code=303)
    flash(request, msg, "error")
    return RedirectResponse(f"{P}/orders/{so_id}/invoice", status_code=303)


# ── manufacturing order requests ─────────────────────────────────
@router.get("/mo-requests", name="commercial_mo_requests")
async def mo_requests(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    status = request.query_params.get("status") or None
    return templates.TemplateResponse("commercial/mo_requests.html", _ctx(
        request, active_page="mo_requests", requests=store.list_mo_requests(cid, status), status_filter=status or ""))


@router.get("/mo-requests/{rid}", name="commercial_mo_request_detail")
async def mo_request_detail(rid: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    r = store.get_mo_request(cid, rid)
    if not r:
        flash(request, "Request not found", "error")
        return RedirectResponse(f"{P}/mo-requests", status_code=303)
    return templates.TemplateResponse("commercial/mo_request_detail.html", _ctx(request, active_page="mo_requests", req=r))


@router.post("/mo-requests/{rid}/update", name="commercial_mo_request_update")
async def mo_request_update(rid: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    ok = store.update_mo_request(cid, rid, await _form(request), _actor(request))
    flash(request, "Request updated" if ok else "Failed to update request", "success" if ok else "error")
    return RedirectResponse(f"{P}/mo-requests/{rid}", status_code=303)


@router.get("/mo-requests/{rid}/print", name="commercial_mo_request_print")
@router.get("/mo-requests/{rid}/pdf", name="commercial_mo_request_pdf")
@router.get("/mo-requests/{rid}/xlsx", name="commercial_mo_request_xlsx")
async def mo_request_print(rid: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    r = store.get_mo_request(cid, rid)
    if not r:
        return RedirectResponse(f"{P}/mo-requests", status_code=303)
    return _doc_response(request, request.url.path.rsplit("/", 1)[-1], _doc_dict("mo_request", r, cid), cid)


# ── deliveries & dispatch ────────────────────────────────────────
@router.get("/deliveries", name="commercial_deliveries")
async def deliveries(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    status = request.query_params.get("status") or None
    return templates.TemplateResponse("commercial/deliveries.html", _ctx(
        request, active_page="deliveries", deliveries=store.list_deliveries(cid, status), dispatches=store.list_dispatches(cid)[:100],
        status_filter=status or "", dispatch_statuses=L.DISPATCH_STATUSES))


@router.get("/deliveries/dispatches/{xid}/print", name="commercial_dispatch_print")
@router.get("/deliveries/dispatches/{xid}/pdf", name="commercial_dispatch_pdf")
async def dispatch_print(xid: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    x = store.get_dispatch(cid, xid)
    if not x:
        return RedirectResponse(f"{P}/deliveries", status_code=303)
    d = store.get_delivery(cid, x["di_id"]) or {}
    x["lines"] = d.get("lines", [])
    return _doc_response(request, request.url.path.rsplit("/", 1)[-1], _doc_dict("dispatch", x, cid), cid)


@router.post("/deliveries/dispatches/{xid}/update", name="commercial_dispatch_update")
async def dispatch_update(xid: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    ok, msg = store.update_dispatch(cid, xid, await _form(request), _actor(request))
    flash(request, msg, "success" if ok else "error")
    x = store.get_dispatch(cid, xid)
    return RedirectResponse(f"{P}/deliveries/{x['di_id']}" if x else f"{P}/deliveries", status_code=303)


@router.get("/deliveries/{di_id}", name="commercial_delivery_detail")
async def delivery_detail(di_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    d = store.get_delivery(cid, di_id)
    if not d:
        flash(request, "Delivery instruction not found", "error")
        return RedirectResponse(f"{P}/deliveries", status_code=303)
    return templates.TemplateResponse("commercial/delivery_detail.html", _ctx(
        request, active_page="deliveries", delivery=d, events=store.events_for(cid, "delivery", di_id), dispatch_statuses=L.DISPATCH_STATUSES))


@router.get("/deliveries/{di_id}/print", name="commercial_delivery_print")
@router.get("/deliveries/{di_id}/pdf", name="commercial_delivery_pdf")
@router.get("/deliveries/{di_id}/xlsx", name="commercial_delivery_xlsx")
async def delivery_print(di_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    d = store.get_delivery(cid, di_id)
    if not d:
        return RedirectResponse(f"{P}/deliveries", status_code=303)
    return _doc_response(request, request.url.path.rsplit("/", 1)[-1], _doc_dict("delivery", d, cid), cid)


@router.post("/deliveries/{di_id}/dispatch", name="commercial_dispatch_new")
async def dispatch_new(di_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    x = store.create_dispatch(cid, di_id, await _form(request), _actor(request))
    flash(request, f"Dispatch {x['dispatch_no']} created" if x else "Could not create dispatch", "success" if x else "error")
    return RedirectResponse(f"{P}/deliveries/{di_id}", status_code=303)


# ── invoices, receipts ───────────────────────────────────────────
@router.get("/invoices", name="commercial_invoices")
async def invoices(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    f = _filters(request)
    credit_only = request.query_params.get("credit") == "1"
    overdue_only = request.query_params.get("overdue") == "1"
    rows = store.list_invoices(cid, f["status"] or None, f["customer_id"] or None, f["q"] or None, credit_only, f["date_from"] or None,
                               f["date_to"] or None, f["sales_rep_id"] or None, overdue_only)
    totals = {"total": sum((D(r["grand_total"]) for r in rows), Decimal(0)), "paid": sum((D(r["paid_total"]) for r in rows), Decimal(0)),
              "outstanding": sum((D(r["outstanding"]) for r in rows), Decimal(0))}
    return templates.TemplateResponse("commercial/invoices.html", _ctx(
        request, active_page="invoices", invoices=rows, filters=f, totals=totals, credit_only=credit_only, overdue_only=overdue_only,
        statuses=L.INVOICE_STATUSES, **_lookups(cid)))


@router.get("/invoices/new", name="commercial_invoice_new_get")
async def invoice_new_get(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    return templates.TemplateResponse("commercial/invoice_form.html", _ctx(
        request, active_page="invoices", order=None, customer=None, invoice={"invoice_date": date.today().isoformat()}, lines=[],
        is_manager=_priv(request) in ("manager", "admin", "super_admin"), settings=store.get_settings(cid), **_lookups(cid)))


@router.post("/invoices/new", name="commercial_invoice_new_post")
async def invoice_new_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data, lines = await _form_with_lines(request)
    if L.as_bool(data.get("credit_override")) and _priv(request) not in ("manager", "admin", "super_admin"):
        data["credit_override"] = ""
    inv, msg = store.create_invoice(cid, None, None, data, lines=lines, actor=_actor(request))
    if inv:
        flash(request, f"Invoice {inv['invoice_no']} issued", "success")
        return RedirectResponse(f"{P}/invoices/{inv['id']}", status_code=303)
    flash(request, msg, "error")
    return RedirectResponse(f"{P}/invoices/new", status_code=303)


@router.post("/invoices/link-payments", name="commercial_invoices_link_payments")
async def invoices_link_payments(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    n, msg = store.link_mobile_payments(cid, _actor(request))
    flash(request, msg, "success" if n else "info")
    return RedirectResponse(f"{P}/invoices", status_code=303)


@router.get("/invoices/{invoice_id}", name="commercial_invoice_detail")
async def invoice_detail(invoice_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    inv = store.get_invoice(cid, invoice_id)
    if not inv:
        flash(request, "Invoice not found", "error")
        return RedirectResponse(f"{P}/invoices", status_code=303)
    return templates.TemplateResponse("commercial/invoice_detail.html", _ctx(
        request, active_page="invoices", invoice=inv, events=store.events_for(cid, "invoice", invoice_id),
        is_manager=_priv(request) in ("manager", "admin", "super_admin")))


@router.get("/invoices/{invoice_id}/print", name="commercial_invoice_print")
@router.get("/invoices/{invoice_id}/pdf", name="commercial_invoice_pdf")
@router.get("/invoices/{invoice_id}/xlsx", name="commercial_invoice_xlsx")
async def invoice_print(invoice_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    inv = store.get_invoice(cid, invoice_id)
    if not inv:
        return RedirectResponse(f"{P}/invoices", status_code=303)
    return _doc_response(request, request.url.path.rsplit("/", 1)[-1], _doc_dict("invoice", inv, cid), cid)


@router.post("/invoices/{invoice_id}/receipt", name="commercial_invoice_receipt")
async def invoice_receipt(invoice_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    r = store.add_receipt(cid, invoice_id, await _form(request), _actor(request))
    flash(request, f"Receipt {r['receipt_no']} recorded" if r else "Failed to record payment — amount must be positive", "success" if r else "error")
    return RedirectResponse(f"{P}/invoices/{invoice_id}", status_code=303)


@router.post("/invoices/{invoice_id}/cancel", name="commercial_invoice_cancel")
async def invoice_cancel(invoice_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    data = await _form(request)
    ok = store.cancel_invoice(cid, invoice_id, _actor(request), data.get("reason", ""))
    flash(request, "Invoice cancelled" if ok else "Invoice cannot be cancelled (payments recorded?)", "success" if ok else "error")
    return RedirectResponse(f"{P}/invoices/{invoice_id}", status_code=303)


@router.get("/invoices/{invoice_id}/return", name="commercial_invoice_return_get")
async def invoice_return_get(invoice_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    inv = store.get_invoice(cid, invoice_id)
    if not inv:
        return RedirectResponse(f"{P}/invoices", status_code=303)
    return templates.TemplateResponse("commercial/return_form.html", _ctx(request, active_page="returns", invoice=inv))


@router.post("/invoices/{invoice_id}/return", name="commercial_invoice_return_post")
async def invoice_return_post(invoice_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    r = store.create_return(cid, invoice_id, _qty_map(data), data, _actor(request))
    if r:
        flash(request, f"Return {r['rn_no']} raised — awaiting approval", "success")
        return RedirectResponse(f"{P}/returns/{r['id']}", status_code=303)
    flash(request, "Enter at least one quantity to return", "error")
    return RedirectResponse(f"{P}/invoices/{invoice_id}/return", status_code=303)


@router.get("/receipts", name="commercial_receipts")
async def receipts(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    f = _filters(request)
    method = request.query_params.get("method") or None
    rows = store.list_receipts(cid, f["date_from"] or None, f["date_to"] or None, method)
    return templates.TemplateResponse("commercial/receipts.html", _ctx(
        request, active_page="invoices", receipts=rows, filters=f, method=method or "",
        total=sum((D(r["amount"]) for r in rows), Decimal(0))))


# ── returns / credit notes ───────────────────────────────────────
@router.get("/returns", name="commercial_returns")
async def returns(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    status = request.query_params.get("status") or None
    return templates.TemplateResponse("commercial/returns.html", _ctx(
        request, active_page="returns", returns=store.list_returns(cid, status), status_filter=status or "", statuses=L.RETURN_STATUSES))


@router.get("/returns/{rid}", name="commercial_return_detail")
async def return_detail(rid: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    r = store.get_return(cid, rid)
    if not r:
        flash(request, "Return not found", "error")
        return RedirectResponse(f"{P}/returns", status_code=303)
    return templates.TemplateResponse("commercial/return_detail.html", _ctx(
        request, active_page="returns", ret=r, is_manager=_priv(request) in ("manager", "admin", "super_admin")))


@router.post("/returns/{rid}/decide", name="commercial_return_decide")
async def return_decide(rid: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    data = await _form(request)
    ok, msg = store.decide_return(cid, rid, data.get("decision") == "approve", _actor(request))
    flash(request, msg, "success" if ok else "error")
    return RedirectResponse(f"{P}/returns/{rid}", status_code=303)


@router.get("/returns/{rid}/print", name="commercial_return_print")
@router.get("/returns/{rid}/pdf", name="commercial_return_pdf")
async def return_print(rid: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    r = store.get_return(cid, rid)
    if not r:
        return RedirectResponse(f"{P}/returns", status_code=303)
    return _doc_response(request, request.url.path.rsplit("/", 1)[-1], _doc_dict("credit_note", r, cid), cid)


# ── forecasts & commissions ──────────────────────────────────────
@router.get("/forecasts", name="commercial_forecasts")
async def forecasts(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    periods = store.forecast_periods(cid)
    period = request.query_params.get("period") or (periods[0] if periods else L.next_period(L.period_of(date.today())))
    rows = store.list_forecasts(cid, period)
    return templates.TemplateResponse("commercial/forecasts.html", _ctx(
        request, active_page="forecasts", forecasts=rows, periods=periods, period=period,
        total_qty=sum((D(r["forecast_qty"]) for r in rows), Decimal(0)), total_value=sum((D(r["forecast_value"]) for r in rows), Decimal(0)),
        next_period=L.next_period(L.period_of(date.today())), **_lookups(cid)))


@router.post("/forecasts/refresh", name="commercial_forecasts_refresh")
async def forecasts_refresh(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    period = (data.get("period") or "").strip() or None
    if period and not L.PERIOD_RE.match(period):
        flash(request, "Period must be YYYY-MM", "error")
        return RedirectResponse(f"{P}/forecasts", status_code=303)
    n = store.refresh_forecasts(cid, period, _actor(request))
    flash(request, f"Forecast refreshed — {n} item(s) projected from invoiced history" if n else "No invoiced history yet to forecast from", "success" if n else "info")
    return RedirectResponse(f"{P}/forecasts" + (f"?period={period}" if period else ""), status_code=303)


@router.post("/forecasts/manual", name="commercial_forecast_manual")
async def forecast_manual(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    ok = store.save_manual_forecast(cid, data, _actor(request))
    flash(request, "Manual forecast saved" if ok else "Period must be YYYY-MM", "success" if ok else "error")
    return RedirectResponse(f"{P}/forecasts?period={data.get('period', '')}", status_code=303)


@router.post("/forecasts/{fid}/delete", name="commercial_forecast_delete")
async def forecast_delete(fid: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    store.delete_forecast(cid, fid)
    return RedirectResponse(f"{P}/forecasts", status_code=303)


@router.get("/commissions", name="commercial_commissions")
async def commissions(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    qp = request.query_params
    rows = store.list_commissions(cid, qp.get("period") or None, qp.get("sales_rep_id") or None, qp.get("status") or None)
    by_rep: Dict[str, dict] = {}
    for r in rows:
        k = f"{r['sales_rep_id']}|{r['period']}"
        b = by_rep.setdefault(k, {"sales_rep_id": r["sales_rep_id"], "sales_rep_name": r["sales_rep_name"], "period": r["period"],
                                  "base": Decimal(0), "commission": Decimal(0), "accrued": Decimal(0), "count": 0})
        b["base"] += D(r["base_amount"]); b["commission"] += D(r["commission_amount"]); b["count"] += 1
        if r["status"] == "accrued":
            b["accrued"] += D(r["commission_amount"])
    return templates.TemplateResponse("commercial/commissions.html", _ctx(
        request, active_page="commissions", commissions=rows, summary=list(by_rep.values()),
        filters={"period": qp.get("period") or "", "sales_rep_id": qp.get("sales_rep_id") or "", "status": qp.get("status") or ""},
        is_manager=_priv(request) in ("manager", "admin", "super_admin"), **_lookups(cid)))


@router.post("/commissions/pay", name="commercial_commissions_pay")
async def commissions_pay(request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    data = await _form(request)
    n = store.mark_commissions_paid(cid, data.get("sales_rep_id", ""), data.get("period", ""))
    flash(request, f"{n} commission line(s) marked paid", "success" if n else "info")
    return RedirectResponse(f"{P}/commissions?period={data.get('period', '')}", status_code=303)


# ── marketing ────────────────────────────────────────────────────
@router.get("/marketing", name="commercial_marketing")
async def marketing(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    m = store.marketing_dashboard(cid)
    chart = {"seg_labels": [r["segment"] for r in m["revenue_by_segment"]], "seg_values": [r["revenue"] for r in m["revenue_by_segment"]],
             "acq_labels": [a["period"] for a in m["acquisition"]], "acq_new": [a["new"] for a in m["acquisition"]],
             "acq_ret": [a["returning"] for a in m["acquisition"]]}
    return templates.TemplateResponse("commercial/marketing_dashboard.html", _ctx(
        request, active_page="marketing", m=m, chart_json=json.dumps(chart, default=_json_default), leads=store.list_leads(cid)[:10]))


@router.get("/marketing/campaigns", name="commercial_campaigns")
async def campaigns(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    status = request.query_params.get("status") or None
    return templates.TemplateResponse("commercial/campaigns.html", _ctx(
        request, active_page="campaigns", campaigns=store.list_campaigns(cid, status), status_filter=status or ""))


@router.get("/marketing/campaigns/new", name="commercial_campaign_new_get")
async def campaign_new_get(request: Request, user=Depends(login_required)):
    return templates.TemplateResponse("commercial/campaign_form.html", _ctx(
        request, active_page="campaigns", campaign={"status": "planned"}, is_edit=False))


@router.post("/marketing/campaigns/new", name="commercial_campaign_new_post")
async def campaign_new_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    c = store.create_campaign(cid, await _form(request))
    if c:
        flash(request, "Campaign created", "success")
        return RedirectResponse(f"{P}/marketing/campaigns/{c['id']}", status_code=303)
    flash(request, "Campaign name is required", "error")
    return RedirectResponse(f"{P}/marketing/campaigns/new", status_code=303)


@router.get("/marketing/campaigns/{cid_}", name="commercial_campaign_detail")
async def campaign_detail(cid_: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    c = store.get_campaign(cid, cid_)
    if not c:
        flash(request, "Campaign not found", "error")
        return RedirectResponse(f"{P}/marketing/campaigns", status_code=303)
    return templates.TemplateResponse("commercial/campaign_detail.html", _ctx(request, active_page="campaigns", campaign=c))


@router.get("/marketing/campaigns/{cid_}/edit", name="commercial_campaign_edit_get")
async def campaign_edit_get(cid_: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    c = store.get_campaign(cid, cid_)
    if not c:
        return RedirectResponse(f"{P}/marketing/campaigns", status_code=303)
    return templates.TemplateResponse("commercial/campaign_form.html", _ctx(request, active_page="campaigns", campaign=c, is_edit=True))


@router.post("/marketing/campaigns/{cid_}/edit", name="commercial_campaign_edit_post")
async def campaign_edit_post(cid_: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    ok = store.update_campaign(cid, cid_, await _form(request))
    flash(request, "Campaign updated" if ok else "Failed to update campaign", "success" if ok else "error")
    return RedirectResponse(f"{P}/marketing/campaigns/{cid_}", status_code=303)


@router.post("/marketing/events/new", name="commercial_event_new")
async def event_new(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    e = store.add_marketing_event(cid, data)
    flash(request, "Event recorded" if e else "Event name is required", "success" if e else "error")
    back = f"{P}/marketing/campaigns/{data['campaign_id']}" if data.get("campaign_id") else f"{P}/marketing/events"
    return RedirectResponse(back, status_code=303)


@router.get("/marketing/events", name="commercial_events")
async def events(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    return templates.TemplateResponse("commercial/events.html", _ctx(
        request, active_page="events", events=store.list_marketing_events(cid), campaigns=store.list_campaigns(cid)))


@router.post("/marketing/events/{eid}/delete", name="commercial_event_delete")
async def event_delete(eid: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    store.delete_marketing_event(cid, eid)
    return RedirectResponse(f"{P}/marketing/events", status_code=303)


@router.get("/marketing/leads", name="commercial_leads")
async def leads(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    qp = request.query_params
    return templates.TemplateResponse("commercial/leads.html", _ctx(
        request, active_page="leads", leads=store.list_leads(cid, qp.get("status") or None, qp.get("campaign_id") or None, qp.get("q") or None),
        campaigns=store.list_campaigns(cid), status_filter=qp.get("status") or "", q=qp.get("q") or ""))


@router.post("/marketing/leads/new", name="commercial_lead_new")
async def lead_new(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    lead = store.create_lead(cid, data, _actor(request))
    if lead and data.get("campaign_id"):
        pass
    flash(request, "Lead added" if lead else "Lead name is required", "success" if lead else "error")
    return RedirectResponse(f"{P}/marketing/leads", status_code=303)


@router.post("/marketing/leads/{lid}/status", name="commercial_lead_status")
async def lead_status(lid: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    ok = store.set_lead_status(cid, lid, data.get("status", ""))
    flash(request, "Lead updated" if ok else "Invalid status", "success" if ok else "error")
    return RedirectResponse(f"{P}/marketing/leads", status_code=303)


@router.post("/marketing/leads/{lid}/convert", name="commercial_lead_convert")
async def lead_convert(lid: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    c = store.convert_lead(cid, lid, _actor(request))
    if c:
        flash(request, f"Lead converted — customer {c['name']} created", "success")
        return RedirectResponse(f"{P}/customers/{c['id']}", status_code=303)
    flash(request, "Could not convert lead", "error")
    return RedirectResponse(f"{P}/marketing/leads", status_code=303)


@router.get("/marketing/competitors", name="commercial_competitors")
async def competitors(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    return templates.TemplateResponse("commercial/competitors.html", _ctx(
        request, active_page="competitors", competitors=store.list_competitors(cid), products=store.list_products(cid, active_only=True)))


@router.post("/marketing/competitors/new", name="commercial_competitor_new")
async def competitor_new(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    c = store.add_competitor(cid, await _form(request), _actor(request))
    flash(request, "Observation recorded" if c else "Competitor name is required", "success" if c else "error")
    return RedirectResponse(f"{P}/marketing/competitors", status_code=303)


@router.post("/marketing/competitors/{cid_}/delete", name="commercial_competitor_delete")
async def competitor_delete(cid_: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    store.delete_competitor(cid, cid_)
    return RedirectResponse(f"{P}/marketing/competitors", status_code=303)


# ── reports ──────────────────────────────────────────────────────
@router.get("/reports", name="commercial_reports")
async def reports(request: Request, user=Depends(login_required)):
    return templates.TemplateResponse("commercial/reports_index.html", _ctx(
        request, active_page="reports", reports=[(k, v[0]) for k, v in store.REPORTS.items()]))


@router.get("/reports/{key}", name="commercial_report")
async def report(key: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    f = _filters(request)
    res = store.run_report(cid, key, f)
    if not res:
        flash(request, "Unknown report", "error")
        return RedirectResponse(f"{P}/reports", status_code=303)
    return templates.TemplateResponse("commercial/report.html", _ctx(
        request, active_page="reports", report=res, filters=f, periods=store.forecast_periods(cid) if key == "forecast" else [], **_lookups(cid)))


@router.get("/reports/{key}/export", name="commercial_report_export")
async def report_export(key: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    f = _filters(request)
    res = store.run_report(cid, key, f)
    if not res:
        return RedirectResponse(f"{P}/reports", status_code=303)
    meta = [(k.replace("_", " ").title(), v) for k, v in f.items() if v]
    path = _xlsx_file(res["title"], res["columns"], res["rows"], res.get("totals"), meta, res.get("extra"))
    return FileResponse(path, filename=f"commercial_{key}_{date.today().isoformat()}.xlsx", media_type=_XLSX)
