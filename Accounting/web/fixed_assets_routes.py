"""
Fixed Assets & Depreciation Routes — /assets
"""
import json
import logging
import os
import tempfile
from datetime import date
from decimal import Decimal

from fastapi import APIRouter, Depends, Request
from fastapi.responses import FileResponse, RedirectResponse

from deps import current_company, flash, login_required, template_context
from fixed_assets_data_store import (
    DISPOSAL_METHODS, METHOD_LABELS, METHODS, PERIOD_RE, STATUSES, D,
    fixed_asset_store, gain_loss, previous_period, q2,
)
from template_engine import templates

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/assets", tags=["fixed_assets"])

_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _ctx(request: Request, **extra) -> dict:
    ctx = template_context(request)
    ctx.update(methods=METHODS, method_labels=METHOD_LABELS, statuses=STATUSES,
               disposal_methods=DISPOSAL_METHODS, active_page=extra.pop("active_page", ""))
    ctx.update(extra)
    return ctx


def _json_default(o):
    if isinstance(o, Decimal):
        return float(o)
    if isinstance(o, date):
        return o.isoformat()
    return str(o)


async def _form(request: Request) -> dict:
    form = await request.form()
    return {k: v for k, v in form.items()}


# ── Dashboard ─────────────────────────────────────────────────────
@router.get("/", name="fixed_assets_dashboard")
async def dashboard(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    seeded = fixed_asset_store.seed_default_categories(cid)
    if seeded:
        flash(request, f"Seeded {seeded} default asset categories (Ethiopian income tax classes)", "success")
    stats = fixed_asset_store.get_dashboard(cid)
    chart = {
        "category_labels": [c["name"] for c in stats["by_category"]],
        "category_cost": [c["cost"] for c in stats["by_category"]],
        "category_accumulated": [c["accumulated"] for c in stats["by_category"]],
        "category_book_value": [c["book_value"] for c in stats["by_category"]],
        "monthly_labels": [m["period"] for m in stats["monthly"]],
        "monthly_totals": [m["total"] for m in stats["monthly"]],
        "status_labels": list(stats["by_status"].keys()),
        "status_counts": list(stats["by_status"].values()),
    }
    return templates.TemplateResponse("assets/dashboard.html", _ctx(
        request, active_page="dashboard", stats=stats,
        chart_json=json.dumps(chart, default=_json_default), today=date.today()))


# ── Static paths MUST be registered before /{asset_id} ────────────
@router.get("/list", name="fixed_assets_list")
async def asset_list(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    qp = request.query_params
    category = qp.get("category") or None
    status = qp.get("status") or None
    search = (qp.get("q") or "").strip() or None
    assets = fixed_asset_store.get_assets(cid, category, status, search)
    totals = {"cost": sum((D(a["cost"]) for a in assets), Decimal(0)),
              "accumulated": sum((D(a["accumulated"]) for a in assets), Decimal(0))}
    totals["book_value"] = totals["cost"] - totals["accumulated"]
    return templates.TemplateResponse("assets/list.html", _ctx(
        request, active_page="list", assets=assets, totals=totals,
        categories=fixed_asset_store.get_categories(cid),
        category_filter=category or "", status_filter=status or "", search=search or ""))


@router.get("/new", name="fixed_assets_new_get")
async def new_asset_get(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    fixed_asset_store.seed_default_categories(cid)
    asset = {"asset_tag": fixed_asset_store.next_asset_tag(cid), "currency": "ETB",
             "method": "straight_line", "useful_life_months": 60,
             "purchase_date": date.today().isoformat(), "in_service_date": date.today().isoformat()}
    return templates.TemplateResponse("assets/form.html", _ctx(
        request, active_page="new", asset=asset, is_edit=False,
        categories=fixed_asset_store.get_categories(cid)))


@router.post("/new", name="fixed_assets_new_post")
async def new_asset_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    a = fixed_asset_store.create_asset(cid, data)
    if a:
        flash(request, f"Asset {a['asset_tag']} created", "success")
        return RedirectResponse(f"/assets/{a['id']}", status_code=303)
    flash(request, "Failed to create asset — tag and name are required and the tag must be unique", "error")
    return RedirectResponse("/assets/new", status_code=303)


# ── Categories ────────────────────────────────────────────────────
@router.get("/categories", name="fixed_assets_categories")
async def categories(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    fixed_asset_store.seed_default_categories(cid)
    return templates.TemplateResponse("assets/categories.html", _ctx(
        request, active_page="categories", categories=fixed_asset_store.get_categories(cid),
        gl_accounts=fixed_asset_store.get_gl_accounts(cid)))


@router.post("/categories/new", name="fixed_assets_category_new")
async def category_new(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    if not (data.get("name") or "").strip():
        flash(request, "Category name is required", "error")
    elif fixed_asset_store.create_category(cid, data):
        flash(request, "Category created", "success")
    else:
        flash(request, "Failed to create category", "error")
    return RedirectResponse("/assets/categories", status_code=303)


@router.post("/categories/{category_id}/edit", name="fixed_assets_category_edit")
async def category_edit(category_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    if fixed_asset_store.update_category(category_id, cid, data):
        flash(request, "Category updated", "success")
    else:
        flash(request, "Failed to update category", "error")
    return RedirectResponse("/assets/categories", status_code=303)


@router.post("/categories/{category_id}/delete", name="fixed_assets_category_delete")
async def category_delete(category_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    ok, err = fixed_asset_store.delete_category(category_id, cid)
    flash(request, "Category deleted" if ok else (err or "Failed to delete category"),
          "success" if ok else "error")
    return RedirectResponse("/assets/categories", status_code=303)


# ── Depreciation run ──────────────────────────────────────────────
def _units_from_form(data: dict) -> dict:
    return {k[len("units_"):]: v for k, v in data.items() if k.startswith("units_") and v not in ("", None)}


@router.get("/depreciation", name="fixed_assets_depreciation")
async def depreciation_run(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    period = (request.query_params.get("period") or previous_period()).strip()
    preview = None
    if request.query_params.get("preview") and PERIOD_RE.match(period):
        preview = fixed_asset_store.preview_period(cid, period)
    elif request.query_params.get("preview"):
        flash(request, "Period must be in YYYY-MM format", "error")
    return templates.TemplateResponse("assets/depreciation_run.html", _ctx(
        request, active_page="depreciation", period=period, preview=preview,
        preview_total=sum((p["amount"] for p in (preview or []) if not p["existing"] and not p["skip_reason"]),
                          Decimal(0)),
        summaries=fixed_asset_store.get_period_summaries(cid)))


@router.post("/depreciation/preview", name="fixed_assets_depreciation_preview")
async def depreciation_preview(request: Request, user=Depends(login_required)):
    data = await _form(request)
    period = (data.get("period") or "").strip()
    return RedirectResponse(f"/assets/depreciation?period={period}&preview=1", status_code=303)


@router.post("/depreciation/run", name="fixed_assets_depreciation_run")
async def depreciation_confirm(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    period = (data.get("period") or "").strip()
    if not PERIOD_RE.match(period):
        flash(request, "Period must be in YYYY-MM format", "error")
        return RedirectResponse("/assets/depreciation", status_code=303)
    res = fixed_asset_store.run_period(cid, period, _units_from_form(data))
    if res.get("error"):
        flash(request, f"Depreciation run failed: {res['error']}", "error")
    else:
        msg = (f"Depreciation for {period}: {res['inserted']} asset(s) recorded, "
               f"total {q2(res['total']):,} — {res['skipped']} skipped")
        if res["fully_depreciated"]:
            msg += f"; {res['fully_depreciated']} now fully depreciated"
        flash(request, msg, "success")
    return RedirectResponse(f"/assets/depreciation?period={period}&preview=1", status_code=303)


@router.post("/depreciation/post", name="fixed_assets_depreciation_post")
async def depreciation_post_gl(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    period = (data.get("period") or "").strip()
    if not PERIOD_RE.match(period):
        flash(request, "Invalid period", "error")
        return RedirectResponse("/assets/depreciation", status_code=303)
    res = fixed_asset_store.post_period_to_gl(cid, period, request.session.get("username", ""))
    if res.get("error"):
        flash(request, f"Post to GL ({period}): {res['error']}", "error")
    else:
        msg = f"Posted {res['posted']} depreciation line(s) for {period} to journal {res['entry_id'][:8]}…"
        if res["skipped"]:
            msg += f" ({res['skipped']} skipped — category has no GL accounts)"
        flash(request, msg, "success")
    return RedirectResponse(f"/assets/depreciation?period={period}", status_code=303)


# ── Asset register report ─────────────────────────────────────────
def _register_rows(cid: str, as_of: str):
    rows = fixed_asset_store.get_register(cid, as_of)
    totals = {"cost": Decimal(0), "accumulated": Decimal(0), "book_value": Decimal(0), "period_charge": Decimal(0)}
    for r in rows:
        for k in totals:
            totals[k] += D(r.get(k))
    return rows, totals


@router.get("/register", name="fixed_assets_register")
async def register(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    as_of = (request.query_params.get("as_of") or "").strip()
    rows, totals = _register_rows(cid, as_of)
    return templates.TemplateResponse("assets/register_report.html", _ctx(
        request, active_page="register", rows=rows, totals=totals, as_of=as_of,
        generated=date.today()))


@router.get("/register/export", name="fixed_assets_register_export")
async def register_export(request: Request, user=Depends(login_required)):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    cid = current_company(request)
    as_of = (request.query_params.get("as_of") or "").strip()
    rows, totals = _register_rows(cid, as_of)
    wb = Workbook()
    ws = wb.active
    ws.title = "Asset Register"
    ws.append([f"Fixed Asset Register — company {cid}" + (f" — as of {as_of}" if as_of else "")])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    headers = ["Asset tag", "Name", "Category", "Serial", "Location", "Custodian", "Purchase date",
               "In service", "Method", "Life (months)", "Status", "Currency", "Cost", "Salvage",
               "Accumulated depreciation", "Net book value", "Period charge", "Last period"]
    ws.append(headers)
    for c in ws[3]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r["asset_tag"], r["name"], r.get("category_name") or "", r.get("serial_number") or "",
                   r.get("location") or "", r.get("custodian") or "",
                   r["purchase_date"].isoformat() if r.get("purchase_date") else "",
                   r["in_service_date"].isoformat() if r.get("in_service_date") else "",
                   METHOD_LABELS.get(r["method"], r["method"]), r["useful_life_months"], r["status"],
                   r["currency"], float(D(r["cost"])), float(D(r["salvage_value"])),
                   float(D(r["accumulated"])), float(D(r["book_value"])), float(D(r["period_charge"])),
                   r.get("last_period") or ""])
    ws.append([])
    ws.append(["TOTAL", "", "", "", "", "", "", "", "", "", "", "", float(totals["cost"]), "",
               float(totals["accumulated"]), float(totals["book_value"]), float(totals["period_charge"]), ""])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for col, width in zip("ABCDEFGHIJKLMNOPQR", (12, 30, 24, 16, 16, 16, 12, 12, 18, 8, 14, 8, 14, 12, 18, 16, 14, 10)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=4, min_col=13, max_col=17):
        for c in row:
            c.number_format = "#,##0.00"
    fd, path = tempfile.mkstemp(suffix=".xlsx"); os.close(fd)
    wb.save(path)
    name = f"asset_register_{as_of or date.today().isoformat()}.xlsx"
    return FileResponse(path, filename=name, media_type=_XLSX)


# ── Asset detail / edit / dispose / maintenance ───────────────────
@router.get("/{asset_id}", name="fixed_assets_detail")
async def asset_detail(asset_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    a = fixed_asset_store.get_asset(asset_id, cid)
    if not a:
        flash(request, "Asset not found", "error")
        return RedirectResponse("/assets/list", status_code=303)
    return templates.TemplateResponse("assets/detail.html", _ctx(
        request, active_page="list", asset=a, schedule=fixed_asset_store.full_schedule(a),
        maintenance=fixed_asset_store.get_maintenance(asset_id),
        disposal=fixed_asset_store.get_disposal(asset_id), today=date.today()))


@router.get("/{asset_id}/edit", name="fixed_assets_edit_get")
async def asset_edit_get(asset_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    a = fixed_asset_store.get_asset(asset_id, cid)
    if not a:
        flash(request, "Asset not found", "error")
        return RedirectResponse("/assets/list", status_code=303)
    return templates.TemplateResponse("assets/form.html", _ctx(
        request, active_page="list", asset=a, is_edit=True,
        categories=fixed_asset_store.get_categories(cid)))


@router.post("/{asset_id}/edit", name="fixed_assets_edit_post")
async def asset_edit_post(asset_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    if fixed_asset_store.update_asset(asset_id, cid, data):
        flash(request, "Asset updated", "success")
        return RedirectResponse(f"/assets/{asset_id}", status_code=303)
    flash(request, "Failed to update asset", "error")
    return RedirectResponse(f"/assets/{asset_id}/edit", status_code=303)


@router.post("/{asset_id}/status", name="fixed_assets_status")
async def asset_status(asset_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    status = data.get("status", "")
    if status not in ("active", "fully_depreciated"):
        flash(request, "Use the disposal form to dispose or write off an asset", "error")
    elif fixed_asset_store.set_status(asset_id, cid, status):
        flash(request, f"Asset marked {status.replace('_', ' ')}", "success")
    else:
        flash(request, "Failed to update status", "error")
    return RedirectResponse(f"/assets/{asset_id}", status_code=303)


@router.get("/{asset_id}/dispose", name="fixed_assets_dispose_get")
async def asset_dispose_get(asset_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    a = fixed_asset_store.get_asset(asset_id, cid)
    if not a:
        flash(request, "Asset not found", "error")
        return RedirectResponse("/assets/list", status_code=303)
    if a["status"] in ("disposed", "written_off"):
        flash(request, "Asset has already been disposed", "error")
        return RedirectResponse(f"/assets/{asset_id}", status_code=303)
    proceeds = D(request.query_params.get("proceeds"))
    return templates.TemplateResponse("assets/disposal_form.html", _ctx(
        request, active_page="list", asset=a, today=date.today(),
        proceeds=proceeds, preview_gain_loss=gain_loss(proceeds, a["book_value"])))


@router.post("/{asset_id}/dispose", name="fixed_assets_dispose_post")
async def asset_dispose_post(asset_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    d = fixed_asset_store.dispose_asset(asset_id, cid, data)
    if d:
        gl = D(d["gain_loss"])
        word = "gain" if gl > 0 else "loss" if gl < 0 else "no gain or loss"
        flash(request, f"Asset disposed ({d['method']}) — {word} of {abs(gl):,.2f}", "success")
        return RedirectResponse(f"/assets/{asset_id}", status_code=303)
    flash(request, "Failed to dispose asset", "error")
    return RedirectResponse(f"/assets/{asset_id}/dispose", status_code=303)


@router.post("/{asset_id}/maintenance", name="fixed_assets_maintenance_add")
async def maintenance_add(asset_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    if not fixed_asset_store.get_asset(asset_id, cid):
        flash(request, "Asset not found", "error")
        return RedirectResponse("/assets/list", status_code=303)
    if not (data.get("description") or "").strip():
        flash(request, "Maintenance description is required", "error")
    elif fixed_asset_store.add_maintenance(asset_id, cid, data):
        flash(request, "Maintenance entry logged", "success")
    else:
        flash(request, "Failed to log maintenance", "error")
    return RedirectResponse(f"/assets/{asset_id}", status_code=303)


@router.post("/{asset_id}/maintenance/{entry_id}/delete", name="fixed_assets_maintenance_delete")
async def maintenance_delete(asset_id: str, entry_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    if fixed_asset_store.delete_maintenance(entry_id, cid):
        flash(request, "Maintenance entry removed", "success")
    else:
        flash(request, "Failed to remove entry", "error")
    return RedirectResponse(f"/assets/{asset_id}", status_code=303)
