"""
Bank Reconciliation Routes — /bank  (route names bank_*; templates use url_for('bank.x'))

Accounts CRUD, statement import (Excel/CSV + template), reconciliation
workbench (auto/manual match, adjustments, exclude), reconciliation statement
(HTML / Excel / PDF), ageing, bank-charge summary and cash-position reports.
"""
from __future__ import annotations

import io
import logging
import os
import tempfile
from datetime import date
from decimal import Decimal
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, Depends, Request
from fastapi.responses import FileResponse, RedirectResponse, Response

import bank_recon_logic as L
from bank_recon_data_store import bank_store
from bank_recon_logic import D, q2
from deps import current_company, flash, login_required, require_auth, template_context, validate_upload
from template_engine import templates

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/bank", tags=["bank_reconciliation"])
manager_required = require_auth("manager")

_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_UPLOAD_EXTS = {"xlsx", "xls", "csv"}
P = "/bank"


@router.on_event("startup")
async def _startup():
    bank_store.ensure_schema()


def _actor(request: Request) -> str:
    return request.session.get("username", "") or ""


def _ctx(request: Request, **extra) -> dict:
    ctx = template_context(request)
    ctx.update(active_page=extra.pop("active_page", ""), bank_names=L.BANK_NAMES,
               adjustment_kinds=L.ADJUSTMENT_KINDS, adjustment_labels=L.ADJUSTMENT_LABELS,
               today=date.today())
    ctx.update(extra)
    return ctx


async def _form(request: Request) -> dict:
    form = await request.form()
    return {k: v for k, v in form.items()}


def _money(v) -> str:
    return f"{q2(v):,.2f}"


def _d(v) -> str:
    dd = L.parse_date(v)
    return dd.isoformat() if dd else ""


# ── Dashboard ─────────────────────────────────────────────────────
@router.get("/", name="bank_dashboard")
async def dashboard(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    return templates.TemplateResponse("bank/dashboard.html", _ctx(
        request, active_page="dashboard", stats=bank_store.get_dashboard(cid)))


# ── Accounts (static paths before /{id}) ──────────────────────────
@router.get("/accounts", name="bank_accounts")
async def accounts(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    accs = bank_store.get_accounts(cid)
    balances = {r["account_id"]: r for r in _safe_balances(cid)}
    for a in accs:
        b = balances.get(a["id"])
        a["ledger_balance"] = b["ledger_balance"] if b else None
        a["statement_balance"] = b["statement_balance"] if b else None
        a["statement_date"] = b["statement_date"] if b else None
    return templates.TemplateResponse("bank/accounts.html", _ctx(request, active_page="accounts", accounts=accs))


def _safe_balances(cid: str) -> list:
    try:
        from bank_recon_data_store import bank_balances
        return bank_balances(cid)
    except Exception as e:  # pragma: no cover
        logger.error("bank_balances: %s", e); return []


@router.get("/accounts/new", name="bank_account_new_get")
async def account_new_get(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    return templates.TemplateResponse("bank/account_form.html", _ctx(
        request, active_page="accounts", account={"currency": "ETB", "is_active": True}, is_edit=False,
        gl_accounts=bank_store.get_gl_accounts(cid, asset_only=True)))


@router.post("/accounts/new", name="bank_account_new_post")
async def account_new_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    data["created_by"] = _actor(request)
    a = bank_store.create_account(cid, data)
    if a:
        flash(request, f"Bank account '{a['name']}' created", "success")
        return RedirectResponse(f"{P}/accounts", status_code=303)
    flash(request, "Failed to create bank account — a name is required", "error")
    return RedirectResponse(f"{P}/accounts/new", status_code=303)


@router.get("/accounts/{account_id}/edit", name="bank_account_edit_get")
async def account_edit_get(account_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    a = bank_store.get_account(account_id, cid)
    if not a:
        flash(request, "Bank account not found", "error")
        return RedirectResponse(f"{P}/accounts", status_code=303)
    return templates.TemplateResponse("bank/account_form.html", _ctx(
        request, active_page="accounts", account=a, is_edit=True,
        gl_accounts=bank_store.get_gl_accounts(cid, asset_only=True)))


@router.post("/accounts/{account_id}/edit", name="bank_account_edit_post")
async def account_edit_post(account_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    if bank_store.update_account(account_id, cid, data):
        flash(request, "Bank account updated", "success")
        return RedirectResponse(f"{P}/accounts", status_code=303)
    flash(request, "Failed to update bank account", "error")
    return RedirectResponse(f"{P}/accounts/{account_id}/edit", status_code=303)


@router.post("/accounts/{account_id}/toggle", name="bank_account_toggle")
async def account_toggle(account_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    if bank_store.toggle_account(account_id, cid):
        flash(request, "Bank account status changed", "success")
    else:
        flash(request, "Failed to change account status", "error")
    return RedirectResponse(f"{P}/accounts", status_code=303)


# ── Statements ────────────────────────────────────────────────────
@router.get("/statements", name="bank_statements")
async def statements(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    account_id = (request.query_params.get("account") or "").strip() or None
    return templates.TemplateResponse("bank/statements.html", _ctx(
        request, active_page="statements", statements=bank_store.get_statements(cid, account_id),
        accounts=bank_store.get_accounts(cid), account_filter=account_id or ""))


def _import_ctx(request: Request, cid: str, **extra) -> dict:
    return _ctx(request, active_page="import", accounts=bank_store.get_accounts(cid, active_only=True),
                template_columns=L.TEMPLATE_COLUMNS, template_notes=L.TEMPLATE_NOTES,
                header_aliases={k: v[:6] for k, v in L.HEADER_ALIASES.items()},
                report=None, selected_account="", **extra)


@router.get("/statements/import", name="bank_statement_import_get")
async def statement_import_get(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    return templates.TemplateResponse("bank/statement_import.html", _import_ctx(
        request, cid, selected_account=(request.query_params.get("account") or "")))


@router.get("/statements/template", name="bank_statement_template")
async def statement_template(request: Request, user=Depends(login_required)):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"
    ws.append(list(L.TEMPLATE_COLUMNS))
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.append([L.TEMPLATE_SAMPLE.get(c, "") for c in L.TEMPLATE_COLUMNS])
    ws.append(["2026-09-03", "Customer transfer — INV-2026-0042", "FT26246ABC", "", 85000.00, 572500.00])
    ws.append(["2026-09-30", "Service charge", "", 150.00, "", 572350.00])
    for col, width in zip("ABCDEF", (14, 44, 18, 14, 14, 16)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=6):
        for c in row:
            c.number_format = "#,##0.00"
    notes = wb.create_sheet("Field Descriptions")
    notes.append(["Column", "Notes"])
    for c in notes[1]:
        c.font = Font(bold=True)
    for col in L.TEMPLATE_COLUMNS:
        notes.append([col, L.TEMPLATE_NOTES.get(col, "")])
    notes.append([])
    notes.append(["Also accepted", "Txn Date / Value Date / Narration / Particulars / Withdrawal / Deposit / "
                                   "Transaction ID / Cheque No — CBE, Awash, Dashen and Abyssinia exports import as-is."])
    notes.column_dimensions["A"].width = 16
    notes.column_dimensions["B"].width = 100
    fd, path = tempfile.mkstemp(suffix=".xlsx"); os.close(fd)
    wb.save(path)
    return FileResponse(path, filename="bank_statement_template.xlsx", media_type=_XLSX)


def _read_rows(filename: str, content: bytes) -> List[dict]:
    """Excel/CSV → list of {header: cell}. Finds the header row by looking for
    a date-ish and an amount-ish column within the first 15 rows (bank exports
    often have a title block above the table)."""
    name = filename.lower()
    if name.endswith(".csv"):
        import csv
        text = content.decode("utf-8-sig", errors="replace")
        matrix = [row for row in csv.reader(io.StringIO(text))]
    else:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        matrix = [list(r) for r in ws.iter_rows(values_only=True)]
    header_idx = 0
    for i, row in enumerate(matrix[:15]):
        mapping = L.detect_columns([c for c in row if c not in (None, "")])
        if not L.missing_columns(mapping):
            header_idx = i; break
    if not matrix:
        return []
    headers = [str(h).strip() if h not in (None, "") else f"col{j}" for j, h in enumerate(matrix[header_idx])]
    rows = []
    for row in matrix[header_idx + 1:]:
        row = list(row) + [None] * (len(headers) - len(row))
        rows.append({headers[j]: row[j] for j in range(len(headers))})
    return rows


@router.post("/statements/import", name="bank_statement_import_post")
async def statement_import_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    form = await request.form()
    account_id = (form.get("account_id") or "").strip()
    upload = form.get("file")
    if not account_id or not bank_store.get_account(account_id, cid):
        flash(request, "Choose the bank account the statement belongs to", "error")
        return RedirectResponse(f"{P}/statements/import", status_code=303)
    if upload is None or not getattr(upload, "filename", ""):
        flash(request, "Please choose a statement file (.xlsx, .xls or .csv)", "error")
        return RedirectResponse(f"{P}/statements/import?account={account_id}", status_code=303)
    content = await upload.read()
    ok, err = validate_upload(upload.filename, content, allowed_exts=_UPLOAD_EXTS)
    if not ok:
        flash(request, err, "error")
        return RedirectResponse(f"{P}/statements/import?account={account_id}", status_code=303)
    try:
        rows = _read_rows(upload.filename, content)
    except Exception as e:
        flash(request, f"Could not read the file: {e}", "error")
        return RedirectResponse(f"{P}/statements/import?account={account_id}", status_code=303)
    report = bank_store.import_statement(
        cid, account_id, upload.filename, rows, _actor(request),
        closing_balance=form.get("closing_balance") or None, opening_balance=form.get("opening_balance") or None,
        period_start=form.get("period_start") or None, period_end=form.get("period_end") or None)
    if report["imported"]:
        flash(request, f"Imported {report['imported']} statement line(s) from {upload.filename}", "success")
    if report["duplicates"]:
        flash(request, f"{report['duplicates']} line(s) skipped as already imported", "warning")
    if report["errors"]:
        shown = "; ".join(report["errors"][:5]) + (f" (+{len(report['errors']) - 5} more)" if len(report["errors"]) > 5 else "")
        flash(request, f"{len(report['errors'])} problem(s) — {shown}", "warning")
    if not report["total"] and not report["errors"]:
        flash(request, "The file contained no data rows", "warning")
    try:
        from siem_data_store import siem_store
        siem_store.log_upload_event(
            request, module="bank", endpoint=f"{P}/statements/import", filename=upload.filename,
            records_imported=report["imported"],
            status="success" if not report["errors"] else ("partial" if report["imported"] else "failed"),
            details=f"bank statement: {report['imported']} ok, {report['duplicates']} dup, {len(report['errors'])} errors")
    except Exception:
        pass
    if report.get("statement_id"):
        bank_store.refresh_reconciliation(report["statement_id"], cid)
        return RedirectResponse(f"{P}/reconcile/{report['statement_id']}", status_code=303)
    return templates.TemplateResponse("bank/statement_import.html", _import_ctx(
        request, cid, selected_account=account_id, report=report))


@router.get("/statements/{statement_id}", name="bank_statement_detail")
async def statement_detail(statement_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    st = bank_store.get_statement(statement_id, cid)
    if not st:
        flash(request, "Statement not found", "error")
        return RedirectResponse(f"{P}/statements", status_code=303)
    lines = bank_store.get_statement_lines(statement_id, cid)
    counts = {s: sum(1 for l in lines if l["status"] == s) for s in L.LINE_STATUSES}
    return templates.TemplateResponse("bank/statement_detail.html", _ctx(
        request, active_page="statements", statement=st, lines=lines, counts=counts,
        totals=L.statement_totals(lines), recon=bank_store.get_reconciliation(statement_id, cid)))


@router.post("/statements/{statement_id}/balances", name="bank_statement_balances")
async def statement_balances(statement_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    if bank_store.update_statement_balances(statement_id, cid, data.get("opening_balance"), data.get("closing_balance")):
        bank_store.refresh_reconciliation(statement_id, cid)
        flash(request, "Statement balances updated", "success")
    else:
        flash(request, "Failed to update statement balances", "error")
    return RedirectResponse(f"{P}/statements/{statement_id}", status_code=303)


@router.post("/statements/{statement_id}/delete", name="bank_statement_delete")
async def statement_delete(statement_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    ok, err = bank_store.delete_statement(statement_id, cid)
    flash(request, "Statement deleted" if ok else (err or "Failed to delete statement"), "success" if ok else "error")
    return RedirectResponse(f"{P}/statements", status_code=303)


# ── Reports (static, before /reconcile/{id}) ──────────────────────
def _xlsx_file(title: str, headers: List[str], rows: List[list], totals: Optional[list] = None,
               meta: Optional[List[str]] = None, money_cols: Optional[List[int]] = None) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = title[:30].replace("/", "-")
    ws.append([title]); ws["A1"].font = Font(bold=True, size=13)
    for m in (meta or []):
        ws.append([m])
    ws.append([])
    ws.append(headers)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    first_data = ws.max_row + 1
    for r in rows:
        ws.append([float(v) if isinstance(v, Decimal) else (v.isoformat() if isinstance(v, date) else v) for v in r])
    if totals:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in totals])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
    for j, h in enumerate(headers):
        ws.column_dimensions[chr(65 + j) if j < 26 else "AA"].width = 40 if "escription" in h else 16
    for col in (money_cols or []):
        for row in ws.iter_rows(min_row=first_data, min_col=col + 1, max_col=col + 1):
            for c in row:
                c.number_format = "#,##0.00"
    fd, path = tempfile.mkstemp(suffix=".xlsx"); os.close(fd)
    wb.save(path)
    return path


def _pdf_bytes(title: str, meta: List[str], columns: List[tuple], rows: List[list], totals: List[tuple],
               footer: str = "Generated by EBMS — Bank Reconciliation") -> bytes:
    """columns: [(label, relative width, align)]. reportlab via erca_forms.PdfDoc; built-in fallback otherwise."""
    try:
        from erca_forms import PdfDoc, _PAGE_W
    except Exception:  # pragma: no cover
        return _plain_pdf(title, meta, [c[0] for c in columns], rows, totals)
    pdf = PdfDoc()
    pdf.text(_PAGE_W / 2, 50, title, 14, True, "center")
    pdf.y = 62
    for m in meta:
        pdf.text(40, pdf.y + 9, m, 8.5); pdf.y += 12
    pdf.y += 6
    total_w = _PAGE_W - 80
    scale = total_w / sum(w for _, w, _ in columns)
    cols = [(lab, w * scale, al) for lab, w, al in columns]
    pdf.table(cols, rows, size=8)
    pdf.y += 8
    for lab, val in totals:
        bold = lab.isupper()
        pdf.text(_PAGE_W - 220, pdf.y + 10, lab, 9, bold)
        pdf.text(_PAGE_W - 40, pdf.y + 10, val, 9, bold, "right")
        pdf.y += 13
    pdf.footer(footer[:160])
    return pdf.render()


def _plain_pdf(title: str, meta: List[str], headers: List[str], rows: List[list], totals: List[tuple]) -> bytes:
    lines = [title] + meta + [""] + [" | ".join(headers)] + [" | ".join(str(c) for c in r) for r in rows] + [""]
    lines += [f"{k}: {v}" for k, v in totals]
    ops = ["BT /F1 9 Tf 40 800 Td 12 TL"]
    for ln in lines[:60]:
        ln = ln.encode("latin-1", "replace").decode("latin-1").replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        ops.append(f"({ln}) Tj T*")
    ops.append("ET")
    stream = "\n".join(ops)
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream", "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out = io.BytesIO(); out.write(b"%PDF-1.4\n"); offs = []
    for i, o in enumerate(objs, 1):
        offs.append(out.tell()); out.write(f"{i} 0 obj\n{o}\nendobj\n".encode("latin-1"))
    x = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for o in offs:
        out.write(f"{o:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{x}\n%%EOF\n".encode())
    return out.getvalue()


def _pdf_response(data: bytes, filename: str) -> Response:
    return Response(content=data, media_type="application/pdf",
                    headers={"Content-Disposition": f'inline; filename="{filename}"', "Cache-Control": "private, no-store"})


@router.get("/reports/ageing", name="bank_report_ageing")
async def report_ageing(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    account_id = (request.query_params.get("account") or "").strip() or None
    as_of = L.parse_date(request.query_params.get("as_of")) or date.today()
    rep = bank_store.ageing(cid, account_id, as_of)
    if request.query_params.get("format") == "xlsx":
        rows = []
        for b in rep["buckets"]:
            for it in rep["statement"][b]["items"]:
                rows.append(["Statement", b, it.get("account_name") or "", it["txn_date"], it.get("description") or "",
                             it.get("reference") or "", q2(it.get("debit")), q2(it.get("credit")), it.get("age_days")])
            for it in rep["ledger"][b]["items"]:
                rows.append(["Ledger", b, it.get("account_name") or "", it["txn_date"], it.get("entry_description") or it.get("description") or "",
                             it.get("reference") or "", q2(it.get("debit")), q2(it.get("credit")), it.get("age_days")])
        path = _xlsx_file("Unreconciled items ageing", ["Side", "Bucket", "Account", "Date", "Description", "Reference",
                                                         "Debit", "Credit", "Age (days)"], rows,
                          meta=[f"As of {as_of.isoformat()}"], money_cols=[6, 7])
        return FileResponse(path, filename=f"bank_ageing_{as_of.isoformat()}.xlsx", media_type=_XLSX)
    return templates.TemplateResponse("bank/report_ageing.html", _ctx(
        request, active_page="report_ageing", report=rep, accounts=bank_store.get_accounts(cid),
        account_filter=account_id or "", as_of=as_of))


@router.get("/reports/charges", name="bank_report_charges")
async def report_charges(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    qp = request.query_params
    today = date.today()
    date_from = L.parse_date(qp.get("from")) or today.replace(month=1, day=1)
    date_to = L.parse_date(qp.get("to")) or today
    rep = bank_store.charges_summary(cid, date_from, date_to)
    if qp.get("format") == "xlsx":
        rows = [[r["adj_date"], r["account_name"], r["bank_name"], L.ADJUSTMENT_LABELS.get(r["kind"], r["kind"]),
                 r["description"], r["gl_account_code"], q2(r["amount"]), "yes" if r["journal_entry_id"] else "no", r["created_by"]]
                for r in rep["rows"]]
        path = _xlsx_file("Bank charges & interest", ["Date", "Account", "Bank", "Kind", "Description", "GL account",
                                                      "Amount (+ raises bank)", "Posted", "By"], rows,
                          totals=["TOTAL", "", "", "", "", "", q2(rep["total_interest"] + rep["total_other"] - rep["total_charges"]), "", ""],
                          meta=[f"{date_from.isoformat()} → {date_to.isoformat()}"], money_cols=[6])
        return FileResponse(path, filename=f"bank_charges_{date_from.isoformat()}_{date_to.isoformat()}.xlsx", media_type=_XLSX)
    return templates.TemplateResponse("bank/report_charges.html", _ctx(
        request, active_page="report_charges", report=rep, date_from=date_from, date_to=date_to))


@router.get("/reports/cash-position", name="bank_report_cash_position")
async def report_cash_position(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    rep = bank_store.cash_position(cid)
    if request.query_params.get("format") == "xlsx":
        rows = [[r["account"]["name"], r["account"]["bank_name"], r["account"]["account_number"], r["account"]["currency"],
                 r["account"]["gl_account_code"], r["ledger_balance"], r["statement_balance"] if r["statement_balance"] is not None else "",
                 r["statement_date"] or "", r["difference"] if r["difference"] is not None else "", r["unmatched"]]
                for r in rep["rows"]]
        path = _xlsx_file("Cash position", ["Account", "Bank", "Number", "Currency", "GL", "Ledger balance",
                                            "Last statement balance", "Statement date", "Difference", "Unmatched lines"], rows,
                          totals=["TOTAL", "", "", "", "", rep["total_ledger"], rep["total_bank"], "", rep["difference"], ""],
                          meta=[f"As of {rep['as_of'].isoformat()}"], money_cols=[5, 6, 8])
        return FileResponse(path, filename=f"cash_position_{rep['as_of'].isoformat()}.xlsx", media_type=_XLSX)
    return templates.TemplateResponse("bank/report_cash_position.html", _ctx(
        request, active_page="report_cash", report=rep))


# ── Reconciliation workbench ──────────────────────────────────────
def _workbench_data(cid: str, statement_id: str):
    st = bank_store.get_statement(statement_id, cid)
    if not st:
        return None
    recon = bank_store.refresh_reconciliation(statement_id, cid) or {}
    all_lines = bank_store.get_statement_lines(statement_id, cid)
    unmatched = [l for l in all_lines if l["status"] == "unmatched"]
    matched = [l for l in all_lines if l["status"] != "unmatched"]
    ledger = bank_store.ledger_lines(cid, st.get("gl_account_code") or "", None,
                                     _shift(st.get("period_end"), 10), True)
    payments = bank_store.payment_candidates(cid, _shift(st.get("period_start"), -10), _shift(st.get("period_end"), 10))
    for l in unmatched:
        l["candidates"] = [{"score": sc, "kind": "journal", "line": ln} for sc, ln in L.candidates(l, ledger, 3)]
        l["candidates"] += [{"score": sc, "kind": "payment", "line": ln} for sc, ln in L.candidates(l, payments, 2)]
        l["candidates"].sort(key=lambda c: -c["score"])
    # ledger lines matched to lines on THIS statement (for the "matched" pane)
    matched_ledger: Dict[str, dict] = {}
    for ln in bank_store.ledger_lines(cid, st.get("gl_account_code") or "", st.get("period_start"),
                                      _shift(st.get("period_end"), 10), False):
        if ln.get("reconciled_against"):
            matched_ledger[ln["reconciled_against"]] = ln
    for l in matched:
        l["ledger"] = matched_ledger.get(l["id"])
    return {"statement": st, "recon": recon, "unmatched": unmatched, "matched": matched, "ledger": ledger,
            "payments": payments, "ledger_totals": L.outstanding_totals(ledger),
            "unmatched_totals": L.statement_totals(unmatched)}


def _shift(d, days):
    from datetime import timedelta
    dd = L.parse_date(d)
    return dd + timedelta(days=days) if dd else None


@router.get("/reconcile/{statement_id}", name="bank_workbench")
async def workbench(statement_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = _workbench_data(cid, statement_id)
    if not data:
        flash(request, "Statement not found", "error")
        return RedirectResponse(f"{P}/statements", status_code=303)
    return templates.TemplateResponse("bank/workbench.html", _ctx(
        request, active_page="statements", gl_accounts=bank_store.get_gl_accounts(cid),
        ledger_available=bank_store.ledger_available(), **data))


@router.post("/reconcile/{statement_id}/auto", name="bank_auto_match")
async def auto_match(statement_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    res = bank_store.run_auto_match(statement_id, cid, _actor(request) or "auto-match")
    if res.get("error"):
        flash(request, res["error"], "error")
    else:
        msg = f"Auto-match: {res['matched']} ledger match(es)"
        if res["payments"]:
            msg += f", {res['payments']} payment match(es)"
        msg += f", {res['skipped']} left for review (of {res['scanned']} unmatched)"
        flash(request, msg, "success" if (res["matched"] or res["payments"]) else "info")
    return RedirectResponse(f"{P}/reconcile/{statement_id}", status_code=303)


@router.post("/reconcile/{statement_id}/match", name="bank_match")
async def match(statement_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    line_id = (data.get("line_id") or "").strip()
    target = (data.get("target") or "").strip()      # "journal:<line_id>:<entry_id>" or "payment:<id>"
    parts = target.split(":")
    if not line_id or len(parts) < 2 or parts[0] not in ("journal", "payment"):
        flash(request, "Pick a ledger line or payment to match", "error")
        return RedirectResponse(f"{P}/reconcile/{statement_id}", status_code=303)
    kind, target_id = parts[0], parts[1]
    entry_id = parts[2] if len(parts) > 2 else ""
    ok, err = bank_store.match_line(line_id, cid, kind, target_id, _actor(request), entry_id)
    flash(request, "Line matched" if ok else (err or "Failed to match line"), "success" if ok else "error")
    return RedirectResponse(f"{P}/reconcile/{statement_id}", status_code=303)


@router.post("/reconcile/{statement_id}/unmatch", name="bank_unmatch")
async def unmatch(statement_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    ok, err = bank_store.unmatch_line((data.get("line_id") or "").strip(), cid, _actor(request))
    flash(request, "Line unmatched" if ok else (err or "Failed to unmatch line"), "success" if ok else "error")
    return RedirectResponse(f"{P}/reconcile/{statement_id}#matched", status_code=303)


@router.post("/reconcile/{statement_id}/exclude", name="bank_exclude")
async def exclude(statement_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    ok, err = bank_store.exclude_line((data.get("line_id") or "").strip(), cid, _actor(request))
    flash(request, "Line excluded from the reconciliation" if ok else (err or "Failed to exclude line"),
          "success" if ok else "error")
    return RedirectResponse(f"{P}/reconcile/{statement_id}", status_code=303)


@router.post("/reconcile/{statement_id}/adjust", name="bank_adjust")
async def adjust(statement_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    data = await _form(request)
    adj, err = bank_store.create_adjustment(statement_id, cid, data, _actor(request))
    if adj:
        label = L.ADJUSTMENT_LABELS.get(adj["kind"], adj["kind"])
        if adj.get("posted"):
            flash(request, f"{label} of {_money(abs(D(adj['amount'])))} recorded and posted to journal "
                           f"{(adj.get('journal_entry_id') or '')[:8]}…", "success")
        else:
            flash(request, f"{label} of {_money(abs(D(adj['amount'])))} recorded — NOT posted to the general ledger "
                           f"(no GL account chosen or ledger unavailable); book it manually", "warning")
    else:
        flash(request, err or "Failed to record adjustment", "error")
    return RedirectResponse(f"{P}/reconcile/{statement_id}", status_code=303)


@router.post("/reconcile/{statement_id}/adjust/{adjustment_id}/delete", name="bank_adjust_delete")
async def adjust_delete(statement_id: str, adjustment_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    ok, err = bank_store.delete_adjustment(adjustment_id, cid)
    flash(request, "Adjustment removed" if ok else (err or "Failed to remove adjustment"), "success" if ok else "error")
    return RedirectResponse(f"{P}/reconcile/{statement_id}", status_code=303)


@router.post("/reconcile/{statement_id}/status", name="bank_recon_status")
async def recon_status(statement_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    data = await _form(request)
    status = (data.get("status") or "").strip()
    notes = data.get("notes")
    ok, err = bank_store.set_reconciliation_status(statement_id, cid, status, _actor(request),
                                                   notes if notes is not None and notes != "" else None)
    if ok:
        flash(request, {"closed": "Reconciliation closed", "balanced": "Reconciliation marked balanced",
                        "open": "Reconciliation reopened"}.get(status, "Status updated"), "success")
    else:
        flash(request, err or "Failed to update status", "error")
    return RedirectResponse(f"{P}/reconcile/{statement_id}", status_code=303)


# ── Reconciliation statement report ───────────────────────────────
def _recon_report(cid: str, statement_id: str):
    st = bank_store.get_statement(statement_id, cid)
    if not st:
        return None, None
    recon = bank_store.refresh_reconciliation(statement_id, cid) or {}
    return st, recon


@router.get("/reconcile/{statement_id}/report", name="bank_recon_report")
async def recon_report(statement_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    st, recon = _recon_report(cid, statement_id)
    if not st:
        flash(request, "Statement not found", "error")
        return RedirectResponse(f"{P}/statements", status_code=303)
    fmt = request.query_params.get("format") or "html"
    title = "Bank Reconciliation Statement"
    meta = [f"{st['bank_name']} — {st['account_name']} ({st['account_number']})",
            f"Period {_d(st['period_start'])} → {_d(st['period_end'])}   As of {_d(recon.get('as_of'))}",
            f"GL account {st.get('gl_account_code') or '—'}   Status: {recon.get('status', 'open')}"]
    summary = [
        ("Balance per bank statement", recon.get("statement_balance")),
        ("Add: deposits in transit (ledger debits not on statement)", recon.get("outstanding_deposits")),
        ("Less: outstanding payments (ledger credits not on statement)", recon.get("outstanding_payments")),
        ("ADJUSTED BANK BALANCE", recon.get("adjusted_bank_balance")),
        ("Balance per general ledger", recon.get("ledger_balance")),
        ("DIFFERENCE", recon.get("difference")),
        ("Unmatched statement credits (not yet booked)", recon.get("unmatched_statement_credits")),
        ("Unmatched statement debits (not yet booked)", recon.get("unmatched_statement_debits")),
        ("UNEXPLAINED DIFFERENCE", recon.get("unexplained_difference")),
    ]
    out_rows = [[ln["entry_date"], ln.get("reference") or "", ln.get("entry_description") or ln.get("description") or "",
                 q2(ln["debit"]), q2(ln["credit"])] for ln in recon.get("outstanding_lines", [])]
    unm_rows = [[ln["txn_date"], ln.get("reference") or "", ln.get("description") or "", q2(ln["debit"]), q2(ln["credit"])]
                for ln in recon.get("unmatched_lines", [])]
    if fmt == "xlsx":
        rows = [[lab, q2(val)] for lab, val in summary] + [[], ["OUTSTANDING LEDGER ITEMS", ""]]
        rows += [["Date", "Reference", "Description", "Debit", "Credit"]] + out_rows
        rows += [[], ["UNMATCHED STATEMENT LINES", ""], ["Date", "Reference", "Description", "Debit", "Credit"]] + unm_rows
        path = _xlsx_file(title, ["Item", "Amount", "", "", ""], rows, meta=meta, money_cols=[1, 3, 4])
        return FileResponse(path, filename=f"bank_reconciliation_{_d(st['period_end'])}.xlsx", media_type=_XLSX)
    if fmt == "pdf":
        cols = [("Date", 0.8, "left"), ("Reference", 1.0, "left"), ("Description", 2.6, "left"), ("Debit", 0.9, "right"), ("Credit", 0.9, "right")]
        rows = [["OUTSTANDING LEDGER ITEMS", "", "", "", ""]] + [[_d(r[0]), r[1], r[2][:60], _money(r[3]) if r[3] else "", _money(r[4]) if r[4] else ""] for r in out_rows]
        rows += [["UNMATCHED STATEMENT LINES", "", "", "", ""]] + [[_d(r[0]), r[1], r[2][:60], _money(r[3]) if r[3] else "", _money(r[4]) if r[4] else ""] for r in unm_rows]
        data = _pdf_bytes(title, meta, cols, rows, [(lab, _money(val)) for lab, val in summary])
        return _pdf_response(data, f"bank_reconciliation_{_d(st['period_end'])}.pdf")
    return templates.TemplateResponse("bank/recon_statement.html", _ctx(
        request, active_page="statements", statement=st, recon=recon, summary=summary, generated=date.today()))
