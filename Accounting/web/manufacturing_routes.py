"""
Manufacturing / Production Routes — /manufacturing
"""
from __future__ import annotations

import logging
import os
import tempfile
from datetime import date, timedelta
from decimal import Decimal

from fastapi import APIRouter, Depends, Request
from fastapi.responses import FileResponse, RedirectResponse

from deps import current_company, flash, login_required, require_auth, template_context
from template_engine import templates
import manufacturing_data_store as mfg
from manufacturing_data_store import (
    CALENDAR_KINDS, CAPACITY_UNITS, D, ISSUE_KINDS, MACHINE_STATUSES, ORDER_STATUSES, ORDER_TYPES, PLAN_BASIS,
    PLAN_PERIODS, PROCESS_FIELDS, PRODUCT_TYPES, PRODUCT_UNITS, TDS_FIELDS, WC_TYPES, manufacturing_store as store,
    params_from_form, params_text, variance,
)

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/manufacturing", tags=["manufacturing"])
manager_required = require_auth("manager")

_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
BASE = "/manufacturing"

mfg._install_approval_hook()


# ── helpers ───────────────────────────────────────────────────────
def _ctx(request: Request, active_page: str = "", **extra) -> dict:
    ctx = template_context(request)
    ctx.update(active_page=active_page, order_statuses=ORDER_STATUSES, order_types=ORDER_TYPES, wc_types=WC_TYPES,
               capacity_units=CAPACITY_UNITS, product_units=PRODUCT_UNITS, product_types=PRODUCT_TYPES,
               machine_statuses=MACHINE_STATUSES, calendar_kinds=CALENDAR_KINDS, plan_periods=PLAN_PERIODS,
               plan_basis=PLAN_BASIS, issue_kinds=ISSUE_KINDS, tds_fields=TDS_FIELDS, process_fields=PROCESS_FIELDS,
               params_text=params_text, today=date.today())
    ctx.update(extra)
    return ctx


async def _form(request: Request) -> dict:
    form = await request.form()
    return {k: v for k, v in form.items()}


def _actor(request: Request) -> str:
    return request.session.get("username", "")


def _filters(request: Request) -> dict:
    qp = request.query_params
    f = {k: (qp.get(k) or "").strip() for k in ("date_from", "date_to", "plant_id", "work_center_id", "machine_id", "group")}
    if not f["date_from"] and not f["date_to"]:
        f["date_from"] = (date.today() - timedelta(days=30)).isoformat()
    return f


def _filter_ctx(request: Request, cid: str) -> dict:
    return dict(filters=_filters(request), plants=store.get_plants(cid), work_centers=store.get_work_centers(cid),
                machines=store.get_machines(cid))


def _redirect(path: str) -> RedirectResponse:
    return RedirectResponse(path, status_code=303)


def _ok(request: Request, ok, success: str, failure: str) -> None:
    flash(request, success if ok else failure, "success" if ok else "error")


# ── dashboard / process map / settings ────────────────────────────
@router.get("/", name="manufacturing_dashboard")
async def dashboard(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    store.seed_defaults(cid)
    return templates.TemplateResponse("manufacturing/dashboard.html", _ctx(
        request, "dashboard", stats=store.dashboard(cid), variances=store.cost_variance_by_product(cid)))


@router.get("/process-map", name="manufacturing_process_map")
async def process_map(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    ref = (request.query_params.get("ref") or "").strip()
    return templates.TemplateResponse("manufacturing/process_map.html", _ctx(
        request, "process_map", ref=ref, steps=store.process_status(cid, ref),
        recent_orders=store.get_orders(cid, limit=15)))


@router.get("/settings", name="manufacturing_settings")
async def settings_get(request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    return templates.TemplateResponse("manufacturing/settings.html", _ctx(
        request, "settings", settings=store.get_settings(cid), plants=store.get_plants(cid),
        gl_accounts=store.get_gl_accounts(cid)))


@router.post("/settings", name="manufacturing_settings_post")
async def settings_post(request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    _ok(request, store.save_settings(cid, await _form(request)), "Settings saved", "Failed to save settings")
    return _redirect(f"{BASE}/settings")


# ── master data: plants / work centres / machines ─────────────────
@router.get("/plants", name="manufacturing_plants")
async def plants(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    edit_id = request.query_params.get("edit")
    return templates.TemplateResponse("manufacturing/plants.html", _ctx(
        request, "plants", plants=store.get_plants(cid), item=store.get_plant(edit_id, cid) if edit_id else {}))


@router.post("/plants/save", name="manufacturing_plant_save")
async def plant_save(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    _ok(request, store.save_plant(cid, data, data.get("id") or None), "Plant saved", "Failed to save plant — name is required")
    return _redirect(f"{BASE}/plants")


@router.get("/work-centers", name="manufacturing_work_centers")
async def work_centers(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    edit_id = request.query_params.get("edit")
    return templates.TemplateResponse("manufacturing/work_centers.html", _ctx(
        request, "work_centers", work_centers=store.get_work_centers(cid), plants=store.get_plants(cid),
        item=store.get_work_center(edit_id, cid) if edit_id else {}))


@router.post("/work-centers/save", name="manufacturing_work_center_save")
async def work_center_save(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    _ok(request, store.save_work_center(cid, data, data.get("id") or None), "Work centre saved", "Failed to save work centre")
    return _redirect(f"{BASE}/work-centers")


@router.get("/machines", name="manufacturing_machines")
async def machines(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    edit_id = request.query_params.get("edit")
    return templates.TemplateResponse("manufacturing/machines.html", _ctx(
        request, "machines", machines=store.get_machines(cid), work_centers=store.get_work_centers(cid),
        item=store.get_machine(edit_id, cid) if edit_id else {}))


@router.post("/machines/save", name="manufacturing_machine_save")
async def machine_save(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    _ok(request, store.save_machine(cid, data, data.get("id") or None), "Machine saved", "Failed to save machine")
    return _redirect(f"{BASE}/machines")


@router.post("/machines/{machine_id}/status", name="manufacturing_machine_status")
async def machine_status(machine_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    _ok(request, store.set_machine_status(machine_id, cid, data.get("status", "")), "Machine status updated", "Invalid status")
    return _redirect(f"{BASE}/machines")


# ── products / TDS / BOM / routing ────────────────────────────────
@router.get("/products", name="manufacturing_products")
async def products(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    qp = request.query_params
    return templates.TemplateResponse("manufacturing/products.html", _ctx(
        request, "products", products=store.get_products(cid, qp.get("type") or None, (qp.get("q") or "").strip() or None),
        type_filter=qp.get("type") or "", search=(qp.get("q") or "").strip()))


@router.get("/products/new", name="manufacturing_product_new_get")
async def product_new_get(request: Request, user=Depends(login_required)):
    return templates.TemplateResponse("manufacturing/product_form.html", _ctx(request, "products", product={}, is_edit=False))


@router.post("/products/new", name="manufacturing_product_new_post")
async def product_new_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    pid = store.save_product(cid, await _form(request))
    if pid:
        flash(request, "Product created", "success")
        return _redirect(f"{BASE}/products/{pid}")
    flash(request, "Failed to create product — code and name are required and the code must be unique", "error")
    return _redirect(f"{BASE}/products/new")


@router.get("/products/{product_id}", name="manufacturing_product_detail")
async def product_detail(product_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    p = store.get_product(product_id, cid)
    if not p:
        flash(request, "Product not found", "error")
        return _redirect(f"{BASE}/products")
    return templates.TemplateResponse("manufacturing/product_detail.html", _ctx(
        request, "products", product=p, tds_list=store.get_tds_list(product_id), boms=store.get_boms(product_id),
        routings=store.get_routings(product_id), material_costs=store.get_material_costs(cid)))


@router.get("/products/{product_id}/edit", name="manufacturing_product_edit_get")
async def product_edit_get(product_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    p = store.get_product(product_id, cid)
    if not p:
        flash(request, "Product not found", "error")
        return _redirect(f"{BASE}/products")
    return templates.TemplateResponse("manufacturing/product_form.html", _ctx(request, "products", product=p, is_edit=True))


@router.post("/products/{product_id}/edit", name="manufacturing_product_edit_post")
async def product_edit_post(product_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    ok = store.save_product(cid, data, product_id)
    if ok and data.get("mirror_inventory"):
        store.mirror_to_inventory(product_id, cid)
    _ok(request, ok, "Product updated", "Failed to update product")
    return _redirect(f"{BASE}/products/{product_id}")


@router.post("/products/{product_id}/tds/new", name="manufacturing_tds_new")
async def tds_new(product_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    tid = store.create_tds(cid, product_id, params_from_form(data, TDS_FIELDS), data.get("notes", ""), _actor(request))
    _ok(request, tid, "TDS version created (draft)", "Failed to create TDS")
    return _redirect(f"{BASE}/products/{product_id}#tds")


@router.post("/tds/{tds_id}/approve", name="manufacturing_tds_approve")
async def tds_approve(tds_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    t = store.get_tds(tds_id, cid)
    _ok(request, t and store.approve_tds(tds_id, cid, _actor(request)), "TDS approved — earlier version superseded", "Failed to approve TDS")
    return _redirect(f"{BASE}/products/{t['product_id']}#tds" if t else f"{BASE}/products")


@router.post("/products/{product_id}/bom/new", name="manufacturing_bom_new")
async def bom_new(product_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    bid = store.create_bom(cid, product_id, data, data.get("copy_from") or None)
    if bid:
        flash(request, "New BOM version created (draft) — add lines then activate", "success")
        return _redirect(f"{BASE}/bom/{bid}")
    flash(request, "Failed to create BOM", "error")
    return _redirect(f"{BASE}/products/{product_id}")


@router.get("/bom/{bom_id}", name="manufacturing_bom_detail")
async def bom_detail(bom_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    b = store.get_bom(bom_id, cid)
    if not b:
        flash(request, "BOM not found", "error")
        return _redirect(f"{BASE}/products")
    costs = store.material_cost_map(cid)
    return templates.TemplateResponse("manufacturing/bom_detail.html", _ctx(
        request, "products", bom=b, product=store.get_product(b["product_id"], cid), work_centers=store.get_work_centers(cid),
        products=store.get_products(cid, "semi_finished"), tds_list=store.get_tds_list(b["product_id"]),
        unit_cost=mfg.planned_material_cost(b, 1, costs), costs=costs))


@router.post("/bom/{bom_id}/lines/add", name="manufacturing_bom_line_add")
async def bom_line_add(bom_id: str, request: Request, user=Depends(login_required)):
    _ok(request, store.add_bom_line(bom_id, await _form(request)), "BOM line added", "Material name or code is required")
    return _redirect(f"{BASE}/bom/{bom_id}")


@router.post("/bom/{bom_id}/lines/{line_id}/delete", name="manufacturing_bom_line_delete")
async def bom_line_delete(bom_id: str, line_id: str, request: Request, user=Depends(login_required)):
    _ok(request, store.delete_bom_line(line_id, bom_id), "BOM line removed", "Failed to remove line")
    return _redirect(f"{BASE}/bom/{bom_id}")


@router.post("/bom/{bom_id}/status", name="manufacturing_bom_status")
async def bom_status(bom_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    data = await _form(request)
    _ok(request, store.set_bom_status(bom_id, cid, data.get("status", "")), "BOM status updated", "Failed to update BOM status")
    return _redirect(f"{BASE}/bom/{bom_id}")


@router.post("/products/{product_id}/routing/new", name="manufacturing_routing_new")
async def routing_new(product_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    rid = store.create_routing(cid, product_id, data.get("notes", ""))
    if rid:
        flash(request, "Routing version created (draft) — add operations then activate", "success")
        return _redirect(f"{BASE}/routing/{rid}")
    flash(request, "Failed to create routing", "error")
    return _redirect(f"{BASE}/products/{product_id}")


@router.get("/routing/{routing_id}", name="manufacturing_routing_detail")
async def routing_detail(routing_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    r = store.get_routing(routing_id, cid)
    if not r:
        flash(request, "Routing not found", "error")
        return _redirect(f"{BASE}/products")
    return templates.TemplateResponse("manufacturing/routing_detail.html", _ctx(
        request, "products", routing=r, product=store.get_product(r["product_id"], cid), work_centers=store.get_work_centers(cid),
        hours_per_1000=mfg.routing_hours(r, 1000)))


@router.post("/routing/{routing_id}/ops/add", name="manufacturing_routing_op_add")
async def routing_op_add(routing_id: str, request: Request, user=Depends(login_required)):
    data = await _form(request)
    _ok(request, store.add_routing_op(routing_id, data, params_from_form(data, PROCESS_FIELDS)), "Operation added", "Operation name is required")
    return _redirect(f"{BASE}/routing/{routing_id}")


@router.post("/routing/{routing_id}/ops/{op_id}/delete", name="manufacturing_routing_op_delete")
async def routing_op_delete(routing_id: str, op_id: str, request: Request, user=Depends(login_required)):
    _ok(request, store.delete_routing_op(op_id, routing_id), "Operation removed", "Failed to remove operation")
    return _redirect(f"{BASE}/routing/{routing_id}")


@router.post("/routing/{routing_id}/status", name="manufacturing_routing_status")
async def routing_status(routing_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    data = await _form(request)
    _ok(request, store.set_routing_status(routing_id, cid, data.get("status", "")), "Routing status updated", "Failed to update routing")
    return _redirect(f"{BASE}/routing/{routing_id}")


# ── calendar / shifts / downtime categories / materials ───────────
@router.get("/calendar", name="manufacturing_calendar")
async def calendar_page(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    year = int(request.query_params.get("year") or date.today().year)
    return templates.TemplateResponse("manufacturing/calendar.html", _ctx(
        request, "calendar", entries=store.get_calendar(cid, year), year=year, plants=store.get_plants(cid)))


@router.post("/calendar/new", name="manufacturing_calendar_new")
async def calendar_new(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    _ok(request, store.add_calendar_entry(cid, await _form(request)), "Calendar entry added", "Date is required")
    return _redirect(f"{BASE}/calendar")


@router.post("/calendar/{entry_id}/delete", name="manufacturing_calendar_delete")
async def calendar_delete(entry_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    _ok(request, store.delete_calendar_entry(entry_id, cid), "Calendar entry removed", "Failed to remove entry")
    return _redirect(f"{BASE}/calendar")


@router.get("/shifts", name="manufacturing_shifts")
async def shifts(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    edit_id = request.query_params.get("edit")
    items = store.get_shifts(cid)
    return templates.TemplateResponse("manufacturing/shifts.html", _ctx(
        request, "shifts", shifts=items, work_centers=store.get_work_centers(cid),
        item=next((s for s in items if s["id"] == edit_id), {}) if edit_id else {}, shift_hours=mfg.shift_hours))


@router.post("/shifts/save", name="manufacturing_shift_save")
async def shift_save(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    _ok(request, store.save_shift(cid, data, data.get("id") or None), "Shift saved", "Shift name is required")
    return _redirect(f"{BASE}/shifts")


@router.get("/downtime-categories", name="manufacturing_downtime_categories")
async def downtime_categories(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    store.seed_defaults(cid)
    return templates.TemplateResponse("manufacturing/downtime_categories.html", _ctx(
        request, "downtime_categories", categories=store.get_downtime_categories(cid), scrap_types=store.get_scrap_types(cid)))


@router.post("/downtime-categories/new", name="manufacturing_downtime_category_new")
async def downtime_category_new(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    _ok(request, store.add_downtime_category(cid, data.get("name", "")), "Category added", "Category name is required")
    return _redirect(f"{BASE}/downtime-categories")


@router.post("/downtime-categories/{category_id}/reasons/new", name="manufacturing_downtime_reason_new")
async def downtime_reason_new(category_id: str, request: Request, user=Depends(login_required)):
    data = await _form(request)
    _ok(request, store.add_downtime_reason(category_id, data.get("name", "")), "Reason added", "Reason name is required")
    return _redirect(f"{BASE}/downtime-categories")


@router.post("/scrap-types/new", name="manufacturing_scrap_type_new")
async def scrap_type_new(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    _ok(request, store.add_scrap_type(cid, data.get("name", "")), "Scrap type added", "Name is required")
    return _redirect(f"{BASE}/downtime-categories")


@router.get("/materials", name="manufacturing_materials")
async def materials(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    return templates.TemplateResponse("manufacturing/materials.html", _ctx(request, "materials", costs=store.get_material_costs(cid)))


@router.post("/materials/save", name="manufacturing_material_save")
async def material_save(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    _ok(request, store.save_material_cost(cid, await _form(request)), "Material cost saved", "Material code is required")
    return _redirect(f"{BASE}/materials")


@router.post("/materials/{code}/delete", name="manufacturing_material_delete")
async def material_delete(code: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    _ok(request, store.delete_material_cost(cid, code), "Material cost removed", "Failed to remove")
    return _redirect(f"{BASE}/materials")


# ── planning ──────────────────────────────────────────────────────
@router.get("/plans", name="manufacturing_plans")
async def plans(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    status = request.query_params.get("status") or None
    return templates.TemplateResponse("manufacturing/plans.html", _ctx(
        request, "plans", plans=store.get_plans(cid, status), status_filter=status or ""))


@router.get("/plans/new", name="manufacturing_plan_new_get")
async def plan_new_get(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    return templates.TemplateResponse("manufacturing/plan_form.html", _ctx(request, "plans", plants=store.get_plants(cid), plan={}))


@router.post("/plans/new", name="manufacturing_plan_new_post")
async def plan_new_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    pid = store.create_plan(cid, await _form(request), _actor(request))
    if pid:
        flash(request, "Production plan created — add product lines", "success")
        return _redirect(f"{BASE}/plans/{pid}")
    flash(request, "Failed to create plan", "error")
    return _redirect(f"{BASE}/plans/new")


@router.get("/plans/{plan_id}", name="manufacturing_plan_detail")
async def plan_detail(plan_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    p = store.get_plan(plan_id, cid)
    if not p:
        flash(request, "Plan not found", "error")
        return _redirect(f"{BASE}/plans")
    return templates.TemplateResponse("manufacturing/plan_detail.html", _ctx(
        request, "plans", plan=p, lines=store.plan_vs_actual(cid, p), products=store.get_products(cid, active_only=True),
        work_centers=store.get_work_centers(cid)))


@router.post("/plans/{plan_id}/lines/add", name="manufacturing_plan_line_add")
async def plan_line_add(plan_id: str, request: Request, user=Depends(login_required)):
    _ok(request, store.add_plan_line(plan_id, await _form(request)), "Plan line added", "Product is required")
    return _redirect(f"{BASE}/plans/{plan_id}")


@router.post("/plans/{plan_id}/lines/{line_id}/delete", name="manufacturing_plan_line_delete")
async def plan_line_delete(plan_id: str, line_id: str, request: Request, user=Depends(login_required)):
    _ok(request, store.delete_plan_line(line_id, plan_id), "Plan line removed", "Failed to remove line")
    return _redirect(f"{BASE}/plans/{plan_id}")


@router.post("/plans/{plan_id}/status", name="manufacturing_plan_status")
async def plan_status(plan_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    data = await _form(request)
    _ok(request, store.set_plan_status(plan_id, cid, data.get("status", "")), "Plan status updated", "Invalid status")
    return _redirect(f"{BASE}/plans/{plan_id}")


@router.get("/capacity", name="manufacturing_capacity")
async def capacity(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    qp = request.query_params
    today = date.today()
    start = qp.get("period_start") or today.replace(day=1).isoformat()
    end = qp.get("period_end") or ((today.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)).isoformat()
    plant_id = qp.get("plant_id") or None
    try:
        rows = store.compute_capacity(cid, date.fromisoformat(start), date.fromisoformat(end), plant_id)
    except ValueError:
        rows = []
        flash(request, "Invalid period", "error")
    return templates.TemplateResponse("manufacturing/capacity.html", _ctx(
        request, "capacity", rows=rows, period_start=start, period_end=end, plant_id=plant_id or "",
        plants=store.get_plants(cid), work_centers=store.get_work_centers(cid), saved=store.get_capacity_plans(cid)))


@router.post("/capacity/save", name="manufacturing_capacity_save")
async def capacity_save(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    _ok(request, store.save_capacity_plan(cid, await _form(request)), "Capacity plan saved", "Work centre is required")
    return _redirect(f"{BASE}/capacity")


@router.get("/rm-plans", name="manufacturing_rm_plans")
async def rm_plans(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    return templates.TemplateResponse("manufacturing/rm_plans.html", _ctx(
        request, "rm_plans", plans=store.get_rm_plans(cid), production_plans=store.get_plans(cid)))


@router.post("/rm-plans/new", name="manufacturing_rm_plan_new")
async def rm_plan_new(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    pid = store.create_rm_plan(cid, await _form(request), _actor(request))
    if pid:
        flash(request, "Raw material plan created — explode it from the production plan", "success")
        return _redirect(f"{BASE}/rm-plans/{pid}")
    flash(request, "Failed to create raw material plan", "error")
    return _redirect(f"{BASE}/rm-plans")


@router.get("/rm-plans/{plan_id}", name="manufacturing_rm_plan_detail")
async def rm_plan_detail(plan_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    p = store.get_rm_plan(plan_id, cid)
    if not p:
        flash(request, "Raw material plan not found", "error")
        return _redirect(f"{BASE}/rm-plans")
    costs = store.material_cost_map(cid)
    total = sum((D(l["to_procure_qty"]) * costs.get(l["material_code"], Decimal(0)) for l in p["lines"]), Decimal(0))
    return templates.TemplateResponse("manufacturing/rm_plan_detail.html", _ctx(
        request, "rm_plans", plan=p, costs=costs, estimated_value=total))


@router.post("/rm-plans/{plan_id}/explode", name="manufacturing_rm_plan_explode")
async def rm_plan_explode(plan_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    n = store.explode_rm_plan(plan_id, cid)
    if n < 0:
        flash(request, "Could not explode — plan must be in draft status", "error")
    else:
        flash(request, f"Exploded {n} material line(s) from the production plan × active BOMs", "success")
    return _redirect(f"{BASE}/rm-plans/{plan_id}")


@router.post("/rm-plans/{plan_id}/lines/add", name="manufacturing_rm_line_add")
async def rm_line_add(plan_id: str, request: Request, user=Depends(login_required)):
    _ok(request, store.add_rm_line(plan_id, await _form(request)), "Material line added", "Material code is required")
    return _redirect(f"{BASE}/rm-plans/{plan_id}")


@router.post("/rm-plans/{plan_id}/lines/{line_id}/delete", name="manufacturing_rm_line_delete")
async def rm_line_delete(plan_id: str, line_id: str, request: Request, user=Depends(login_required)):
    _ok(request, store.delete_rm_line(line_id, plan_id), "Line removed", "Failed to remove line")
    return _redirect(f"{BASE}/rm-plans/{plan_id}")


@router.post("/rm-plans/{plan_id}/submit", name="manufacturing_rm_plan_submit")
async def rm_plan_submit(plan_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    res = store.submit_rm_plan(plan_id, cid, _actor(request))
    if not res["ok"]:
        flash(request, res.get("error") or "Submission failed", "error")
    else:
        msg = f"Submitted to Property Administration — {res['store_reqs']} store requisition(s)"
        msg += f", purchase requisition {res['purchase_req'][:8]}…" if res.get("purchase_req") else ""
        if res.get("manual"):
            msg += " — manual step required for: " + ", ".join(res["manual"])
        msg += " — awaiting approval" if res.get("approval") else " — no approval workflow configured, plan auto-approved"
        flash(request, msg, "success")
    return _redirect(f"{BASE}/rm-plans/{plan_id}")


@router.post("/rm-plans/{plan_id}/approve", name="manufacturing_rm_plan_approve")
async def rm_plan_approve(plan_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    _ok(request, store.approve_rm_plan(plan_id, cid, _actor(request)), "Raw material plan approved", "Failed to approve")
    return _redirect(f"{BASE}/rm-plans/{plan_id}")


# ── production orders ─────────────────────────────────────────────
@router.get("/orders", name="manufacturing_orders")
async def orders(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    qp = request.query_params
    return templates.TemplateResponse("manufacturing/orders.html", _ctx(
        request, "orders", orders=store.get_orders(cid, qp.get("status") or None, qp.get("type") or None,
                                                   (qp.get("q") or "").strip() or None, qp.get("past_due") == "1"),
        status_filter=qp.get("status") or "", type_filter=qp.get("type") or "", search=(qp.get("q") or "").strip(),
        past_due_filter=qp.get("past_due") == "1"))


@router.get("/orders/new", name="manufacturing_order_new_get")
async def order_new_get(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    qp = request.query_params
    order = {"order_type": qp.get("type") or "make_to_stock", "source_ref": qp.get("source_ref") or "",
             "customer_name": qp.get("customer_name") or "", "priority": "normal"}
    return templates.TemplateResponse("manufacturing/order_form.html", _ctx(
        request, "orders", order=order, is_edit=False, products=store.get_products(cid, active_only=True), plants=store.get_plants(cid),
        tds_list=[], boms=[], routings=[]))


@router.post("/orders/new", name="manufacturing_order_new_post")
async def order_new_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    o = store.create_order(cid, await _form(request), _actor(request))
    if o:
        flash(request, f"Production order {o['order_no']} planned", "success")
        return _redirect(f"{BASE}/orders/{o['id']}")
    flash(request, "Failed to create order — product and a positive quantity are required", "error")
    return _redirect(f"{BASE}/orders/new")


def _order_or_redirect(request: Request, order_id: str):
    cid = current_company(request)
    o = store.get_order(order_id, cid)
    if not o:
        flash(request, "Production order not found", "error")
    return cid, o


@router.get("/orders/{order_id}", name="manufacturing_order_detail")
async def order_detail(order_id: str, request: Request, user=Depends(login_required)):
    cid, o = _order_or_redirect(request, order_id)
    if not o:
        return _redirect(f"{BASE}/orders")
    det = store.order_details(order_id)
    var_abs, var_pct = variance(o["planned_cost"], o["actual_cost"])
    tot_in = sum((D(l["input_qty"]) for l in det["logs"]), Decimal(0))
    tot_out = sum((D(l["output_qty"]) for l in det["logs"]), Decimal(0))
    return templates.TemplateResponse("manufacturing/order_detail.html", _ctx(
        request, "orders", order=o, cost_variance=var_abs, cost_variance_pct=var_pct, yield_pct=mfg.yield_pct(tot_out, tot_in),
        machines=store.get_machines(cid), shifts=store.get_shifts(cid, active_only=True), work_centers=store.get_work_centers(cid),
        scrap_types=store.get_scrap_types(cid), downtime_categories=store.get_downtime_categories(cid),
        approval=_approval_status(cid, order_id), **det))


def _approval_status(cid: str, order_id: str):
    try:
        from approval_hooks import approval_status
        return approval_status("production_order", order_id, cid)
    except Exception:
        return None


@router.get("/orders/{order_id}/edit", name="manufacturing_order_edit_get")
async def order_edit_get(order_id: str, request: Request, user=Depends(login_required)):
    cid, o = _order_or_redirect(request, order_id)
    if not o:
        return _redirect(f"{BASE}/orders")
    return templates.TemplateResponse("manufacturing/order_form.html", _ctx(
        request, "orders", order=o, is_edit=True, products=store.get_products(cid), plants=store.get_plants(cid),
        tds_list=store.get_tds_list(o["product_id"]), boms=store.get_boms(o["product_id"]), routings=store.get_routings(o["product_id"])))


@router.post("/orders/{order_id}/edit", name="manufacturing_order_edit_post")
async def order_edit_post(order_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    _ok(request, store.update_order(order_id, cid, await _form(request)), "Order updated", "Failed to update — only planned/released orders can be edited")
    return _redirect(f"{BASE}/orders/{order_id}")


@router.post("/orders/{order_id}/release", name="manufacturing_order_release")
async def order_release(order_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    res = store.request_release(order_id, cid, _actor(request))
    if not res["ok"]:
        flash(request, res.get("error") or "Release failed", "error")
    elif res.get("pending"):
        flash(request, "Release submitted to the approval workflow — the order is released automatically on approval", "success")
    else:
        flash(request, "Order released — operations and materials generated", "success")
    return _redirect(f"{BASE}/orders/{order_id}")


@router.post("/orders/{order_id}/start", name="manufacturing_order_start")
async def order_start(order_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    _ok(request, store.start_order(order_id, cid, _actor(request)), "Production started", "Order must be released first")
    return _redirect(f"{BASE}/orders/{order_id}")


@router.post("/orders/{order_id}/confirm", name="manufacturing_order_confirm")
async def order_confirm(order_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    res = store.confirm_order(order_id, cid, await _form(request), _actor(request))
    if res["ok"]:
        msg = f"Order confirmed — FG transfer recorded (GL: {res['gl_status']}"
        msg += ", inventory updated)" if res.get("inventory") else ", inventory not linked)"
        flash(request, msg, "success")
    else:
        flash(request, res.get("error") or "Confirmation failed", "error")
    return _redirect(f"{BASE}/orders/{order_id}")


@router.post("/orders/{order_id}/close", name="manufacturing_order_close")
async def order_close(order_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    _ok(request, store.close_order(order_id, cid, _actor(request)), "Order closed", "Only confirmed orders can be closed")
    return _redirect(f"{BASE}/orders/{order_id}")


@router.post("/orders/{order_id}/cancel", name="manufacturing_order_cancel")
async def order_cancel(order_id: str, request: Request, user=Depends(manager_required)):
    cid = current_company(request)
    data = await _form(request)
    _ok(request, store.cancel_order(order_id, cid, _actor(request), data.get("reason", "")), "Order cancelled", "Only planned/released orders can be cancelled")
    return _redirect(f"{BASE}/orders/{order_id}")


@router.post("/orders/{order_id}/operations/{op_id}", name="manufacturing_order_operation_update")
async def order_operation_update(order_id: str, op_id: str, request: Request, user=Depends(login_required)):
    data = await _form(request)
    _ok(request, store.update_operation(op_id, order_id, data, params_from_form(data, PROCESS_FIELDS)), "Operation updated", "Failed to update operation")
    return _redirect(f"{BASE}/orders/{order_id}#operations")


@router.post("/orders/{order_id}/log", name="manufacturing_order_log")
async def order_log(order_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    _ok(request, store.add_production_log(cid, order_id, await _form(request)), "Production log recorded", "Failed to record log")
    return _redirect(f"{BASE}/orders/{order_id}#logs")


@router.post("/orders/{order_id}/issue", name="manufacturing_order_issue")
async def order_issue(order_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    res = store.add_material_issue(cid, order_id, await _form(request), _actor(request))
    if res["ok"]:
        flash(request, "Material movement recorded" + (f" (GL: {res['gl_status']})" if res.get("gl_status") else ""), "success")
    else:
        flash(request, res.get("error") or "Failed", "error")
    return _redirect(f"{BASE}/orders/{order_id}#materials")


@router.post("/orders/{order_id}/labor", name="manufacturing_order_labor")
async def order_labor(order_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    data["order_id"] = order_id
    ok = store.add_labor_log(cid, data)
    if ok:
        store.recompute_actual_cost(order_id, cid)
    _ok(request, ok, "Labour time recorded", "Failed to record labour")
    return _redirect(f"{BASE}/orders/{order_id}#labor")


@router.post("/orders/{order_id}/downtime", name="manufacturing_order_downtime")
async def order_downtime(order_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    data["order_id"] = order_id
    _ok(request, store.add_downtime(cid, data, _actor(request)), "Downtime recorded", "Failed to record downtime")
    return _redirect(f"{BASE}/orders/{order_id}#downtime")


@router.get("/orders/{order_id}/bom-print", name="manufacturing_order_bom_print")
async def order_bom_print(order_id: str, request: Request, user=Depends(login_required)):
    cid, o = _order_or_redirect(request, order_id)
    if not o:
        return _redirect(f"{BASE}/orders")
    det = store.order_details(order_id)
    return templates.TemplateResponse("manufacturing/bom_print.html", _ctx(
        request, "orders", order=o, bom=store.get_bom(o["bom_id"]) if o.get("bom_id") else None,
        tds=store.get_tds(o["tds_id"], cid) if o.get("tds_id") else None, operations=det["operations"], materials=det["materials"]))


# ── shop-floor lists ──────────────────────────────────────────────
@router.get("/logs", name="manufacturing_logs")
async def logs(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    fc = _filter_ctx(request, cid)
    return templates.TemplateResponse("manufacturing/logs.html", _ctx(
        request, "logs", logs=store.get_production_logs(cid, fc["filters"]),
        open_orders=store.get_orders(cid, limit=200), shifts=store.get_shifts(cid, active_only=True),
        scrap_types=store.get_scrap_types(cid), **fc))


@router.post("/logs/new", name="manufacturing_log_new")
async def log_new(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    data = await _form(request)
    if not data.get("order_id"):
        flash(request, "Select a production order", "error")
    else:
        _ok(request, store.add_production_log(cid, data["order_id"], data), "Production log recorded", "Failed to record log")
    return _redirect(f"{BASE}/logs")


@router.get("/downtime", name="manufacturing_downtime")
async def downtime(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    store.seed_defaults(cid)
    fc = _filter_ctx(request, cid)
    return templates.TemplateResponse("manufacturing/downtime.html", _ctx(
        request, "downtime", downtime=store.get_downtime(cid, fc["filters"]), categories=store.get_downtime_categories(cid),
        shifts=store.get_shifts(cid, active_only=True), open_orders=store.get_orders(cid, limit=200), **fc))


@router.post("/downtime/new", name="manufacturing_downtime_new")
async def downtime_new(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    _ok(request, store.add_downtime(cid, await _form(request), _actor(request)), "Downtime recorded", "Failed to record downtime")
    return _redirect(f"{BASE}/downtime")


@router.post("/labor/new", name="manufacturing_labor_new")
async def labor_new(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    _ok(request, store.add_labor_log(cid, await _form(request)), "Labour time recorded", "Failed to record labour")
    return _redirect(f"{BASE}/logs")


# ── reports ───────────────────────────────────────────────────────
REPORTS = {
    "machine-performance": "Performance per machine",
    "rm-converted": "Raw material converted to finished goods",
    "fg-delivered": "Finished goods delivered to store",
    "rm-status": "Raw material status per order",
    "scrap": "Scrap generated",
    "utilisation": "Machine utilisation and downtime",
    "plan-vs-actual": "Plan vs actual, efficiency, yield and wastage",
    "cycle": "Production order cycle and cost variance",
    "journals": "Consumption and output journals",
}


def _n(v):
    return float(D(v))


def _report_data(cid: str, name: str, f: dict):
    """→ (columns [(key,label,kind)], rows, totals dict). kind: text|num|int|date|pct"""
    if name == "machine-performance":
        rows = store.report_machine_performance(cid, f)
        cols = [("machine_code", "Machine", "text"), ("product", "Product", "text"), ("order_no", "Order", "text"),
                ("diameter", "Diameter (mm)", "text"), ("size_mm2", "Cross-section (mm²)", "num"), ("input_kg", "Input", "num"),
                ("output_kg", "Output", "num"), ("ratio_pct", "Output/input %", "pct"), ("plan_qty", "Plan", "num"),
                ("plan_vs_actual", "Actual − plan", "num"), ("length_m", "Length (m)", "num"), ("from_date", "From", "date"), ("to_date", "To", "date")]
    elif name == "rm-converted":
        rows = store.report_rm_converted(cid, f)
        cols = [("order_no", "Production order", "text"), ("size_mm2", "Size (mm²)", "num"), ("description", "Description", "text"),
                ("color", "Colour", "text"), ("output_rolls", "Output rolls", "int"), ("under_length_rolls", "Under-length rolls", "int"),
                ("length_m", "Length (m)", "num"), ("weight_kg", "Weight (kg)", "num"), ("input_kg", "Input (kg)", "num")]
    elif name == "fg-delivered":
        rows = store.report_fg_delivered(cid, f)
        cols = [("transferred_at", "Date", "date"), ("order_no", "Order", "text"), ("size_mm2", "Size (mm²)", "num"),
                ("description", "Description", "text"), ("color", "Colour", "text"), ("rolls", "Rolls", "int"), ("length_m", "Length (m)", "num"),
                ("under_length_rolls", "Under-length", "int"), ("weight_kg", "Kg", "num"), ("ref_no", "Ref no", "text"), ("to_store", "To store", "text")]
    elif name == "rm-status":
        mats, rows = store.report_rm_status(cid, f)
        cols = [("order_no", "Order", "text"), ("size_mm2", "Size (mm²)", "num"), ("product_name", "Product", "text")]
        for m in mats:
            cols += [(f"{m}__issued", f"{m} issued", "num"), (f"{m}__consumed", f"{m} consumed", "num"), (f"{m}__returned", f"{m} returned", "num")]
        cols += [("total_issued", "Total issued", "num"), ("total_consumed", "Total consumed", "num"), ("total_returned", "Total returned", "num"),
                 ("drum_refs", "Drum ref no", "text")]
    elif name == "scrap":
        types, rows = store.report_scrap(cid, f)
        cols = [("order_no", "Order", "text"), ("size_mm2", "Size (mm²)", "num"), ("description", "Description", "text"), ("color", "Colour", "text")]
        cols += [(t, f"{t} (kg)", "num") for t in types] + [("total_kg", "Total (kg)", "num")]
    elif name == "utilisation":
        rows = store.report_utilisation(cid, f, f.get("group") or "day")
        cats = sorted({c for r in rows for c in r["downtime_by_category"]})
        for r in rows:
            for c in cats:
                r[f"dt_{c}"] = r["downtime_by_category"].get(c, Decimal(0))
        cols = [("period", "Period", "date"), ("machine_code", "Machine", "text"), ("work_center", "Work centre", "text"),
                ("available_hours", "Available h", "num"), ("run_hours", "Run h", "num"), ("downtime_hours", "Downtime h", "num")]
        cols += [(f"dt_{c}", f"{c} h", "num") for c in cats]
        cols += [("utilisation_pct", "Utilisation %", "pct"), ("output_qty", "Output", "num"), ("yield_pct", "Yield %", "pct")]
    elif name == "plan-vs-actual":
        rows = store.report_plan_vs_actual(cid, f)
        cols = [("work_center", "Line / work centre", "text"), ("machine_code", "Machine", "text"), ("product_code", "SKU", "text"),
                ("product", "Variant", "text"), ("shift", "Shift", "text"), ("planned_qty", "Planned", "num"), ("output_qty", "Actual", "num"),
                ("efficiency_pct", "Efficiency %", "pct"), ("input_qty", "Input", "num"), ("yield_pct", "Yield %", "pct"),
                ("scrap_qty", "Scrap", "num"), ("wastage_pct", "Wastage %", "pct"), ("rework_qty", "Rework", "num")]
    elif name == "cycle":
        rows = store.report_cycle(cid, f)
        cols = [("order_no", "Order", "text"), ("product_code", "Product", "text"), ("order_type", "Type", "text"), ("status", "Status", "text"),
                ("created_at", "Planned", "date"), ("released_at", "Released", "date"), ("confirmed_at", "Confirmed", "date"), ("closed_at", "Closed", "date"),
                ("days_to_release", "Days to release", "num"), ("days_in_production", "Days in production", "num"), ("days_total", "Total days", "num"),
                ("qty_ordered", "Ordered", "num"), ("qty_produced", "Produced", "num"), ("planned_cost", "Planned cost", "num"),
                ("actual_cost", "Actual cost", "num"), ("cost_variance", "Variance", "num"), ("cost_variance_pct", "Variance %", "pct")]
    elif name == "journals":
        rows = store.report_journals(cid, f)
        cols = [("posted_at", "Date", "date"), ("order_no", "Order", "text"), ("kind", "Journal", "text"), ("item", "Item", "text"),
                ("qty", "Qty", "num"), ("unit", "Unit", "text"), ("amount", "Amount", "num"), ("gl_status", "GL status", "text"), ("gl_entry_id", "GL entry", "text")]
    else:
        return None, [], {}
    totals = {k: sum((D(r.get(k)) for r in rows), Decimal(0)) for k, _, kind in cols if kind in ("num", "int")}
    return cols, rows, totals


@router.get("/reports", name="manufacturing_reports")
async def reports_index(request: Request, user=Depends(login_required)):
    return templates.TemplateResponse("manufacturing/reports_index.html", _ctx(request, "reports", reports=REPORTS))


@router.get("/reports/{name}", name="manufacturing_report")
async def report(name: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    if name not in REPORTS:
        flash(request, "Unknown report", "error")
        return _redirect(f"{BASE}/reports")
    fc = _filter_ctx(request, cid)
    cols, rows, totals = _report_data(cid, name, fc["filters"])
    return templates.TemplateResponse("manufacturing/report.html", _ctx(
        request, "reports", report_name=name, report_title=REPORTS[name], columns=cols, rows=rows, totals=totals,
        generated=date.today(), **fc))


@router.get("/reports/{name}/export", name="manufacturing_report_export")
async def report_export(name: str, request: Request, user=Depends(login_required)):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    cid = current_company(request)
    if name not in REPORTS:
        return _redirect(f"{BASE}/reports")
    f = _filters(request)
    cols, rows, totals = _report_data(cid, name, f)
    wb = Workbook()
    ws = wb.active
    ws.title = name[:30]
    ws.append([f"{REPORTS[name]} — company {cid}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Period {f.get('date_from') or '…'} to {f.get('date_to') or '…'} — generated {date.today().isoformat()}"])
    ws.append([])
    ws.append([label for _, label, _ in cols])
    for c in ws[4]:
        c.font = Font(bold=True)
    for r in rows:
        line = []
        for key, _, kind in cols:
            v = r.get(key)
            if kind in ("num", "pct"):
                line.append(_n(v))
            elif kind == "int":
                line.append(int(v or 0))
            elif kind == "date":
                line.append(v.isoformat() if hasattr(v, "isoformat") else (v or ""))
            else:
                line.append("" if v is None else str(v))
        ws.append(line)
    ws.append([])
    ws.append([("TOTAL" if i == 0 else (_n(totals[k]) if k in totals else "")) for i, (k, _, _) in enumerate(cols)])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for i, (_, label, _) in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = max(12, min(40, len(label) + 4))
    ws.append([]); ws.append([f"EBMS Manufacturing — {REPORTS[name]} — page footer"])
    fd, path = tempfile.mkstemp(suffix=".xlsx"); os.close(fd)
    wb.save(path)
    return FileResponse(path, filename=f"mfg_{name}_{date.today().isoformat()}.xlsx", media_type=_XLSX)
