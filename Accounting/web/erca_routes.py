"""
ERCA Tax Outputs routes — VAT return, withholding, gapless e-invoice numbering.
Prefix /erca, route names erca_*.
"""
from __future__ import annotations

import io
import json
import logging
from datetime import date, datetime
from decimal import Decimal

from fastapi import APIRouter, Depends, Request
from fastapi.responses import JSONResponse, RedirectResponse, Response

from deps import current_company, flash, login_required, template_context
from template_engine import templates
from erca_data_store import erca_store
import erca_forms as forms

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/erca", tags=["erca"])

_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ── helpers ─────────────────────────────────────────────────────────────────
def _actor(request: Request) -> str:
    return request.session.get("username", "") or ""


def _ctx(request: Request, **extra) -> dict:
    ctx = template_context(request)
    ctx.update(profile=erca_store.get_profile(current_company(request)),
               months=[(i, forms.period_bounds(2000, i)["label"].split()[0]) for i in range(1, 13)],
               today=date.today(), **extra)
    return ctx


def _ym(request: Request):
    """(year, month) from ?year&month, defaulting to the period currently due."""
    q = request.query_params
    y, m = forms.current_period()
    try:
        y = int(q.get("year") or y)
        m = int(q.get("month") or m)
        if not 1 <= m <= 12:
            raise ValueError
    except ValueError:
        y, m = forms.current_period()
    return y, m


def _date(v):
    try:
        return datetime.strptime(str(v)[:10], "%Y-%m-%d").date() if v else None
    except ValueError:
        return None


def _pdf(content: bytes, filename: str) -> Response:
    return Response(content=content, media_type="application/pdf",
                    headers={"Content-Disposition": f'inline; filename="{filename}"'})


def _xlsx(wb, filename: str) -> Response:
    buf = io.BytesIO()
    wb.save(buf)
    return Response(content=buf.getvalue(), media_type=_XLSX,
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def _sheet(wb, title: str, header: list, rows: list, first: bool = False):
    from openpyxl.styles import Font
    ws = wb.active if first else wb.create_sheet()
    ws.title = title[:31]
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([float(v) if isinstance(v, Decimal) else (v.isoformat() if hasattr(v, "isoformat") else v)
                   for v in r])
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
    return ws


def _json_default(o):
    if isinstance(o, Decimal):
        return str(o)
    if hasattr(o, "isoformat"):
        return o.isoformat()
    return str(o)


# ── dashboard ───────────────────────────────────────────────────────────────
@router.get("/", name="erca_dashboard")
async def dashboard(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    erca_store.ensure_default_series(cid)
    data = erca_store.dashboard(cid)
    return templates.TemplateResponse("erca/dashboard.html", _ctx(request, **data))


# ── company profile ─────────────────────────────────────────────────────────
@router.get("/profile", name="erca_profile")
async def profile_get(request: Request, user=Depends(login_required)):
    ctx = _ctx(request)
    tin_ok, tin_msg = forms.validate_tin(ctx["profile"].get("tin")) if ctx["profile"].get("tin") else (True, "")
    ctx.update(tin_ok=tin_ok, tin_msg=tin_msg, categories=forms.TAXPAYER_CATEGORIES)
    return templates.TemplateResponse("erca/profile.html", ctx)


@router.post("/profile", name="erca_profile_post")
async def profile_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    form = await request.form()
    data = {k: v for k, v in form.items()}
    ok, msg = forms.validate_tin(data.get("tin"))
    if data.get("tin") and not ok:
        flash(request, f"TIN rejected: {msg}", "error")
        return RedirectResponse("/erca/profile", status_code=303)
    if erca_store.save_profile(cid, data):
        flash(request, "Taxpayer profile saved" + (f" ({msg})" if ok and msg != "ok" and data.get("tin") else ""), "success")
    else:
        flash(request, "Failed to save profile", "error")
    return RedirectResponse("/erca/profile", status_code=303)


# ── VAT returns ─────────────────────────────────────────────────────────────
@router.get("/vat-returns", name="erca_vat_returns")
async def vat_returns(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    rets = erca_store.get_vat_returns(cid)
    for r in rets:
        r["period"] = forms.period_bounds(r["period_year"], r["period_month"])
    filed = {(r["period_year"], r["period_month"]) for r in rets if r["status"] == "filed"}
    y, m = forms.current_period()
    ctx = _ctx(request, returns=rets, unfiled=forms.unfiled_periods(filed, lookback=12),
               default_year=y, default_month=m, years=list(range(date.today().year - 3, date.today().year + 1)))
    return templates.TemplateResponse("erca/vat_return.html", ctx)


@router.post("/vat-returns/compute", name="erca_vat_return_compute")
async def vat_return_compute(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    form = await request.form()
    try:
        y, m = int(form.get("year")), int(form.get("month"))
        forms.period_bounds(y, m)
    except (TypeError, ValueError):
        flash(request, "Invalid period", "error")
        return RedirectResponse("/erca/vat-returns", status_code=303)
    ret = erca_store.compute_vat_return(cid, y, m)
    if not ret:
        flash(request, "Failed to compute VAT return", "error")
        return RedirectResponse("/erca/vat-returns", status_code=303)
    if ret["status"] != "draft":
        flash(request, f"Period already {ret['status']} — not recomputed", "info")
    else:
        flash(request, f"VAT return for {forms.period_bounds(y, m)['label']} computed", "success")
    return RedirectResponse(f"/erca/vat-returns/{ret['id']}", status_code=303)


def _load_return(request: Request, return_id: str):
    cid = current_company(request)
    ret = erca_store.get_vat_return(cid, return_id)
    if ret:
        ret["period"] = forms.period_bounds(ret["period_year"], ret["period_month"])
        lines = ret.get("lines") or {}
        ret["boxes"] = [{"box": b, "label": en, "label_am": am, "key": k, "amount": forms.money(lines.get(k, 0))}
                        for b, en, am, k in forms.VAT_BOXES]
    return cid, ret


@router.get("/vat-returns/{return_id}", name="erca_vat_return_detail")
async def vat_return_detail(return_id: str, request: Request, user=Depends(login_required)):
    cid, ret = _load_return(request, return_id)
    if not ret:
        flash(request, "VAT return not found", "error")
        return RedirectResponse("/erca/vat-returns", status_code=303)
    return templates.TemplateResponse("erca/vat_return_detail.html", _ctx(request, ret=ret))


@router.post("/vat-returns/{return_id}/recompute", name="erca_vat_return_recompute")
async def vat_return_recompute(return_id: str, request: Request, user=Depends(login_required)):
    cid, ret = _load_return(request, return_id)
    if not ret:
        flash(request, "VAT return not found", "error")
        return RedirectResponse("/erca/vat-returns", status_code=303)
    if ret["status"] != "draft":
        flash(request, "Only draft returns can be recomputed — reopen it first", "error")
    elif erca_store.compute_vat_return(cid, ret["period_year"], ret["period_month"]):
        flash(request, "Return recomputed from current VAT records", "success")
    else:
        flash(request, "Recompute failed", "error")
    return RedirectResponse(f"/erca/vat-returns/{return_id}", status_code=303)


@router.post("/vat-returns/{return_id}/finalize", name="erca_vat_return_finalize")
async def vat_return_finalize(return_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    if erca_store.finalize_vat_return(cid, return_id, _actor(request)):
        flash(request, "VAT return finalized — figures are now locked", "success")
    else:
        flash(request, "Could not finalize (already finalized/filed?)", "error")
    return RedirectResponse(f"/erca/vat-returns/{return_id}", status_code=303)


@router.post("/vat-returns/{return_id}/reopen", name="erca_vat_return_reopen")
async def vat_return_reopen(return_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    if erca_store.reopen_vat_return(cid, return_id):
        flash(request, "VAT return reopened as draft", "success")
    else:
        flash(request, "Only finalized (not yet filed) returns can be reopened", "error")
    return RedirectResponse(f"/erca/vat-returns/{return_id}", status_code=303)


@router.post("/vat-returns/{return_id}/file", name="erca_vat_return_file")
async def vat_return_file(return_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    form = await request.form()
    if erca_store.mark_vat_return_filed(cid, return_id, _date(form.get("filed_on")), form.get("erca_reference") or ""):
        flash(request, "VAT return marked as filed with ERCA", "success")
    else:
        flash(request, "Finalize the return before marking it filed", "error")
    return RedirectResponse(f"/erca/vat-returns/{return_id}", status_code=303)


@router.get("/vat-returns/{return_id}/pdf", name="erca_vat_return_pdf")
async def vat_return_pdf(return_id: str, request: Request, user=Depends(login_required)):
    cid, ret = _load_return(request, return_id)
    if not ret:
        flash(request, "VAT return not found", "error")
        return RedirectResponse("/erca/vat-returns", status_code=303)
    pdf = forms.build_vat_return_pdf(erca_store.get_profile(cid), ret)
    return _pdf(pdf, f"vat_return_{ret['period_year']}_{ret['period_month']:02d}.pdf")


@router.get("/vat-returns/{return_id}/excel", name="erca_vat_return_excel")
async def vat_return_excel(return_id: str, request: Request, user=Depends(login_required)):
    cid, ret = _load_return(request, return_id)
    if not ret:
        flash(request, "VAT return not found", "error")
        return RedirectResponse("/erca/vat-returns", status_code=303)
    from openpyxl import Workbook
    wb = Workbook()
    p = ret["period"]
    prof = erca_store.get_profile(cid)
    _sheet(wb, "VAT Return", ["Box", "Description", "Amharic", "Amount (ETB)"],
           [(b["box"], b["label"], b["label_am"], b["amount"]) for b in ret["boxes"]], first=True)
    _sheet(wb, "Header", ["Field", "Value"], [
        ("Taxpayer", prof.get("taxpayer_name")), ("TIN", prof.get("tin")), ("VAT reg no", prof.get("vat_reg_no")),
        ("Tax centre", prof.get("tax_centre")), ("Category", prof.get("category")),
        ("Period", p["label"]), ("Ethiopian period", p["ethiopian_label"]), ("Filing deadline", p["deadline"]),
        ("Status", ret["status"]), ("Computed", ret.get("computed_at")), ("ERCA reference", ret.get("erca_reference") or ""),
    ])
    return _xlsx(wb, f"vat_return_{ret['period_year']}_{ret['period_month']:02d}.xlsx")


# ── withholding register ────────────────────────────────────────────────────
@router.get("/withholding", name="erca_withholding")
async def withholding_list(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    q = request.query_params
    y, m = (None, None)
    if q.get("year") or q.get("month"):
        y, m = _ym(request)
    entries = erca_store.get_withholding(cid, y, m)
    ctx = _ctx(request, entries=entries, summary=forms.withholding_summary(entries),
               year=y, month=m, years=list(range(date.today().year - 3, date.today().year + 1)))
    return templates.TemplateResponse("erca/withholding.html", ctx)


@router.get("/withholding/new", name="erca_withholding_new")
async def withholding_new_get(request: Request, user=Depends(login_required)):
    ctx = _ctx(request, entry={}, transaction_types=forms.TRANSACTION_TYPES, rules=_rules())
    return templates.TemplateResponse("erca/withholding_form.html", ctx)


def _rules() -> dict:
    return {"rate_tin": forms.WHT_RATE_TIN, "rate_no_tin": forms.WHT_RATE_NO_TIN, "rate_import": forms.WHT_RATE_IMPORT,
            "threshold_goods": forms.WHT_THRESHOLD_GOODS, "threshold_services": forms.WHT_THRESHOLD_SERVICES}


@router.post("/withholding/new", name="erca_withholding_new_post")
async def withholding_new_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    erca_store.ensure_default_series(cid)
    form = await request.form()
    data = {k: v for k, v in form.items()}
    data["created_by"] = _actor(request)
    if data.get("payee_tin") and not forms.validate_tin(data["payee_tin"])[0]:
        flash(request, f"Payee TIN invalid: {forms.validate_tin(data['payee_tin'])[1]}", "error")
        return RedirectResponse("/erca/withholding/new", status_code=303)
    e = erca_store.add_withholding(cid, data)
    if e:
        flash(request, f"Withholding recorded — ETB {forms.money(e['withheld_amount']):,.2f} at "
                       f"{forms.money(e['withheld_rate']) * 100:.0f}%" + (f", receipt {e['receipt_no']}" if e.get("receipt_no") else ""),
              "success")
        return RedirectResponse("/erca/withholding", status_code=303)
    flash(request, "Failed to record withholding", "error")
    return RedirectResponse("/erca/withholding/new", status_code=303)


@router.get("/withholding/return", name="erca_withholding_return")
async def withholding_return(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    y, m = _ym(request)
    entries = erca_store.get_withholding(cid, y, m)
    ctx = _ctx(request, entries=entries, summary=forms.withholding_summary(entries), year=y, month=m,
               period=forms.period_bounds(y, m), years=list(range(date.today().year - 3, date.today().year + 1)))
    return templates.TemplateResponse("erca/withholding_return.html", ctx)


@router.get("/withholding/return/pdf", name="erca_withholding_return_pdf")
async def withholding_return_pdf(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    y, m = _ym(request)
    pdf = forms.build_withholding_return_pdf(erca_store.get_profile(cid), y, m, erca_store.get_withholding(cid, y, m))
    return _pdf(pdf, f"withholding_return_{y}_{m:02d}.pdf")


@router.get("/withholding/return/excel", name="erca_withholding_return_excel")
async def withholding_return_excel(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    y, m = _ym(request)
    entries = erca_store.get_withholding(cid, y, m)
    s = forms.withholding_summary(entries)
    from openpyxl import Workbook
    wb = Workbook()
    _sheet(wb, "Withholding", ["Receipt", "Payment date", "Payee", "TIN", "Has TIN", "Type", "Gross", "Rate", "Withheld", "Invoice ref"],
           [(e.get("receipt_no"), e.get("payment_date"), e.get("payee_name"), e.get("payee_tin"), e.get("has_tin"),
             e.get("transaction_type"), e.get("gross_amount"), e.get("withheld_rate"), e.get("withheld_amount"), e.get("invoice_ref"))
            for e in entries] + [("TOTAL", "", "", "", "", "", s["total_gross"], "", s["total_withheld"], "")], first=True)
    _sheet(wb, "Summary", ["Rate", "Count", "Gross", "Withheld"],
           [(k, b["count"], b["gross"], b["withheld"]) for k, b in s["buckets"].items()])
    return _xlsx(wb, f"withholding_return_{y}_{m:02d}.xlsx")


@router.get("/withholding/{wid}/edit", name="erca_withholding_edit")
async def withholding_edit_get(wid: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    e = erca_store.get_withholding_entry(cid, wid)
    if not e:
        flash(request, "Withholding entry not found", "error")
        return RedirectResponse("/erca/withholding", status_code=303)
    ctx = _ctx(request, entry=e, transaction_types=forms.TRANSACTION_TYPES, rules=_rules())
    return templates.TemplateResponse("erca/withholding_form.html", ctx)


@router.post("/withholding/{wid}/edit", name="erca_withholding_edit_post")
async def withholding_edit_post(wid: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    form = await request.form()
    data = {k: v for k, v in form.items()}
    if erca_store.update_withholding(cid, wid, data):
        flash(request, "Withholding entry updated", "success")
    else:
        flash(request, "Failed to update entry", "error")
    return RedirectResponse("/erca/withholding", status_code=303)


@router.post("/withholding/{wid}/delete", name="erca_withholding_delete")
async def withholding_delete(wid: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    if erca_store.delete_withholding(cid, wid, _actor(request)):
        flash(request, "Entry deleted — its receipt number is retained in the gaps audit", "success")
    else:
        flash(request, "Failed to delete entry", "error")
    return RedirectResponse("/erca/withholding", status_code=303)


@router.get("/withholding/{wid}/receipt.pdf", name="erca_withholding_receipt_pdf")
async def withholding_receipt_pdf(wid: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    e = erca_store.get_withholding_entry(cid, wid)
    if not e:
        flash(request, "Withholding entry not found", "error")
        return RedirectResponse("/erca/withholding", status_code=303)
    return _pdf(forms.build_withholding_receipt_pdf(erca_store.get_profile(cid), e),
                f"withholding_receipt_{e.get('receipt_no') or wid}.pdf")


# ── invoice series ──────────────────────────────────────────────────────────
@router.get("/series", name="erca_series")
async def series_list(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    erca_store.ensure_default_series(cid)
    ctx = _ctx(request, series=erca_store.get_series(cid), gaps=erca_store.get_gaps_audit(cid, 50))
    return templates.TemplateResponse("erca/invoice_numbers.html", ctx)


@router.get("/series/new", name="erca_series_new")
async def series_new_get(request: Request, user=Depends(login_required)):
    return templates.TemplateResponse("erca/invoice_series_form.html", _ctx(request, series={}, kinds=forms.INVOICE_KINDS))


@router.post("/series/new", name="erca_series_new_post")
async def series_new_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    form = await request.form()
    data = {k: v for k, v in form.items()}
    data.setdefault("is_active", "on")
    if erca_store.create_series(cid, data):
        flash(request, "Number series created", "success")
        return RedirectResponse("/erca/series", status_code=303)
    flash(request, "Failed to create series (code must be unique, kind valid)", "error")
    return RedirectResponse("/erca/series/new", status_code=303)


@router.get("/series/{series_id}/edit", name="erca_series_edit")
async def series_edit_get(series_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    s = erca_store.get_series_one(cid, series_id)
    if not s:
        flash(request, "Series not found", "error")
        return RedirectResponse("/erca/series", status_code=303)
    return templates.TemplateResponse("erca/invoice_series_form.html", _ctx(request, series=s, kinds=forms.INVOICE_KINDS))


@router.post("/series/{series_id}/edit", name="erca_series_edit_post")
async def series_edit_post(series_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    form = await request.form()
    data = {k: v for k, v in form.items()}
    if erca_store.update_series(cid, series_id, data):
        flash(request, "Series updated (next number can only move forward)", "success")
    else:
        flash(request, "Failed to update series", "error")
    return RedirectResponse("/erca/series", status_code=303)


@router.post("/series/{series_id}/toggle", name="erca_series_toggle")
async def series_toggle(series_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    if erca_store.toggle_series(cid, series_id):
        flash(request, "Series active flag toggled", "success")
    else:
        flash(request, "Failed to toggle series", "error")
    return RedirectResponse("/erca/series", status_code=303)


# ── invoices ────────────────────────────────────────────────────────────────
@router.get("/invoices", name="erca_invoices")
async def invoices_list(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    erca_store.ensure_default_series(cid)
    q = request.query_params
    f = {"kind": q.get("kind") or None, "status": q.get("status") or None,
         "start": _date(q.get("start")), "end": _date(q.get("end"))}
    invs = erca_store.get_invoices(cid, kind=f["kind"], status=f["status"], start=f["start"], end=f["end"])
    series = [s for s in erca_store.get_series(cid, active_only=True) if s["kind"] != "withholding_receipt"]
    ctx = _ctx(request, invoices=invs, filters=f, kinds=forms.INVOICE_KINDS, series=series, vat_rate=forms.VAT_RATE,
               mode="list", total=sum((forms.money(i["total"]) for i in invs if i["status"] == "issued"), Decimal("0")))
    return templates.TemplateResponse("erca/invoices.html", ctx)


@router.get("/invoices/new", name="erca_invoice_new")
async def invoice_new_get(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    erca_store.ensure_default_series(cid)
    # The "issue new document" form is the top card of invoices.html; mode='new'
    # renders it expanded with the recent list underneath.
    series = [s for s in erca_store.get_series(cid, active_only=True) if s["kind"] != "withholding_receipt"]
    invs = erca_store.get_invoices(cid, limit=10)
    ctx = _ctx(request, series=series, vat_rate=forms.VAT_RATE, kinds=forms.INVOICE_KINDS,
               invoices=invs, filters={"kind": None, "status": None, "start": None, "end": None},
               total=sum((forms.money(i["total"]) for i in invs if i["status"] == "issued"), Decimal("0")),
               mode="new")
    return templates.TemplateResponse("erca/invoices.html", ctx)


def _items_from_form(form) -> list:
    descs = form.getlist("item_description")
    qtys = form.getlist("item_qty")
    prices = form.getlist("item_unit_price")
    rates = form.getlist("item_vat_rate")
    items = []
    for i, d in enumerate(descs):
        items.append({"description": d,
                      "qty": qtys[i] if i < len(qtys) else 1,
                      "unit_price": prices[i] if i < len(prices) else 0,
                      "vat_rate": rates[i] if i < len(rates) else None})
    return items


@router.post("/invoices/new", name="erca_invoice_new_post")
async def invoice_new_post(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    form = await request.form()
    tin = form.get("customer_tin") or ""
    if tin and not forms.validate_tin(tin)[0]:
        flash(request, f"Customer TIN invalid: {forms.validate_tin(tin)[1]}", "error")
        return RedirectResponse("/erca/invoices/new", status_code=303)
    inv = erca_store.issue_invoice(
        cid, form.get("series_code") or "INV",
        customer_name=form.get("customer_name"), customer_tin=tin,
        items=_items_from_form(form), issued_at=form.get("issued_at") or None,
        linked_income_id=form.get("linked_income_id") or None,
        withholding_type=form.get("withholding_type") or "goods",
        withholding_expected=form.get("withholding_expected") or None,
        created_by=_actor(request))
    if inv:
        flash(request, f"{inv['kind'].replace('_', ' ').title()} {inv['number']} issued", "success")
        return RedirectResponse(f"/erca/invoices/{inv['id']}", status_code=303)
    flash(request, "Failed to issue — check the series is active and at least one line item is filled", "error")
    return RedirectResponse("/erca/invoices/new", status_code=303)


@router.get("/invoices/{invoice_id}", name="erca_invoice_detail")
async def invoice_detail(invoice_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    inv = erca_store.get_invoice(cid, invoice_id)
    if not inv:
        flash(request, "Invoice not found", "error")
        return RedirectResponse("/erca/invoices", status_code=303)
    return templates.TemplateResponse("erca/invoice_detail.html", _ctx(request, invoice=inv))


@router.get("/invoices/{invoice_id}/pdf", name="erca_invoice_pdf")
async def invoice_pdf(invoice_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    inv = erca_store.get_invoice(cid, invoice_id)
    if not inv:
        flash(request, "Invoice not found", "error")
        return RedirectResponse("/erca/invoices", status_code=303)
    return _pdf(forms.build_invoice_pdf(erca_store.get_profile(cid), inv, {"machine_id": inv.get("machine_id")}),
                f"{inv['number']}.pdf")


@router.post("/invoices/{invoice_id}/void", name="erca_invoice_void")
async def invoice_void(invoice_id: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    form = await request.form()
    if erca_store.void_invoice(cid, invoice_id, form.get("reason") or "", _actor(request)):
        flash(request, "Document voided — the number is retained and will never be reused", "success")
    else:
        flash(request, "Could not void (already voided?)", "error")
    return RedirectResponse(f"/erca/invoices/{invoice_id}", status_code=303)


# ── audit export ────────────────────────────────────────────────────────────
def _audit_range(request: Request):
    q = request.query_params
    efy = forms.ethiopian_fiscal_year(date.today())
    fy_start, fy_end = forms.fiscal_year_bounds(efy)
    return _date(q.get("start")) or fy_start, _date(q.get("end")) or min(fy_end, date.today())


@router.get("/audit", name="erca_audit")
async def audit_page(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    start, end = _audit_range(request)
    data = erca_store.audit_export(cid, start, end)
    ctx = _ctx(request, start=start, end=end, audit=data, fiscal_year=forms.ethiopian_fiscal_year(date.today()),
               backend=forms.pdf_backend())
    return templates.TemplateResponse("erca/audit_export.html", ctx)


@router.get("/audit/export.xlsx", name="erca_audit_excel")
async def audit_excel(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    start, end = _audit_range(request)
    data = erca_store.audit_export(cid, start, end)
    from openpyxl import Workbook
    wb = Workbook()
    _sheet(wb, "Invoices", ["Number", "Series", "Kind", "Issued", "Customer", "Customer TIN", "Subtotal", "VAT", "Total",
                            "Withholding expected", "Status", "Void reason", "Hash"],
           [(i["number"], i.get("series_code"), i["kind"], i["issued_at"], i["customer_name"], i["customer_tin"],
             i["subtotal"], i["vat_amount"], i["total"], i["withholding_expected"], i["status"], i.get("void_reason") or "", i["hash"])
            for i in data["invoices"]], first=True)
    _sheet(wb, "Line items", ["Number", "#", "Description", "Qty", "Unit price", "VAT rate", "Line total"],
           [(i["number"], n, it.get("description"), it.get("qty"), it.get("unit_price"), it.get("vat_rate"), it.get("line_total"))
            for i in data["invoices"] for n, it in enumerate(i.get("items") or [], start=1)])
    _sheet(wb, "Voided numbers", ["Series", "Number", "Reason", "Actor", "When"],
           [(g.get("series_code"), g["number"], g["reason"], g.get("actor"), g["created_at"]) for g in data["voided"]])
    _sheet(wb, "Chain", ["Checked", "OK", "Broken numbers", "Range start", "Range end", "Generated", "TIN"],
           [(data["chain"]["checked"], data["chain"]["ok"], ", ".join(data["chain"]["broken"]), start, end,
             data["generated_at"], data["profile"].get("tin"))])
    return _xlsx(wb, f"erca_audit_{start}_{end}.xlsx")


@router.get("/audit/export.json", name="erca_audit_json")
async def audit_json(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    start, end = _audit_range(request)
    data = erca_store.audit_export(cid, start, end)
    body = json.dumps(data, default=_json_default, indent=2)
    return Response(content=body, media_type="application/json",
                    headers={"Content-Disposition": f'attachment; filename="erca_audit_{start}_{end}.json"'})


# ── JSON API ────────────────────────────────────────────────────────────────
@router.get("/api/next-number/{series_code}", name="erca_api_next_number")
async def api_next_number(series_code: str, request: Request, user=Depends(login_required)):
    cid = current_company(request)
    n = erca_store.peek_next_number(cid, series_code)
    if n is None:
        return JSONResponse({"success": False, "error": "unknown series"}, status_code=404)
    return {"success": True, "series_code": series_code.upper(), "next_number": n}


@router.post("/api/invoices", name="erca_api_issue_invoice")
async def api_issue_invoice(request: Request, user=Depends(login_required)):
    cid = current_company(request)
    try:
        data = await request.json()
    except Exception:
        return JSONResponse({"success": False, "error": "Invalid JSON"}, status_code=400)
    series_code = data.pop("series_code", "INV")
    data["created_by"] = _actor(request)
    inv = erca_store.issue_invoice(cid, series_code, **data)
    if not inv:
        return JSONResponse({"success": False, "error": "issue failed"}, status_code=400)
    return JSONResponse({"success": True, "invoice": json.loads(json.dumps(inv, default=_json_default))})


@router.get("/api/withholding/preview", name="erca_api_withholding_preview")
async def api_withholding_preview(request: Request, user=Depends(login_required)):
    q = request.query_params
    calc = forms.withholding_for(q.get("gross") or 0, (q.get("has_tin") or "true").lower() in ("1", "true", "yes", "on"),
                                 q.get("transaction_type") or "goods")
    return JSONResponse(json.loads(json.dumps(calc, default=_json_default)))
